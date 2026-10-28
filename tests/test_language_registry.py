import contextlib
import io
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import numpy as np
from openpyxl import Workbook, load_workbook

import core_tools.language as languages
import pipeline
from core_tools.dialogue_context import DialogueContextAnalyzer
from core_tools.term_filter import TermFilter
from core_tools.term_scanner import TermScanner
from core_tools.tm_matcher import TMMatcher
from core_tools.tm_semantic import SemanticTMRetriever
from pipeline_service import PipelineService
from prompts import build


ROOT = Path(__file__).resolve().parents[1]
KEY = 'mdnew_forestconcert_mdnew_decoration_3_1'


class LanguageRegistryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.profiles = self.root / 'language'
        self.prompts = self.root / 'prompts'
        shutil.copytree(ROOT / 'profiles/language', self.profiles)
        shutil.copytree(ROOT / 'prompts', self.prompts)
        for name, value in [('LANGUAGE_DIR', self.profiles), ('PROMPTS_ROOT', self.prompts)]:
            p = patch.object(languages, name, value)
            p.start()
            self.addCleanup(p.stop)
        network = patch('requests.sessions.Session.request', side_effect=AssertionError('NO_NETWORK'))
        network.start()
        self.addCleanup(network.stop)
        self.config = self.root / 'models.json'
        self.config.write_text(json.dumps({k: {'api_url': 'https://example.invalid/' + k,
                                             'api_key': 'TEST_ONLY', 'model_name': 'test'}
                                         for k in ('llm', 'embedding', 'rerank')}))

    def catalog(self):
        return json.loads((self.profiles / '_catalog.json').read_text())

    def save_catalog(self, value):
        (self.profiles / '_catalog.json').write_text(json.dumps(value, ensure_ascii=False))

    def add_pair(self):
        # A synthetic third pair proves source handling does not choose its target.
        profile = {'source_lang': 'zh-CN', 'target_lang': 'de', 'columns': {},
                   'termbases': [], 'tms': [], 'prompt_source': 'zh-de', 'prompt_output': 'dist/zh-de'}
        (self.profiles / 'zh-de.json').write_text(json.dumps(profile))
        shutil.copytree(self.prompts / 'zh-en-US', self.prompts / 'zh-de')
        (self.prompts / 'zh-de/base.md').write_text('TARGET_DE_SENTINEL\n{{SCENE_TITLE}}\n{{SCENE_RULES}}\n')
        return languages.LanguagePair('zh', 'de')

    @staticmethod
    def scene(payload):
        row = payload['dialogue'][0]
        return {'scene_structure': '同段角色交流。', 'situation_tone': '围绕修复大门展开。',
                'evidence': [{'key': row['key'], 'quote': row['source_text']}],
                'material_evidence': [], 'warnings': []}

    def source(self):
        book = Workbook()
        book.active.append(['key', 'speaker', 'game_text_cn', 'game_text_de'])
        book.active.append([KEY, 'Chad', '请修复城堡大门。', 'KEEP_TARGET'])
        book.active.freeze_panes = 'C2'
        path = self.root / 'input.xlsx'
        book.save(path)
        book.close()
        return path

    def libraries(self):
        glossary = self.root / 'terms.csv'
        glossary.write_text('ZH,DE\n城堡,Schloss\n城堡大门,Schlosstor\n', encoding='utf-8')
        memory = self.root / 'memory.tmx'
        memory.write_text('<tmx><body><tu><tuv xml:lang="zh-CN"><seg>请修复城堡大门。</seg></tuv>'
                          '<tuv xml:lang="de-DE"><seg>Repariere das Schlosstor.</seg></tuv></tu></body></tmx>')
        return glossary, memory

    def test_only_registered_pairs_are_available_and_old_aliases_still_work(self):
        self.assertEqual(set(languages.language_profiles()), {'en-de', 'zh-en-US'})
        self.assertEqual(languages.LanguagePair('en_US', 'de-DE').key, 'en-de')
        zh = languages.LanguagePair('cn', 'en')
        self.assertEqual(zh.key, 'zh-en-US')
        self.assertTrue(zh.tm_language('zh-Hans-CN', source=True))
        self.assertFalse(zh.tm_language('en-GB'))
        with self.assertRaisesRegex(ValueError, '当前支持'):
            languages.LanguagePair('zh', 'de')
        with self.assertRaisesRegex(ValueError, '未登记'):
            languages.language_profile('../en-de')

    def test_one_pair_registration_reaches_api_cli_and_prompt_builder(self):
        pair = self.add_pair()
        import app
        import pipeline_api
        service = PipelineService(workspace=self.root / 'service')
        with patch.object(pipeline_api, 'service', service):
            client = app.app.test_client()
            options = client.get('/api/pipeline/bootstrap').json['language_pairs']
            option = next(o for o in options if o['key'] == 'zh-de')
            self.assertEqual((option['source_name'], option['target_name']), ('简体中文', '德语'))
            self.assertEqual(option['target_column'], 'game_text_de')
            result = client.post('/api/pipeline/projects', json={'name': '配置扩展测试', 'src': 'zh', 'tgt': 'de'})
            self.assertEqual(result.status_code, 200, result.json)
            self.assertEqual(result.json['project']['tgt'], 'de')
        with patch.object(pipeline, 'WORKSPACE', self.root / 'cli'):
            runner = pipeline.Pipeline('wc', 'zh-de', 'batch')
            self.assertEqual(languages.language_pair(runner.context()['language']).key, pair.key)
            with contextlib.redirect_stdout(io.StringIO()) as help_text, self.assertRaises(SystemExit) as exit_info:
                pipeline.main(['--help'])
            self.assertEqual(exit_info.exception.code, 0)
            self.assertIn('zh-de', help_text.getvalue())
        with patch.object(build, 'ROOT', self.prompts), contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(build.main(['--lang', 'zh-de']), 0)
        outputs = list((self.prompts / 'dist/zh-de').glob('*.md'))
        self.assertEqual(len(outputs), 4)
        self.assertTrue(all('TARGET_DE_SENTINEL' in p.read_text() for p in outputs))

    def test_target_changes_keep_chinese_matching_but_use_german_terms_and_rules(self):
        pair = self.add_pair()
        glossary, memory = self.libraries()
        scanner = TermScanner(glossary, language=pair)
        self.assertEqual([(m['source_text'], m['target_text']) for m in scanner.find_matches('请修复城堡大门。')],
                         [('城堡大门', ['Schlosstor'])])
        tm = TMMatcher(memory, language=pair)
        reference = tm.format_matches(tm.find_matches('请修复城堡大门。'))
        self.assertIn('【德】Repariere', reference)
        self.assertNotIn('【英】', reference)
        response = Mock(ok=True)
        response.json.return_value = {'choices': [{'message': {'content': json.dumps({
            'decisions': [{'id': 0, 'decision': 'keep', 'reason': '本句指城堡大门'}]})}}]}
        record = dict(source_text='请修复城堡大门。', context='', category='活动对白', key=KEY,
                      speaker='Chad', emotion='', candidates=['城堡大门 (Schlosstor)'])
        with patch('core_tools.term_filter.requests.post', return_value=response) as request:
            TermFilter(self.config, language=pair).filter_terms_with_llm(record)
        messages = request.call_args.kwargs['json']['messages']
        self.assertIn('简体中文→德语', messages[0]['content'])
        self.assertIn('德语格', messages[0]['content'])
        self.assertIn('中文连续文本', messages[0]['content'])
        self.assertNotIn('美式英语', messages[0]['content'])

    def test_configured_pair_runs_six_stages_without_rewriting_original_cells(self):
        pair = self.add_pair()
        glossary, memory = self.libraries()
        original = self.source()
        ctx = {'game': {'speakers': ['Chad'], 'key_rules': str(ROOT / 'profiles/game/wc_key_rules.json')},
               'language': {**pair.as_dict(), 'termbases': [str(glossary)], 'tms': [str(memory)]},
               'config_path': self.config, 'batch_dir': self.root / 'batch',
               'tm_cache_dir': self.root / 'tm_cache'}
        source = original
        with (patch.object(DialogueContextAnalyzer, '_request', side_effect=self.scene) as scene,
             patch.object(TermFilter, 'filter_terms_with_llm', side_effect=lambda row: [
                 {'term': t, 'decision': 'keep', 'reason': '测试'} for t in row['candidates']])):
            for key, _, filename, run in pipeline.STAGES:
                target = self.root / filename
                result = run(ctx, source, target)
                self.assertEqual(result['state'], 'completed', (key, result))
                source = target
        self.assertEqual(scene.call_count, 1)
        before, after = load_workbook(original), load_workbook(source)
        try:
            self.assertEqual([[c.value for c in row] for row in before.active],
                             [[c.value for c in row[:4]] for row in after.active])
            self.assertEqual(after.active.freeze_panes, 'C2')
            headers = {c.value: c.column for c in after.active[1]}
            self.assertIn('【场景结构】', after.active.cell(2, headers['Context']).value)
            self.assertIn('Schlosstor', after.active.cell(2, headers['生效术语']).value)
            self.assertIn('【德】', after.active.cell(2, headers['TM参考']).value)
        finally:
            before.close()
            after.close()

    def test_default_pair_is_shared_by_web_and_cli(self):
        catalog = self.catalog()
        catalog['default_pair'] = 'zh-en-US'
        self.save_catalog(catalog)
        service = PipelineService(workspace=self.root / 'default')
        self.assertEqual(service.bootstrap()['project']['tgt'], 'en-US')
        self.assertEqual(service.create_project('默认语向')['project']['src'], 'zh-CN')
        with patch.object(pipeline, 'WORKSPACE', self.root / 'cli'), contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(pipeline.main([str(self.source()), '--stages', 'normalize', '--batch', 'default']), 0)
        manifest = json.loads((self.root / 'cli/default/run.json').read_text())
        self.assertEqual(manifest['language'], 'zh-en-US')

    def test_changed_language_rules_invalidate_context_and_tm_caches(self):
        pair = languages.LanguagePair('zh', 'en')
        line = dict(key=KEY, speaker='Chad', emotion='', bubble='', source_text='请修复城堡大门。')
        with (patch.object(SemanticTMRetriever, 'embed', return_value=np.array([[1., .5]])) as embed,
              patch.object(DialogueContextAnalyzer, '_request', side_effect=self.scene) as scene):
            for version in range(3):
                if version == 1:
                    catalog = self.catalog()
                    catalog['sources']['zh-CN']['term_guidance'] += ' 配置变化。'
                    self.save_catalog(catalog)
                pair = languages.LanguagePair('zh', 'en')
                SemanticTMRetriever(self.config, self.root / 'cache', language=pair).prepare(['同句'], 'same-hash')
                DialogueContextAnalyzer(self.config, cache_dir=self.root / 'scenes', language=pair).analyze('g', [line], [], [])
            self.assertEqual(embed.call_count, 2)
            self.assertEqual(scene.call_count, 2)

    def test_legacy_project_and_preview_keep_their_language_when_default_changes(self):
        service = PipelineService(workspace=self.root / 'legacy')
        before = service.projects_file.read_bytes()
        catalog = self.catalog()
        catalog['default_pair'] = 'zh-en-US'
        self.save_catalog(catalog)
        book = Workbook()
        book.active.append(['key', 'game_text_en', 'game_text_de'])
        book.active.append(['title', 'Original English', 'KEEP'])
        path = self.root / 'english.xlsx'
        book.save(path)
        book.close()
        summary = service._summarize('normalize', 'wc', {}, 'normalize', {}, path, path, 0)
        self.assertEqual(summary['preview']['head'][2], 'game_text_en')
        self.assertEqual(summary['preview']['rows'][0][2], 'Original English')
        self.assertEqual(service.bootstrap()['project']['src'], 'en')
        self.assertEqual(service.projects_file.read_bytes(), before)

    def test_bad_config_fails_explicitly_instead_of_using_another_language(self):
        original = self.catalog()
        for field, value, message in [('strategy', 'not-implemented', '策略尚未实现'),
                                      ('negative_pattern', '[', '否定表达规则'),
                                      ('aliases', ['zh-CN', 'en'], '别名'),
                                      ('bm25_floor', -1, '检索或情绪')]:
            with self.subTest(field=field):
                catalog = json.loads(json.dumps(original))
                catalog['sources']['zh-CN'][field] = value
                self.save_catalog(catalog)
                with self.assertRaisesRegex(ValueError, message):
                    languages.language_profiles()
        self.save_catalog(original)
        profile = languages.language_profile('en-de')
        profile['prompt_output'] = '../../outside'
        with self.assertRaisesRegex(ValueError, 'prompts'):
            languages.prompt_locations(profile)
        with self.assertRaisesRegex(ValueError, '未实现'):
            languages.source_tokens('test', 'unknown')

    def test_default_columns_are_shared_and_prompt_outputs_cannot_collide(self):
        pair = self.add_pair()
        path = self.profiles / 'zh-de.json'
        profile = json.loads(path.read_text())
        profile['columns'] = {'source': '中文台词', 'target': '德文稿', 'emotion': '表情'}
        path.write_text(json.dumps(profile, ensure_ascii=False))
        pair = languages.LanguagePair('zh', 'de')
        option = next(o for o in languages.supported_language_pairs() if o['key'] == pair.key)
        self.assertEqual(option['source_column'], pair.aliases('source')[0])
        self.assertEqual(option['target_column'], pair.aliases('target')[0])
        self.assertEqual(option['emotion_columns'], ['表情'])
        self.assertEqual(languages.language_pair(pair, {'target': '自定义'}).aliases('target'), ('自定义',))
        profile['columns']['emotion'] = '中文台词'
        path.write_text(json.dumps(profile, ensure_ascii=False))
        with self.assertRaisesRegex(ValueError, '同一列'):
            languages.language_profiles()
        profile['columns'] = {}
        profile['prompt_output'] = 'dist/zh-en-US'
        path.write_text(json.dumps(profile))
        with self.assertRaisesRegex(ValueError, '输出目录.*重复'):
            languages.language_profiles()

    def test_existing_prompts_rebuild_without_changing_their_contents(self):
        before = {str(p.relative_to(self.prompts)): p.read_bytes() for p in (self.prompts / 'dist').rglob('*.md')}
        with patch.object(build, 'ROOT', self.prompts), contextlib.redirect_stdout(io.StringIO()):
            for key in languages.language_profiles():
                build.main(['--lang', key])
        after = {str(p.relative_to(self.prompts)): p.read_bytes() for p in (self.prompts / 'dist').rglob('*.md')}
        self.assertEqual(after, before)


if __name__ == '__main__':
    unittest.main()
