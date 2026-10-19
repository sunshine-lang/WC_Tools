import csv
import io
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import time
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core_tools.key_classifier import CATEGORIES, PENDING, KeyClassifier


class KeyClassifierTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)
        config = self.root / 'rules.json'
        config.write_text(json.dumps({
            CATEGORIES[0]: ['album_*'], CATEGORIES[1]: ['Ahalloween_set*_name'],
            CATEGORIES[2]: ['mdnew_*_decoration_name_*'], CATEGORIES[3]: ['story*'],
        }), encoding='utf-8')
        self.classifier = KeyClassifier(str(config))

    def tearDown(self):
        self.temp.cleanup()

    def workbook(self, rows, headers=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers or ['limit', 'key', 'game_text_cn', 'game_text_en', 'game_text_de'])
        for row in rows:
            sheet.append(row)
        path = self.root / 'input.xlsx'
        workbook.save(path)
        return path

    def run_classifier(self, path):
        task = {'state': 'processing'}
        self.classifier.process(str(path), str(self.root / 'output.xlsx'), {'test': task}, 'test')
        return task

    def test_mixed_dialogue_has_no_fixed_start_row_and_preserves_source(self):
        path = self.workbook([
            [None, 'Max', 'happy', '  <i>Hello!</i> {0}  ', None],
            [24, 'Ahalloween_set1_name', '秋日', 'Autumn', None],
            [None, None, None, None, None],
            [200, None, '宣传语', 'Join us!', None],
        ])
        workbook = load_workbook(path)
        workbook.active['D2'].font = Font(bold=True, color='123456')
        workbook.save(path)
        before = path.read_bytes()
        task = self.run_classifier(path)
        self.assertEqual(task['state'], 'completed')
        self.assertEqual([r['category'] for r in task['records']], [CATEGORIES[3], CATEGORIES[1], CATEGORIES[0]])
        self.assertEqual(task['records'][0]['speaker'], 'Max')
        self.assertEqual(task['records'][0]['key'], '')
        self.assertEqual(len({r['id'] for r in task['records']}), 3)
        result = load_workbook(task['result_file']).active
        self.assertEqual(result['D2'].value, '  <i>Hello!</i> {0}  ')
        self.assertEqual(result['D2'].font.color.rgb, '00123456')
        self.assertIsNone(result['F4'].value)
        self.assertEqual([cell.value for cell in result[1]], [
            'limit', 'key', 'game_text_cn', 'game_text_en', 'game_text_de', '分类',
        ])
        self.assertEqual(result['F2'].value, CATEGORIES[3])
        self.assertEqual(path.read_bytes(), before)

    def test_unknown_and_conflicting_rules_are_pending(self):
        rules = self.classifier.load_rules()
        self.assertEqual(self.classifier.classify('unknown', 'Text', '', '', rules)[0], PENDING)
        rules[CATEGORIES[3]].append('Ahalloween*')
        category, basis = self.classifier.classify('Ahalloween_set1_name', 'Autumn', '', '', rules)
        self.assertEqual(category, PENDING)
        self.assertIn('同时命中', basis)

    def test_explicit_speaker_fields_and_missing_english(self):
        path = self.workbook([
            ['new_dialogue_1', 'New Character', 'happy', 'Hello'],
            ['new_dialogue_2', 'New Character', 'happy', None],
        ], ['Key', 'Speaker', 'Animation', 'game_text_en'])
        task = self.run_classifier(path)
        self.assertEqual(task['records'][0]['category'], CATEGORIES[3])
        self.assertEqual(task['records'][0]['key'], 'new_dialogue_1')
        self.assertEqual(task['records'][1]['category'], PENDING)

    def test_formula_source_is_flagged_and_not_replaced(self):
        task = self.run_classifier(self.workbook([[30, 'album_title', '标题', '=1+1', None]]))
        self.assertEqual(task['records'][0]['category'], PENDING)
        self.assertEqual(load_workbook(task['result_file']).active['D2'].value, '=1+1')

    def test_csv_keeps_literal_formulas_and_multiline_text(self):
        path = self.root / 'input.csv'
        with path.open('w', encoding='utf-8-sig', newline='') as file:
            writer = csv.writer(file)
            writer.writerow(['Key', 'en'])
            writer.writerow(['album_title', '=1+1'])
            writer.writerow(['album_desc', '  line one\nline two  '])
        task = self.run_classifier(path)
        sheet = load_workbook(task['result_file']).active
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertEqual(sheet['B2'].value, '=1+1')
        self.assertEqual(sheet['B3'].value, '  line one\nline two  ')

    def test_missing_source_header_returns_useful_error(self):
        task = self.run_classifier(self.workbook([['hello']], ['cn']))
        self.assertEqual(task['state'], 'error')
        self.assertIn('game_text_en', task['status'])

    def test_saved_rules_take_effect_and_invalid_rules_preserve_config(self):
        rules = self.classifier.load_rules()
        rules[CATEGORIES[1]].append('new_card_*')
        self.classifier.save_rules(rules)
        reloaded = KeyClassifier(self.classifier.config_path)
        self.assertEqual(reloaded.classify('new_card_1', 'Card', '', '', reloaded.load_rules())[0], CATEGORIES[1])
        before = Path(self.classifier.config_path).read_bytes()
        with self.assertRaises(ValueError):
            self.classifier.save_rules({'bad': []})
        self.assertEqual(Path(self.classifier.config_path).read_bytes(), before)

    def test_http_upload_preview_and_download(self):
        import app as server
        original_folder = server.app.config['UPLOAD_FOLDER']
        original_classifier = server.key_classifier
        server.app.config['UPLOAD_FOLDER'] = str(self.root)
        server.key_classifier = self.classifier
        task_id = None
        try:
            client = server.app.test_client()
            self.assertEqual(client.post('/api/key_rules', json={'bad': []}).status_code, 400)
            rules = client.get('/api/key_rules').json
            rules[CATEGORIES[1]].append('custom_*')
            self.assertEqual(client.post('/api/key_rules', json=rules).status_code, 200)
            self.assertIn('custom_*', client.get('/api/key_rules').json[CATEGORIES[1]])
            self.assertEqual(client.post('/upload/key_classifier').status_code, 400)
            path = self.workbook([[24, 'unknown', '名称', 'Text', None]])
            response = client.post('/upload/key_classifier', data={'file': (io.BytesIO(path.read_bytes()), 'input.xlsx')})
            task_id = response.json['task_id']
            for _ in range(50):
                progress = client.get('/progress/' + task_id).json
                if progress['state'] != 'processing':
                    break
                time.sleep(.02)
            self.assertEqual(progress['state'], 'completed')
            self.assertNotIn('records', progress)
            result = client.get('/api/classification/' + task_id).json
            self.assertEqual(len(result['records']), 1)
            self.assertEqual(client.post('/api/classification/' + task_id, json={}).status_code, 405)
            download = client.get('/download/' + task_id)
            self.assertEqual(download.status_code, 200)
            sheet = load_workbook(io.BytesIO(download.data)).active
            self.assertEqual(sheet['F1'].value, '分类')
            self.assertEqual(sheet['F2'].value, PENDING)
            self.assertEqual(sheet.max_column, 6)
            download.close()
        finally:
            server.app.config['UPLOAD_FOLDER'] = original_folder
            server.key_classifier = original_classifier
            if task_id:
                server.tasks.pop(task_id, None)


if __name__ == '__main__':
    unittest.main()
