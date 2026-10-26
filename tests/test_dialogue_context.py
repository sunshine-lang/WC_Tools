import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch

import requests
from openpyxl import Workbook, load_workbook

from core_tools.context_processor import ContextProcessor
from core_tools.dialogue_context import DialogueContextAnalyzer, SYSTEM_PROMPT
from core_tools.key_classifier import CATEGORIES
from pipeline import run_context


PREFIX = 'mdnew_forestconcert_mdnew_decoration_'


class DialogueContextTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.config = self.root / 'config.json'
        self.config.write_text(json.dumps({'llm': {'api_url': 'https://example.invalid/chat',
                                                  'api_key': 'TEST_SECRET', 'model_name': 'test-model'}}))
        self.analyzer = DialogueContextAnalyzer(self.config)
        self.network = patch('requests.sessions.Session.request', side_effect=AssertionError('NETWORK_DISABLED'))
        self.network.start()
        self.addCleanup(self.temp.cleanup)
        self.addCleanup(self.network.stop)

    def source(self, rows=None):
        book = Workbook()
        sheet = book.active
        sheet.title = '对白'
        sheet.append(['key', 'note', 'speaker', 'animation', 'bubble_type', 'game_text_cn',
                      'game_text_en', 'game_text_de', '分类'])
        for row in rows if rows is not None else [
            ['Max', PREFIX + '3_10', 'Max', '', '2.0', 'thinking', 'We found the missing key!', 'TARGET_SECRET', CATEGORIES[3]],
            ['Christine', PREFIX + '3_2', 'Christine', 'happy', '1.0', '', 'Let us look for the key.', '', CATEGORIES[3]],
            [PREFIX + '5_1', '', 'Chad', '', '', '', 'Another scene.', '', CATEGORIES[3]],
            [PREFIX + 'name_3', '', '', '', '', '', 'Music Box', '', CATEGORIES[2]],
            ['cat_multiboards_title', '', '', '', '', '', 'Cat event', '', CATEGORIES[0]],
            ['cat_multiboards_pre_desc', '', '', '', '', '', 'Play with cats', '', CATEGORIES[0]],
        ]:
            sheet.append(row)
        sheet.freeze_panes = 'C2'
        book.create_sheet('说明')['A1'] = 'UNCHANGED'
        path = self.root / 'classified.xlsx'
        book.save(path)
        book.close()
        return path

    @staticmethod
    def response(payload):
        line = payload['dialogue'][0]
        return {'scene_structure': '同组角色围绕一个问题展开连续对白。',
                'situation_tone': '语气由提出问题转为回应，翻译时保留互动与情绪变化。',
                'evidence': [{'key': line['key'], 'quote': line['source_text']}],
                'material_evidence': [], 'warnings': []}

    def run_scene(self, path, result=None):
        task = {}
        with patch.object(DialogueContextAnalyzer, '_request', side_effect=result or self.response) as request:
            output = run_context({'game': {}, 'config_path': self.config}, path, self.root / 'output.xlsx', task)
        return output, request

    def test_pipeline_reads_whole_group_once_and_preserves_non_dialogue_and_source(self):
        path = self.source()
        before = path.read_bytes()
        original = load_workbook(path)
        values = [[c.value for c in row] for row in original.active]
        task, request = self.run_scene(path)
        self.assertEqual(task['state'], 'completed', task.get('status'))
        self.assertEqual(request.call_count, 2)
        payload = request.call_args_list[0].args[0]
        self.assertEqual([r['key'] for r in payload['dialogue']], [PREFIX + '3_2', PREFIX + '3_10'])
        self.assertEqual([r['bubble_type'] for r in payload['dialogue']], ['1', '2'])
        self.assertEqual(payload['dialogue'][1]['emotion'], 'thinking')
        self.assertEqual([r['key'] for r in payload['references']], [PREFIX + 'name_3'])
        self.assertNotIn('TARGET_SECRET', json.dumps(payload))
        self.assertNotIn('Another scene.', json.dumps(payload))
        self.assertTrue(payload['input_warnings'])
        records = task['records']
        self.assertEqual(records[0]['context'], records[1]['context'])
        self.assertTrue(records[0]['context'].startswith('【场景结构】\n'))
        self.assertIn('\n\n【情境与基调】\n', records[0]['context'])
        self.assertNotIn('[关联语境参考]', records[0]['context'])
        self.assertIn('We found the missing key!', records[3]['context'])
        self.assertIn('[关联语境参考]', records[4]['context'])
        self.assertIn('未提供气泡类型', records[2]['warning'])
        self.assertEqual(task['dialogue_model_requests'], 2)
        output = load_workbook(task['result_file'])
        self.assertEqual([[c.value for c in row[:9]] for row in output.active], values)
        self.assertEqual(output.active.freeze_panes, 'C2')
        self.assertEqual(output['说明']['A1'].value, 'UNCHANGED')
        self.assertEqual(path.read_bytes(), before)
        original.close()
        output.close()

    def test_unknown_group_and_conflicting_or_duplicate_source_do_not_call_model(self):
        row = ['Max', '', 'Max', '', '', '', 'Unknown scene.', '', CATEGORIES[3]]
        cases = [
            [row],
            [[PREFIX + '3_1', '', 'Max', '', '', '', 'One', '', CATEGORIES[3]],
             [PREFIX + '3_1', '', 'Max', '', '', '', 'Different', '', CATEGORIES[3]]],
            [[PREFIX + '3_1', '', 'Max', '', '', '', 'One', '', CATEGORIES[3]],
             [PREFIX + '3_2', '', 'Max', '', '', '', '=1+1', '', CATEGORIES[3]]],
        ]
        for rows in cases:
            with self.subTest(rows=rows):
                task, request = self.run_scene(self.source(rows))
                self.assertEqual(task['state'], 'completed', task.get('status'))
                request.assert_not_called()
                self.assertTrue(all(r['context'] == '' for r in task['records']))

    def test_model_failure_does_not_publish_a_partial_workbook(self):
        def fail_second(payload):
            if '/5' in payload['group']:
                raise ValueError('模拟场景分析失败')
            return self.response(payload)
        task, _ = self.run_scene(self.source(), fail_second)
        self.assertEqual(task['state'], 'error')
        self.assertNotIn('result_file', task)
        self.assertFalse((self.root / 'output.xlsx').exists())

    def test_retry_reuses_validated_scenes_and_material_changes_invalidate_cache(self):
        path = self.source()
        def fail_second(payload):
            if payload['group'].endswith('/5'):
                raise ValueError('模拟第二组失败')
            return self.response(payload)
        failed, first = self.run_scene(path, fail_second)
        self.assertEqual(failed['state'], 'error')
        self.assertEqual(first.call_count, 2)
        finished, retry = self.run_scene(path)
        self.assertEqual(finished['state'], 'completed', finished.get('status'))
        self.assertEqual(retry.call_count, 1)
        self.assertEqual(finished['dialogue_cached_scenes'], 1)
        self.assertEqual(finished['dialogue_model_requests'], 1)
        with patch.object(DialogueContextAnalyzer, '_request', side_effect=self.response) as changed:
            result = run_context({'game': {}, 'config_path': self.config, 'materials': {'cast': '更新后的人设'}},
                                 path, self.root / 'updated.xlsx')
        self.assertEqual(result['state'], 'completed')
        self.assertEqual(changed.call_count, 2)
        self.assertEqual(result['dialogue_cached_scenes'], 0)

    def test_invalid_model_quote_is_saved_for_diagnosis_but_never_reused(self):
        def invalid(payload):
            result = self.response(payload)
            result['material_evidence'] = [{'material': 'cast', 'quote': '伪造的引用'}]
            return result
        path = self.source()
        failed, _ = self.run_scene(path, invalid)
        self.assertEqual(failed['state'], 'error')
        diagnostics = list((self.root / 'context_scenes').glob('*.failed.json'))
        self.assertEqual(len(diagnostics), 1)
        saved = json.loads(diagnostics[0].read_text())
        self.assertIn('不存在的人物资料', saved['error'])
        self.assertNotIn('TEST_SECRET', diagnostics[0].read_text())
        finished, retried = self.run_scene(path)
        self.assertEqual(finished['state'], 'completed')
        self.assertEqual(retried.call_count, 2)
        self.assertFalse(list((self.root / 'context_scenes').glob('*.failed.json')))

    def test_non_dialogue_does_not_require_model_config(self):
        path = self.source([
            ['cat_multiboards_title', '', '', '', '', '', 'Cats', '', CATEGORIES[0]],
            ['cat_multiboards_pre_desc', '', '', '', '', '', 'Play with cats', '', CATEGORIES[0]],
        ])
        task = run_context({'game': {}, 'config_path': self.root / 'missing.json'}, path, self.root / 'output.xlsx')
        self.assertEqual(task['state'], 'completed', task.get('status'))
        self.assertEqual(task['dialogue_model_requests'], 0)

    def test_analysis_validates_quotes_and_shows_uncertainty_outside_context(self):
        line = {'key': PREFIX + '3_1', 'speaker': 'Max_remote', 'emotion': 'thinking', 'bubble': '1',
                'source_text': 'What should I say?', 'source': 'input.xlsx / 对白!2', 'group': 'forestconcert/dialogue/3'}
        good = self.response({'dialogue': [line]})
        good['warnings'] = ['气泡字段与内心独白的解释存在冲突，请人工核对。']
        with patch.object(self.analyzer, '_request', return_value=good) as request:
            result = self.analyzer.analyze(line['group'], [line], [], [])
        self.assertTrue(result['warnings'])
        self.assertNotIn(result['warnings'][0], result['context'])
        self.assertEqual(request.call_args.args[0]['dialogue'][0]['speaker'], 'Max_remote')
        for replacement in [[], [{'key': 'fake', 'quote': 'What'}],
                            [{'key': line['key'], 'quote': 'Invented source'}]]:
            with self.subTest(replacement=replacement), patch.object(self.analyzer, '_request',
                    return_value=good | {'evidence': replacement}):
                with self.assertRaises(ValueError):
                    self.analyzer.analyze(line['group'], [line], [], [])
        self.assertIn('_remote 后缀不证明', SYSTEM_PROMPT)
        self.assertIn('未提供的年龄', SYSTEM_PROMPT)

    def test_full_character_documents_reach_scene_model_and_quotes_are_validated(self):
        materials = {'cast': '# 人物人设\n' + '完整人物小传。' * 400 + '\n人物尾部依据。',
                     'relationships': '# 人物关系\n' + '完整人物关系。' * 400 + '\n关系尾部依据。',
                     'world': '本游戏背景。', 'style': '当前语种规范。', 'secret': 'DO_NOT_SEND'}
        payloads = []
        def reply(payload):
            payloads.append(payload)
            result = self.response(payload)
            result['material_evidence'] = [{'material': 'cast', 'quote': '人物尾部依据。'},
                                           {'material': 'relationships', 'quote': '关系尾部依据。'}]
            return result
        with patch.object(DialogueContextAnalyzer, '_request', side_effect=reply):
            task = run_context({'game': {}, 'config_path': self.config, 'materials': materials},
                               self.source(), self.root / 'full.xlsx')
        self.assertEqual(task['state'], 'completed', task.get('status'))
        for payload in payloads:
            actual = {key: item['content'] for item in payload['character_materials'] for key in item['keys']}
            self.assertEqual(actual, {k: v for k, v in materials.items() if k != 'secret'})
            self.assertNotIn('DO_NOT_SEND', json.dumps(payload))
        self.assertEqual(len(task['dialogue_scenes'][0]['material_evidence']), 2)
        analyzer = DialogueContextAnalyzer(self.config, materials=materials)
        line = dict(key='scene_1', speaker='Max', emotion='', bubble='', source_text='Hi', source='file', group='g')
        for bad in [{'material': 'cast', 'quote': '不存在的人设原文'},
                    {'material': 'missing', 'quote': '人物尾部依据。'}]:
            result = self.response({'dialogue': [line]}) | {'material_evidence': [bad]}
            with self.subTest(bad=bad), patch.object(analyzer, '_request', return_value=result):
                with self.assertRaisesRegex(ValueError, '不存在的人物资料'):
                    analyzer.analyze('g', [line], [], [])

    def test_duplicate_material_is_sent_once_with_both_labels(self):
        value = '同时包含人物人设和关系的完整资料。'
        analyzer = DialogueContextAnalyzer(self.config, materials={'cast': value, 'relationships': value})
        self.assertEqual(len(analyzer.materials_payload), 1)
        self.assertEqual(analyzer.materials_payload[0]['keys'], ['cast', 'relationships'])
        self.assertEqual(analyzer.materials_payload[0]['content'], value)
        self.assertTrue(any('相同' in warning for warning in analyzer.material_warnings))

    def test_material_quote_tolerates_markdown_layout_but_not_changed_words(self):
        source = '**译名提示**：`Tadpole` 对应 `Krümel`。\n> 甲**不**是乙的父亲。'
        matches = DialogueContextAnalyzer._material_quote_matches
        self.assertTrue(matches('Tadpole 对应 Krümel。', source))
        self.assertTrue(matches('Tadpole 对应 Krümel。 甲不是乙的父亲。', source))
        self.assertFalse(matches('Tadpole 对应 Maus。', source))
        self.assertFalse(matches('甲是乙的父亲。', source))
        self.assertFalse(matches('> \n> ', source))
        self.assertTrue(matches('Max 是女孩', 'Max 是**女孩**（Maxine）'))
        self.assertFalse(matches('Max 是男孩', 'Max 是**女孩**（Maxine）'))

    def test_saved_response_is_revalidated_before_promotion_to_success_cache(self):
        value = '`Tadpole` 对应 `Krümel`。'
        analyzer = DialogueContextAnalyzer(self.config, materials={'relationships': value}, cache_dir=self.root / 'cache')
        line = dict(key='scene_1', speaker='Max', emotion='', bubble='', source_text='Hi', group='g')
        response = self.response({'dialogue': [line]}) | {
            'material_evidence': [{'material': 'relationships', 'quote': 'Tadpole 对应 Krümel。'}]}
        with patch.object(analyzer, '_request', return_value=response), patch.object(analyzer, '_material_quote_matches', return_value=False):
            with self.assertRaises(ValueError):
                analyzer.analyze('g', [line], [], [])
        with patch.object(analyzer, '_request', side_effect=AssertionError('No new API request')):
            result = analyzer.analyze('g', [line], [], [])
        self.assertTrue(result['cache_hit'])
        self.assertFalse(list((self.root / 'cache').glob('*.failed.json')))

    def test_transport_retries_transient_failures_and_rejects_auth_and_malformed_response(self):
        payload = {'dialogue': [{'key': 'x', 'source_text': 'Hi'}]}
        content = self.response(payload)
        response = Mock(status_code=200, json=lambda: {'choices': [{'message': {'content': json.dumps(content)}}]})
        with patch('core_tools.dialogue_context.requests.post', side_effect=[requests.Timeout('TEST_SECRET'),
                Mock(status_code=503), response]) as post, patch('core_tools.dialogue_context.time.sleep'):
            self.assertEqual(self.analyzer._request(payload), content)
            self.assertEqual(post.call_count, 3)
            self.assertFalse(post.call_args.kwargs['allow_redirects'])
        with patch('core_tools.dialogue_context.requests.post', return_value=Mock(status_code=401)) as post:
            with self.assertRaisesRegex(ValueError, 'HTTP 401') as error:
                self.analyzer._request(payload)
            self.assertNotIn('TEST_SECRET', str(error.exception))
            self.assertEqual(post.call_count, 1)
        with patch('core_tools.dialogue_context.requests.post', return_value=Mock(status_code=200, json=lambda: {})):
            with self.assertRaisesRegex(ValueError, 'JSON'):
                self.analyzer._request(payload)


if __name__ == '__main__':
    unittest.main()
