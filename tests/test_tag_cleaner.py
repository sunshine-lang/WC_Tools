import csv
import io
import re
import time
import unittest
from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core_tools.tag_cleaner import TagCleaner
from core_tools.qa_checker import QAChecker


class TagCleanerTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def source(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '活动ui'
        sheet.append(['limit', 'key', 'game_text_cn', 'game_text_en', 'game_text_de', 'Context'])
        sheet.append([100, 'a', '开门', 'Open <b>the door</b> now.', 'Öffne  <b>die Tür</b>  jetzt.', 'KEEP_CONTEXT'])
        sheet.append([None, 'b', '', 'Formula', '=CONCAT("A")', None])
        sheet.append([None, 'c', '', 'Plain', 'Ohne Tags', None])
        sheet['E2'].font = Font(bold=True, color='112233')
        sheet['E2'].number_format = '@'
        sheet.freeze_panes = 'B2'
        sheet.column_dimensions['E'].width = 60
        sheet.row_dimensions[2].height = 42
        other = book.create_sheet('待确认')
        other['A1'] = 'KEEP_OTHER_SHEET'
        other.merge_cells('A1:B1')
        path = self.root / 'translated.xlsx'
        book.save(path)
        book.close()
        return path

    def run_clean(self, path, column='game_text_de', name=None):
        task = {'state': 'processing', 'input_name': name or path.name}
        TagCleaner(column).process(path, self.root / ('cleaned' + path.suffix), {'t': task}, 't')
        return task

    def test_german_and_english_preserve_word_boundaries_tags_parameters_and_newlines(self):
        source = 'Hol  <color=#fff>  {0}  </color>  <sprite name="icon   coin">  Münzen!\n<b> Jetzt </b>\\nWeiter'
        expected = 'Hol <color=#fff> {0} </color> <sprite name="icon   coin"> Münzen!\n<b> Jetzt </b>\\nWeiter'
        for column in ('game_text_de', 'game_text_en'):
            with self.subTest(column=column):
                cleaner = TagCleaner(column)
                self.assertEqual(cleaner.clean_text(source), expected)
                self.assertEqual(cleaner.clean_text(expected), expected)
                self.assertEqual(re.findall(r'<[^>]+>', source), re.findall(r'<[^>]+>', expected))
                plain = 'Öffne <b>die Tür</b> jetzt.'
                self.assertEqual(cleaner.clean_text(plain), plain)
                self.assertEqual(cleaner.clean_text('Ohne  Tags'), 'Ohne  Tags')

    def test_chinese_matches_ro_rule_without_touching_tag_contents(self):
        source = '你好  <color=#fff>  世界 </color> ！\n<sprite name="x  y"> '
        for column in ('game_text_cn', 'cn'):
            cleaner = TagCleaner(column)
            self.assertEqual(cleaner.clean_text(source), re.sub(r'\s*(<[^>]+>)\s*', r'\1', source))
            self.assertIn('<sprite name="x  y">', cleaner.clean_text(source))

    def test_workbook_changes_only_selected_text_preserves_formulas_styles_and_other_sheets(self):
        path = self.source()
        digest = sha256(path.read_bytes()).hexdigest()
        task = self.run_clean(path)
        self.assertEqual(task['state'], 'completed', task)
        self.assertEqual((task['total'], task['changed'], task['skipped_formulas']), (2, 1, 1))
        self.assertEqual(sha256(path.read_bytes()).hexdigest(), digest)
        before, after = load_workbook(path), load_workbook(task['result_file'])
        try:
            self.assertEqual(before.sheetnames, after.sheetnames)
            for sheet in before:
                updated = after[sheet.title]
                self.assertEqual(sheet.freeze_panes, updated.freeze_panes)
                self.assertEqual(str(sheet.merged_cells), str(updated.merged_cells))
                self.assertEqual(sheet.row_dimensions[2].height, updated.row_dimensions[2].height)
                for row in sheet:
                    for cell in row:
                        new = updated.cell(cell.row, cell.column)
                        expected = 'Öffne <b>die Tür</b> jetzt.' if sheet.title == '活动ui' and cell.coordinate == 'E2' else cell.value
                        self.assertEqual((expected, cell.data_type, cell._style), (new.value, new.data_type, new._style))
            self.assertEqual(after.active.column_dimensions['E'].width, 60)
            self.assertEqual(after.active['E3'].data_type, 'f')
            self.assertEqual(task['records'][0]['row'], 2)
        finally:
            before.close(); after.close()

    def test_csv_preserves_na_empty_strings_quoted_newlines_and_other_columns(self):
        path = self.root / 'input.csv'
        rows = [['key', 'game_text_en', 'game_text_de', 'other'],
                ['NA', 'KEEP', 'Hol  <b>{0}</b>  Münzen!\nJetzt', 'NA'],
                ['b', '', '', 'KEEP_EMPTY_TARGET']]
        with path.open('w', encoding='utf-8-sig', newline='') as file:
            csv.writer(file).writerows(rows)
        task = self.run_clean(path)
        self.assertEqual(task['state'], 'completed', task)
        with open(task['result_file'], encoding='utf-8-sig', newline='') as file:
            after = list(csv.reader(file))
        rows[1][2] = 'Hol <b>{0}</b> Münzen!\nJetzt'
        self.assertEqual(after, rows)

    def test_missing_duplicate_column_bad_encoding_and_done_fail_without_publishing(self):
        path = self.source()
        self.assertEqual(self.run_clean(path, 'missing')['state'], 'error')
        self.assertEqual(self.run_clean(path, name='input_done.xlsx')['state'], 'error')
        bad = self.root / 'broken.xlsx'; bad.write_bytes(b'bad')
        self.assertEqual(self.run_clean(bad)['state'], 'error')
        bad = self.root / 'bad.csv'; bad.write_bytes(b'\xff\xff')
        task = self.run_clean(bad)
        self.assertEqual(task['state'], 'error')
        self.assertIn('UTF-8', task['status'])
        duplicate = self.root / 'duplicate.csv'
        duplicate.write_text('game_text_de,game_text_de\na,b\n', encoding='utf-8')
        task = self.run_clean(duplicate)
        self.assertEqual(task['state'], 'error')
        self.assertIn('重复', task['status'])
        self.assertNotIn('result_file', task)

    def test_empty_column_and_overwrite_are_rejected(self):
        with self.assertRaises(ValueError):
            TagCleaner('')
        path = self.source()
        task = {'state': 'processing'}
        TagCleaner('game_text_de').process(path, path, {'t': task}, 't')
        self.assertEqual(task['state'], 'error')
        self.assertIn('不能覆盖', task['status'])

    def test_http_default_column_preview_download_and_qa_handoff(self):
        import app as server
        task_id = None
        path = self.source()
        with patch.dict(server.app.config, {'UPLOAD_FOLDER': str(self.root)}):
            client = server.app.test_client()
            html = client.get('/tag_clean').get_data(as_text=True)
            self.assertLess(html.index('href="/pipeline"'), html.index('href="/tag_clean"'))
            self.assertLess(html.index('href="/tag_clean"'), html.index('href="/qa"'))
            self.assertEqual(client.post('/upload/tag_clean').status_code, 400)
            self.assertEqual(client.post('/upload/tag_clean', data={'file': (io.BytesIO(b'x'), 'bad.txt')}).status_code, 400)
            self.assertEqual(client.post('/upload/tag_clean', data={'file': (io.BytesIO(path.read_bytes()), 'input_done.xlsx')}).status_code, 400)
            try:
                response = client.post('/upload/tag_clean', data={'file': (io.BytesIO(path.read_bytes()), '译后文件.xlsx')})
                self.assertEqual(response.status_code, 200)
                task_id = response.json['task_id']
                for _ in range(100):
                    progress = client.get('/progress/' + task_id).json
                    if progress['state'] != 'processing':
                        break
                    time.sleep(.02)
                self.assertEqual(progress['state'], 'completed', progress)
                self.assertNotIn('records', progress)
                preview = client.get('/api/tag_clean/' + task_id).json
                self.assertEqual((preview['column'], preview['changed']), ('game_text_de', 1))
                response = client.get('/download/' + task_id)
                self.assertIn('_cleaned.xlsx', response.headers['Content-Disposition'])
                output = self.root / 'download.xlsx'; output.write_bytes(response.data)
                response.close()
                qa_task = {'state': 'processing'}
                QAChecker('game_text_en', 'game_text_de').process(output, self.root / 'qa.xlsx', {'q': qa_task}, 'q')
                self.assertEqual(qa_task['state'], 'completed', qa_task)
                self.assertEqual(qa_task['total'], 3)
                self.assertEqual(client.get('/api/tag_clean/missing').status_code, 404)
            finally:
                if task_id:
                    server.tasks.pop(task_id, None)


if __name__ == '__main__':
    unittest.main()
