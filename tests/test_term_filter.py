import io
import json
import time
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core_tools.term_filter import FILTER_COLUMNS, FILTER_POLICY, TermFilter


class TermFilterTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.config = self.root / 'llm.json'
        self.config.write_text(json.dumps({'llm': {'api_key': 'TEST_SECRET', 'api_url': 'https://example.test/chat',
                                                 'model_name': 'test-model', 'max_workers': 2}}))
        self.tool = TermFilter(self.config)

    def tearDown(self):
        self.temp.cleanup()

    def source(self, rows=None, headers=None):
        book = Workbook()
        sheet = book.active
        sheet.title = '活动ui'
        sheet.append(headers or ['key', 'game_text_en', 'game_text_de', '分类', 'Context', '关联术语'])
        for row in rows if rows is not None else [
                ['song', 'Songs in a minor key.', 'KEEP_TARGET', '活动对白', 'Music puzzle', 'key (Schlüssel)'],
                [None, None, None, None, None, None],
                ['found', 'With a key inside!', '', '活动对白', '', 'key (Schlüssel)'],
                ['new', 'Brand new text.', '', '集卡名称', '', '']]:
            sheet.append(row)
        sheet['B2'].font = Font(bold=True, color='112233')
        sheet.freeze_panes = 'B2'
        sheet.row_dimensions[2].height = 42
        sheet.column_dimensions['A'].width = 25
        book.create_sheet('说明')['A1'] = 'KEEP_OTHER_SHEET'
        path = self.root / 'step3.xlsx'
        book.save(path)
        book.close()
        return path

    def run_filter(self, path, name='step3.xlsx'):
        task = {'state': 'processing', 'input_name': name}
        self.tool.process(path, self.root / 'result.xlsx', {'t': task}, 't')
        return task

    @staticmethod
    def response(items):
        return Mock(ok=True, json=lambda: {'choices': [{'message': {'content': json.dumps({'decisions': items})}}]})

    @staticmethod
    def classify(record):
        return [{'term': term, 'decision': 'exclude' if 'minor key' in record['english'] else 'keep',
                 'reason': '音乐调性，不是钥匙' if 'minor key' in record['english'] else '本句指实际钥匙'}
                for term in record['candidates']]

    def test_model_receives_english_context_and_ids_without_target_or_secret_leak(self):
        record = dict(english='minor key', context='Music puzzle', category='活动对白', key='', speaker='Max', emotion='happy',
                      candidates=['key (Schlüssel)', 'Max (Max)'])
        response = self.response([{'id': 1, 'decision': 'keep', 'reason': '人名'},
                                  {'id': 0, 'decision': 'exclude', 'reason': '音乐调性'}])
        with patch('core_tools.term_filter.requests.post', return_value=response) as post:
            result = self.tool.filter_terms_with_llm(record)
        self.assertEqual([d['term'] for d in result], record['candidates'])
        self.assertEqual([d['decision'] for d in result], ['exclude', 'keep'])
        payload = json.loads(post.call_args.kwargs['json']['messages'][1]['content'])
        self.assertEqual(payload['context'], 'Music puzzle')
        self.assertEqual(payload['english'], 'minor key')
        self.assertEqual(payload['speaker'], 'Max')
        self.assertNotIn('TEST_SECRET', json.dumps(payload))
        self.assertNotIn('TEST_SECRET', json.dumps(self.tool.status()))

    def test_customer_note_and_world_material_reach_model_without_changing_source(self):
        key = 'mdnew_forestconcert_mdnew_decoration_3_1'
        rows = [[key, 'Max', 'Max', 'With a key inside!', 'KEEP_TARGET', '活动对白', 'Scene three', 'key (Schlüssel)']]
        path = self.source(rows, ['note', 'key', 'speaker', 'game_text_en', 'game_text_de', '分类', 'Context', '关联术语'])
        self.tool = TermFilter(self.config, materials={'world': 'WORLD_SENTINEL', 'cast': 'CAST_SENTINEL', 'style': 'STYLE_SENTINEL'})
        response = self.response([{'id': 0, 'decision': 'keep', 'reason': '本句指钥匙'}])
        with patch('core_tools.term_filter.requests.post', return_value=response) as post:
            task = self.run_filter(path)
        self.assertEqual(task['state'], 'completed', task.get('status'))
        messages = post.call_args.kwargs['json']['messages']
        self.assertIn('[游戏背景]\nWORLD_SENTINEL', messages[1]['content'])
        self.assertIn('CAST_SENTINEL', messages[1]['content'])
        self.assertIn('STYLE_SENTINEL', messages[1]['content'])
        payload = json.loads(messages[-1]['content'])
        self.assertEqual((payload['key'], payload['speaker']), (key, 'Max'))
        self.assertNotIn('KEEP_TARGET', json.dumps(messages))
        result = load_workbook(task['result_file'])
        self.assertEqual([cell.value for cell in result.active[2]][:7], rows[0][:7])
        result.close()

    def test_full_background_and_separate_character_materials_reach_model(self):
        materials = {key: f'{key}_START\n' + '完整资料。' * 400 + f'\n{key}_END'
                     for key in ('world', 'cast', 'relationships', 'style')}
        tool = TermFilter(self.config, materials=materials)
        record = dict(english='A key', context='', category='活动对白', key='scene_1',
                      speaker='Max', emotion='', candidates=['key (Schlüssel)'])
        response = self.response([{'id': 0, 'decision': 'keep', 'reason': '本句指钥匙'}])
        with patch('core_tools.term_filter.requests.post', return_value=response) as post:
            tool.filter_terms_with_llm(record)
        background = post.call_args.kwargs['json']['messages'][1]['content']
        for key, label in (('world', '游戏背景'), ('cast', '人物人设'),
                           ('relationships', '人物关系'), ('style', '风格指南')):
            self.assertIn(f'[{label}]\n{materials[key]}', background)
            self.assertIn(f'{key}_END', background)

    def test_incomplete_duplicate_invented_and_empty_reason_responses_rejected(self):
        record = dict(english='Hello', context='', category='', key='', speaker='', emotion='', candidates=['Chad (Chad)'])
        for items in [[], [{'id': 2, 'decision': 'keep', 'reason': 'new'}],
                      [{'id': True, 'decision': 'keep', 'reason': 'bad id'}],
                      [{'id': 0, 'decision': 'keep', 'reason': ''}],
                      [{'id': 0, 'decision': 'new', 'reason': 'bad status'}],
                      [{'id': 0, 'decision': 'uncertain', 'reason': '无法判断'}],
                      [{'id': 0, 'decision': 'keep', 'reason': 'ok'}, {'id': 0, 'decision': 'keep', 'reason': 'duplicate'}]]:
            with self.subTest(items=items), patch('core_tools.term_filter.requests.post', return_value=self.response(items)):
                with self.assertRaisesRegex(ValueError, '响应不完整'):
                    self.tool.filter_terms_with_llm(record)

    def test_http_and_timeout_errors_are_safe(self):
        record = dict(english='Hi', context='', category='', key='', speaker='', emotion='', candidates=['Chad (Chad)'])
        with patch('core_tools.term_filter.requests.post', side_effect=requests.Timeout('TEST_SECRET')) as post, \
                patch('time.sleep'):
            with self.assertRaisesRegex(ValueError, '超时') as caught:
                self.tool.filter_terms_with_llm(record)
            self.assertNotIn('TEST_SECRET', str(caught.exception))
            self.assertEqual(post.call_count, 3)
        with patch('core_tools.term_filter.requests.post', return_value=Mock(ok=False, status_code=401)) as post:
            with self.assertRaisesRegex(ValueError, 'HTTP 401'):
                self.tool.filter_terms_with_llm(record)
            self.assertEqual(post.call_count, 1)

    def test_transient_502_retries_row_188_instead_of_aborting_workbook(self):
        path = self.source([['forestconcert_road_desc2', 'Collect Maracas to earn extra event points!', '',
                             '活动 UI 与商店文案', 'Woodland Singalong', 'event (Event)']])
        book = load_workbook(path)
        book.active.move_range('A2:F2', rows=186)
        book.save(path)
        book.close()
        response = self.response([{'id': 0, 'decision': 'keep', 'reason': '本句指游戏活动'}])
        with patch('core_tools.term_filter.requests.post', side_effect=[Mock(ok=False, status_code=502), response]) as post, \
                patch('time.sleep'):
            task = self.run_filter(path)
        self.assertEqual(task['state'], 'completed', task)
        self.assertEqual(post.call_count, 2)
        book = load_workbook(task['result_file'])
        try:
            self.assertEqual(book.active['F188'].value, 'event (Event)')
            self.assertIsNone(book.active['G188'].value)
        finally:
            book.close()

    def test_persistent_502_stops_after_three_attempts_without_output(self):
        path = self.source([['event', 'Collect event points!', '', '活动 UI 与商店文案', '', 'event (Event)']])
        with patch('core_tools.term_filter.requests.post', return_value=Mock(ok=False, status_code=502)) as post, \
                patch('time.sleep') as sleep:
            task = self.run_filter(path)
        self.assertEqual(task['state'], 'error')
        self.assertEqual(post.call_count, 3)
        self.assertEqual([call.args[0] for call in sleep.call_args_list], [1, 2])
        self.assertIn('HTTP 502（已尝试 3 次）', task['status'])
        self.assertNotIn('result_file', task)
        self.assertFalse((self.root / 'result.xlsx').exists())

    def test_transient_transport_failures_and_gateway_errors_share_retry_limit(self):
        record = dict(english='Hi Chad', context='', category='', key='', speaker='', emotion='', candidates=['Chad (Chad)'])
        response = self.response([{'id': 0, 'decision': 'keep', 'reason': '本行指角色人名'}])
        for error in (requests.Timeout('TEST_SECRET'), requests.ConnectionError('TEST_SECRET')):
            with self.subTest(error=type(error).__name__), \
                    patch('core_tools.term_filter.requests.post', side_effect=[error, Mock(ok=False, status_code=502), response]) as post, \
                    patch('time.sleep') as sleep:
                result = self.tool.filter_terms_with_llm(record)
                self.assertEqual(result[0]['term'], 'Chad (Chad)')
                self.assertEqual(post.call_count, 3)
                self.assertEqual([call.args[0] for call in sleep.call_args_list], [1, 2])

    def test_splits_candidates_into_only_two_columns_and_preserves_other_cells(self):
        path = self.source()
        with patch.object(self.tool, 'filter_terms_with_llm', side_effect=self.classify) as model:
            task = self.run_filter(path)
        self.assertEqual(task['state'], 'completed', task)
        self.assertEqual(model.call_count, 2)
        self.assertEqual(task['counts'], {'已筛选': 1, '有排除': 1, '无候选': 1})
        before, after = load_workbook(path), load_workbook(task['result_file'])
        try:
            self.assertEqual(before.sheetnames, after.sheetnames)
            for sheet in before:
                for row in sheet:
                    for cell in row:
                        if sheet.title == '活动ui' and cell.column == 6:
                            continue
                        new = after[sheet.title].cell(cell.row, cell.column)
                        self.assertEqual((cell.value, cell.data_type, cell._style), (new.value, new.data_type, new._style))
                self.assertEqual(sheet.freeze_panes, after[sheet.title].freeze_panes)
                self.assertEqual(sheet.row_dimensions[2].height, after[sheet.title].row_dimensions[2].height)
            self.assertEqual([cell.value for cell in after.active[1]],
                             [cell.value for cell in before.active[1]][:5] + ['生效术语', '无效术语'])
            self.assertEqual(after.active['G2'].value, 'key (Schlüssel)')
            self.assertIsNone(after.active['F2'].value)
            self.assertEqual(after.active['F4'].value, 'key (Schlüssel)')
            self.assertIsNone(after.active['F3'].value)
            self.assertTrue(all(not after.active.column_dimensions[c].hidden for c in ('F', 'G')))
            for record in task['records']:
                active = record['active'].split(' | ') if record['active'] else []
                excluded = record['excluded'].split(' | ') if record['excluded'] else []
                self.assertCountEqual(active + excluded, record['candidates'])
                self.assertFalse(set(active) & set(excluded))
        finally:
            before.close()
            after.close()

    def test_model_failure_does_not_publish_an_incomplete_result(self):
        path = self.source()
        with patch.object(self.tool, 'filter_terms_with_llm', side_effect=ValueError('模型请求超时，请重试')):
            task = self.run_filter(path)
        self.assertEqual(task['state'], 'error')
        self.assertIn('超时', task['status'])
        self.assertNotIn('result_file', task)
        self.assertFalse((self.root / 'result.xlsx').exists())

    def test_abnormal_source_and_candidate_formula_do_not_call_model(self):
        for row in [['bad', '=1+1', '', '待确认', '', 'key (Schlüssel)'],
                    ['formula', 'Hi', '', '活动对白', '', '=1+1'],
                    ['broken', 'Hi', '', '活动对白', '', 'bad data']]:
            with self.subTest(row=row), patch.object(self.tool, 'filter_terms_with_llm') as model:
                task = self.run_filter(self.source([row]))
                self.assertEqual(task['state'], 'error', task)
                self.assertNotIn('result_file', task)
                model.assert_not_called()

    def test_columns_after_candidates_keep_their_data_and_width(self):
        path = self.source([['key (Schlüssel)', 'With a key inside!', 'KEEP_VALUE', '活动对白', 'KEEP_CONTEXT']],
                           ['关联术语', 'game_text_en', '备注', '分类', 'Context'])
        book = load_workbook(path)
        book.active.column_dimensions['C'].width = 48
        book.save(path)
        book.close()
        with patch.object(self.tool, 'filter_terms_with_llm', side_effect=self.classify):
            task = self.run_filter(path)
        self.assertEqual(task['state'], 'completed', task)
        book = load_workbook(task['result_file'])
        try:
            self.assertEqual([c.value for c in book.active[1]],
                             ['game_text_en', '备注', '分类', 'Context', '生效术语', '无效术语'])
            self.assertEqual(book.active['B2'].value, 'KEEP_VALUE')
            self.assertEqual(book.active.column_dimensions['B'].width, 48)
            self.assertEqual(book.active['E2'].value, 'key (Schlüssel)')
        finally:
            book.close()

    def test_missing_model_configuration_fails_instead_of_passing(self):
        self.tool.api_key = ''
        with patch.object(self.tool, 'filter_terms_with_llm') as model:
            task = self.run_filter(self.source())
        self.assertEqual(task['state'], 'error')
        model.assert_not_called()

    def test_wrong_stage_existing_output_bad_file_and_done_are_rejected(self):
        for headers in [['game_text_en'], ['game_text_en', '分类', 'Context', '关联术语', '生效术语'],
                        ['game_text_en', '分类', 'Context', '关联术语', 'TM参考'],
                        ['game_text_en', '分类', 'Context', '关联术语', '关联术语']]:
            with self.subTest(headers=headers), patch.object(self.tool, 'filter_terms_with_llm') as model:
                self.assertEqual(self.run_filter(self.source([['Hi']], headers))['state'], 'error')
                model.assert_not_called()
        self.assertEqual(self.run_filter(self.source(), 'file_done.xlsx')['state'], 'error')
        bad = self.root / 'bad.xlsx'
        bad.write_bytes(b'bad')
        self.assertEqual(self.run_filter(bad)['state'], 'error')



if __name__ == '__main__':
    unittest.main()
