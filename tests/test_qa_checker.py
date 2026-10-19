import io
import time
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from core_tools.qa_checker import QA_POLICY, QAChecker, display_length, numbers, parse_terms

HEADERS = ['limit', 'key', 'game_text_en', 'game_text_de', '分类', 'Context', '关联术语']


class QACheckerTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def save(self, rows, headers=None, name='translated.xlsx', sheet_title='活动ui', extra_sheet=True):
        book = Workbook()
        sheet = book.active
        sheet.title = sheet_title
        sheet.append(headers or HEADERS)
        for row in rows:
            sheet.append(row)
        if extra_sheet:
            book.create_sheet('说明')['A1'] = '本表没有源文列，应被跳过'
        path = self.root / name
        book.save(path)
        book.close()
        return path

    def run_check(self, path, original=None, ratio=1.5, **kwargs):
        task = {'state': 'processing', 'input_name': path.name,
                'original_name': original.name if original else ''}
        checker = QAChecker('game_text_en', 'game_text_de', length_ratio=ratio, **kwargs)
        checker.process(path, self.root / 'report.xlsx', {'t': task}, 't', original)
        return task

    def issues(self, task, check=None):
        return [(i['level'], i['check'], i['message']) for r in task['records'] for i in r['issues']
                if check is None or i['check'] == check]

    def test_reports_missing_translation_and_untranslated_rows(self):
        task = self.run_check(self.save([
            [30, 'a', 'Collect {0} stars!', 'Sammle {0} Sterne!', '', '', ''],
            [30, 'b', 'Keep playing', '', '', '', ''],
            [30, 'c', 'Forest Finale', 'Forest Finale', '', '', ''],
            [30, 'd', '', 'Nur Zieltext', '', '', ''],
        ]))
        self.assertEqual(task['state'], 'completed', task['status'])
        self.assertEqual(task['total'], 4)
        self.assertEqual(task['translated'], 3)
        self.assertEqual(task['counts'], {'错误': 1, '提示': 2, '通过': 1})
        self.assertIn(('错误', '漏译', '有源文但译文为空'), self.issues(task))
        self.assertEqual(self.issues(task, '未翻译'), [('提示', '未翻译', '译文与源文完全相同，请确认是否需要翻译')])
        self.assertEqual(self.issues(task, '源文异常'), [('提示', '源文异常', '源文为空但有译文，请确认是否错行')])

    def test_checks_parameters_tags_newlines_and_numbers(self):
        task = self.run_check(self.save([
            [None, 'p', 'Collect {0} of {1}', 'Sammle {0}', '', '', ''],
            [None, 't', 'Get <b>50</b> coins', 'Hol dir 50 Münzen', '', '', ''],
            [None, 'n', 'Line one\\nLine two', 'Zeile eins Zeile zwei', '', '', ''],
            [None, 'd', 'Win 3 rounds', 'Gewinne 4 Runden', '', '', ''],
            [None, 'k', 'Collect 1,000 coins', 'Sammle 1.000 Münzen', '', '', ''],
        ]))
        self.assertEqual(self.issues(task, '参数'), [('错误', '参数', '缺少 {1}')])
        self.assertEqual(self.issues(task, '标签'), [('错误', '标签', '缺少 </b>、<b>')])
        self.assertEqual(self.issues(task, '换行'),
                         [('错误', '换行', '源文字面 \\n 1 个、实际换行 0 个；译文分别为 0、0 个')])
        self.assertEqual(self.issues(task, '数字'), [('错误', '数字', '缺少 3；多出 4')])
        # 1,000 与 1.000 是同一个数，不应报错
        self.assertEqual([r['key'] for r in task['records']], ['p', 't', 'n', 'd'])

    def test_sprite_tags_may_move_but_other_tags_keep_order(self):
        task = self.run_check(self.save([
            [None, 'a', 'I ate 3 <sprite name=x> today', 'Ich habe heute <sprite name=x> 3 gegessen', '', '', ''],
            [None, 'b', '<b>Now</b> <i>go</i>', '<i>go</i> <b>Now</b>', '', '', ''],
        ]))
        self.assertEqual(self.issues(task, '标签'), [])
        self.assertEqual(self.issues(task, '标签顺序'),
                         [('提示', '标签顺序', '标签数量一致但先后顺序不同（图片标签除外）')])

    def test_limit_and_length_ratio(self):
        rows = [[10, 'a', 'Feline Frenzy', 'Katzenwahnsinn Extra', '', '', ''],
                ['abc', 'b', 'Short', 'Kurz', '', '', ''],
                [None, 'c', 'Collect', 'Sammle alle Sterne im Wald', '', '', ''],
                [None, 'd', 'Hi <b>there</b>', 'Hallo <b>dort</b>', '', '', '']]
        task = self.run_check(self.save(rows))
        self.assertEqual(self.issues(task, '字数限制'),
                         [('错误', '字数限制', '译文 20 字符，超过限制 10'),
                          ('提示', '字数限制', 'limit 值“abc”无法识别，未检查字限')])
        self.assertEqual(self.issues(task, '长度比例'),
                         [('提示', '长度比例', '无字限；译文 26 字符，为源文 7 字符的 371%')])
        self.assertEqual(self.issues(self.run_check(self.save(rows), ratio=0), '长度比例'), [])

    def test_terminology_and_translation_consistency(self):
        task = self.run_check(self.save([
            [None, 'a', 'Open the Safe', 'Öffne den Tresor', '', '', 'Safe (Tresor)'],
            [None, 'b', 'Open the Safe', 'Öffne den Geldschrank', '', '', 'Safe (Tresor)'],
            [None, 'c', 'Collect rewards', 'Belohnungen sammeln', '', '', 'rewards (Belohnung / Preis)'],
            [None, 'd', 'Collect coins', 'Sammle Taler', '', '', 'coins (Münzen)'],
        ]))
        self.assertEqual(self.issues(task, '术语'),
                         [('提示', '术语', '未在译文中找到“Safe”的术语译法：Tresor（词形变化时可忽略）'),
                          ('提示', '术语', '未在译文中找到“coins”的术语译法：Münzen（词形变化时可忽略）')])
        # “Belohnungen sammeln”含词条译法 Belohnung，按字面包含即通过
        inconsistent = self.issues(task, '译法不一致')
        self.assertEqual(len(inconsistent), 2)
        self.assertIn('活动ui 第2行、活动ui 第3行', inconsistent[0][2])

    def test_prefers_filtered_term_column(self):
        headers = ['key', 'game_text_en', 'game_text_de', '关联术语', '生效术语']
        task = self.run_check(self.save([['a', 'Open the Safe', 'Öffne den Tresor', 'Open (Offen)', 'Safe (Tresor)']],
                                        headers=headers))
        self.assertEqual(self.issues(task, '术语'), [])

    def test_compares_with_original_file(self):
        original = self.save([[30, 'a', 'First line', '', '', '', ''],
                              [30, 'b', 'Second line', 'Vorhandene Übersetzung', '', '', ''],
                              [30, 'c', 'Third line', '', '', '', '']], name='original.xlsx')
        translated = self.save([[30, 'a', 'First line', 'Erste Zeile', '', '', ''],
                                [30, 'b', 'Second line', 'Neue Übersetzung', '', '', ''],
                                [30, 'x', 'Third line', 'Dritte Zeile', '', '', ''],
                                [30, 'd', 'Extra line', 'Zusätzlich', '', '', '']])
        task = self.run_check(translated, original)
        self.assertEqual(task['file_issues'],
                         [{'level': '错误', 'check': '行对应',
                           'message': '记录数不一致：原始文件 3 条，译后文件 4 条'}])
        self.assertEqual(self.issues(task, '行对应'),
                         [('提示', '行对应', '原始文件已有译文，本次被修改'),
                          ('错误', '行对应', '与原始文件第 4 行的 Key 或源文不一致，可能错行或源文被修改')])

    def test_accepts_csv_and_reports_formula_targets(self):
        path = self.root / 'translated.csv'
        path.write_text('key,game_text_en,game_text_de\na,Collect {0},Sammle {0}\n', encoding='utf-8-sig')
        self.assertEqual(self.run_check(path)['counts'], {'错误': 0, '提示': 0, '通过': 1})

        book = Workbook()
        sheet = book.active
        sheet.append(['key', 'game_text_en', 'game_text_de'])
        sheet.append(['a', '=CONCAT("A")', '=CONCAT("B")'])
        formulas = self.root / 'formula.xlsx'
        book.save(formulas)
        book.close()
        self.assertEqual(self.issues(self.run_check(formulas)),
                         [('提示', '源文异常', '源文单元格为公式，请核对实际源文'),
                          ('错误', '漏译', '译文单元格为公式，不是文本')])

    def test_missing_columns_and_empty_file_report_errors(self):
        self.assertEqual(self.run_check(self.save([['x']], headers=['other']))['state'], 'error')
        task = self.run_check(self.save([[30, 'a', 'Text', 'Übersetzung', '', '', '']], headers=HEADERS[:3] + ['other'] + HEADERS[4:]))
        self.assertEqual(task['state'], 'error')
        self.assertIn('缺少“game_text_de”列', task['status'])
        duplicated = self.save([[30, 'a', 'Text', 'Ü', '', '', '']], headers=HEADERS + ['game_text_de'])
        self.assertIn('多个', self.run_check(duplicated)['status'])
        broken = self.root / 'broken.xlsx'
        broken.write_bytes(b'not a workbook')
        self.assertEqual(self.run_check(broken)['state'], 'error')

    def test_report_keeps_text_and_lists_every_check(self):
        task = self.run_check(self.save([[5, 'a', '=1+1 source', 'Translated text', '', '', '']]))
        book = load_workbook(self.root / 'report.xlsx')
        try:
            sheet = book['问题清单']
            self.assertEqual([c.value for c in sheet[1]][:3], ['工作表', '原行号', 'Key'])
            self.assertEqual(sheet['A2'].value, '活动ui')
            self.assertEqual(sheet['B2'].value, 2)
            self.assertEqual(sheet['D2'].value, '=1+1 source')
            self.assertEqual(sheet['D2'].data_type, 's')
            overview = dict((row[0].value, row[1].value) for row in book['汇总'].iter_rows(min_row=2))
            self.assertEqual(overview['检查规则'], QA_POLICY)
            self.assertIn('未上传', overview['原始文件'])
            self.assertIn('源文 game_text_en，译文 game_text_de', overview['列设置'])
            self.assertIn('检查项：术语', overview)
            self.assertIn('不评价译文质量', overview['说明'])
        finally:
            book.close()
        # 公式源文同时触发源文异常，且源文中的 1+1 与译文数字不一致
        self.assertEqual(task['check_counts'], {'源文异常': 1, '数字': 1, '字数限制': 1})

    def test_percent_followed_by_letter_is_not_a_parameter(self):
        task = self.run_check(self.save([
            [None, 'a', '100% increase', '100%ige Steigerung', '', '', ''],
            [None, 'b', 'Deals 200% damage', 'Verursacht 200% Schaden', '', '', ''],
            [None, 'c', 'Win %d rounds', 'Gewinne Runden', '', '', ''],
        ]), ratio=0)
        self.assertEqual(self.issues(task, '参数'), [('错误', '参数', '缺少 %d')])

    def test_helpers(self):
        self.assertEqual(display_length('Hi <b>there</b>\\n!'), 9)
        self.assertEqual(numbers('1.000 and 1,000'), {'1000': 2})
        self.assertEqual(parse_terms('Safe (Tresor) | rewards (Belohnung / Preis) | broken'),
                         [('Safe', ['Tresor']), ('rewards', ['Belohnung', 'Preis'])])

    def test_http_upload_preview_and_download(self):
        import app as server
        path = self.save([[5, 'a', 'Collect {0} stars', 'Sammle {0} Sterne und mehr', '', '', '']])
        task_id = None
        with patch.dict(server.app.config, {'UPLOAD_FOLDER': str(self.root)}):
            client = server.app.test_client()
            self.assertEqual(client.get('/qa').status_code, 200)
            self.assertEqual(client.post('/upload/qa', data={}).status_code, 400)
            for data in ({'translated_file': (io.BytesIO(b'x'), 'a.txt')},
                         {'translated_file': (io.BytesIO(b'x'), 'a.xlsx'), 'source_column': ''},
                         {'translated_file': (io.BytesIO(b'x'), 'a.xlsx'), 'source_column': 'en', 'target_column': 'EN'},
                         {'translated_file': (io.BytesIO(b'x'), 'a.xlsx'), 'source_column': 'en', 'target_column': 'de', 'length_ratio': 'abc'},
                         {'translated_file': (io.BytesIO(b'x'), 'a.xlsx'), 'source_column': 'en', 'target_column': 'de', 'length_ratio': '9'}):
                self.assertEqual(client.post('/upload/qa', data=data).status_code, 400, data)
            try:
                response = client.post('/upload/qa', data={
                    'translated_file': (io.BytesIO(path.read_bytes()), 'WC_de_draft.xlsx'),
                    'source_column': 'game_text_en', 'target_column': 'game_text_de',
                    'key_column': 'key', 'limit_column': 'limit', 'length_ratio': '1.5'})
                task_id = response.json['task_id']
                for _ in range(100):
                    progress = client.get('/progress/' + task_id).json
                    if progress['state'] != 'processing':
                        break
                    time.sleep(.02)
                self.assertEqual(progress['state'], 'completed', progress)
                self.assertNotIn('records', progress)
                preview = client.get('/api/qa/' + task_id).json
                self.assertEqual(preview['total'], 1)
                self.assertEqual(preview['counts']['错误'], 1)
                self.assertEqual(preview['records'][0]['issues'][0]['check'], '字数限制')
                report = client.get('/download/' + task_id)
                self.assertIn('attachment', report.headers['Content-Disposition'])
                book = load_workbook(io.BytesIO(report.data))
                self.assertEqual(book['问题清单']['I2'].value, '字数限制')
                book.close()
                report.close()
                self.assertEqual(client.get('/api/qa/missing').status_code, 404)
            finally:
                if task_id:
                    server.tasks.pop(task_id, None)


if __name__ == '__main__':
    unittest.main()
