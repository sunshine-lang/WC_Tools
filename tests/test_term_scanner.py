import io
import time
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core_tools.term_scanner import TermScanner


class TermScannerTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.tb = self.root / 'WC术语库.xlsx'
        self.save_terms([['Max', 'Max'], ['Chad', 'Chad'], ['key', 'Schlüssel'],
                         ['Key Chains', 'Schlüsselanhänger'], ['Coin', 'Münze'],
                         ['coins', 'Münzen'], ['company', 'Firma'],
                         ["Chad's Vacation Wipeout", 'Chads Urlaubsrennen']])

    def tearDown(self):
        self.temp.cleanup()

    def save_terms(self, rows, headers=None):
        book = Workbook()
        book.active.title = '术语表'
        book.active.append(headers or ['EN', 'DE'])
        for row in rows:
            book.active.append(row)
        book.save(self.tb)
        book.close()

    def source(self, rows=None, headers=None):
        book = Workbook()
        sheet = book.active
        sheet.title = '活动ui'
        sheet.append(headers or ['limit', 'key', 'game_text_cn', 'game_text_en', 'game_text_de', '分类', 'Context'])
        for row in rows or [[30, 'test', '参考', "  <i>Chad</i> earns keys! {0}  ", 'OLD_TARGET', '活动 UI 与商店文案', 'KEEP_CONTEXT']]:
            sheet.append(row)
        sheet['D2'].font = Font(bold=True, color='123456')
        sheet['A2'].number_format = '0'
        sheet.row_dimensions[2].height = 40
        sheet.freeze_panes = 'B2'
        book.create_sheet('说明')['A1'] = 'KEEP_OTHER_SHEET'
        path = self.root / 'step2.xlsx'
        book.save(path)
        book.close()
        return path

    def run_scan(self, path, name='step2.xlsx'):
        task = {'state': 'processing', 'input_name': name}
        TermScanner(self.tb).process(path, self.root / 'result.xlsx', {'t': task}, 't')
        return task

    def test_deduplicates_identical_pairs_but_keeps_conflicts_and_sources(self):
        self.save_terms([['safe', 'Tresor'], ['safe', 'Tresor'], ['SAFE', 'sicher']])
        scanner = TermScanner(self.tb)
        self.assertEqual(scanner.status()['row_count'], 3)
        self.assertEqual(scanner.status()['term_count'], 1)
        self.assertEqual(scanner.conflict_count, 1)
        match = scanner.find_matches('A safe.')[0]
        self.assertEqual(match['target_text'], ['Tresor', 'sicher'])
        self.assertEqual(len(match['sources']), 3)
        task = self.run_scan(self.source([[None, 'safe', '', 'A safe.', None, '活动 UI 与商店文案', None]]))
        self.assertEqual(task['counts']['待确认'], 1)
        self.assertIn('Tresor / sicher', task['records'][0]['terms'])

    def test_requires_english_german_text_and_skips_deleted_terms(self):
        self.save_terms([['中文', 'English']], ['cn', 'en'])
        with self.assertRaisesRegex(ValueError, 'EN／DE'):
            TermScanner(self.tb)
        for row in [['key', None], ['key', '=1+1'], ['key', 42]]:
            self.save_terms([row])
            with self.assertRaisesRegex(ValueError, '第2行'):
                TermScanner(self.tb)
        self.save_terms([['Max', 'Max'], ['key', 'Schlüssel']])
        book = load_workbook(self.tb)
        book.active['A3'].font = Font(strike=True)
        book.save(self.tb)
        book.close()
        self.assertEqual(TermScanner(self.tb).scan_text('Max key'), 'Max (Max)')

    def test_boundaries_inline_tags_attributes_and_sprite_objects(self):
        scanner = TermScanner(self.tb)
        text = '<color name="Max">Ma</color>x needs key chains and a key. <sprite name="Chad">'
        self.assertEqual(scanner.scan_text(text), 'Max (Max) | Key Chains (Schlüsselanhänger) | key (Schlüssel)')
        self.assertEqual(scanner.scan_text('maximum monkey keyboard Chadwick'), '')
        self.assertEqual(scanner.scan_text('key<sprite name="Chains">chains'), 'key (Schlüssel)')
        self.assertEqual(scanner.scan_text('<!-- Max -->'), '')
        self.assertEqual(scanner.scan_text('&lt;sprite name="Chad"&gt;'), '')
        self.assertEqual(scanner.scan_text(None), '')

    def test_case_possessives_common_inflections_and_exact_form_priority(self):
        scanner = TermScanner(self.tb)
        self.assertEqual(scanner.scan_text("CHAD’s keys and companies"), 'Chad (Chad) | key (Schlüssel) | company (Firma)')
        self.assertEqual(scanner.scan_text('coins'), 'coins (Münzen)')
        self.assertEqual(scanner.find_matches('keys')[0]['match_type'], '词形')
        self.assertEqual(scanner.find_matches('coins')[0]['match_type'], '原词')
        self.assertEqual(scanner.scan_text("Chad’s Vacation Wipeout"), "Chad's Vacation Wipeout (Chads Urlaubsrennen)")

    def test_preserves_workbook_and_only_appends_one_column(self):
        source = self.source()
        original = source.read_bytes()
        task = self.run_scan(source)
        self.assertEqual(task['state'], 'completed', task)
        self.assertEqual(task['counts'], {'命中术语': 1, '未命中': 0, '待确认': 0, '未使用': 0})
        self.assertEqual(source.read_bytes(), original)
        before, after = load_workbook(source), load_workbook(task['result_file'])
        try:
            for old_sheet in before:
                new_sheet = after[old_sheet.title]
                for row in old_sheet:
                    for old in row:
                        new = new_sheet.cell(old.row, old.column)
                        self.assertEqual((old.value, old.data_type, old._style), (new.value, new.data_type, new._style))
                self.assertEqual(old_sheet.freeze_panes, new_sheet.freeze_panes)
                self.assertEqual(old_sheet.row_dimensions[2].height, new_sheet.row_dimensions[2].height)
            self.assertEqual([c.value for c in after.active[1]], ['limit', 'key', 'game_text_cn', 'game_text_en', 'game_text_de', '分类', 'Context', '关联术语'])
            self.assertEqual(after.active['H2'].value, 'Chad (Chad) | key (Schlüssel)')
        finally:
            before.close()
            after.close()

    def test_formula_blank_and_unmatched_sources_are_distinct_and_preserved(self):
        source = self.source([[None, 'formula', '', '=1+1', None, '待确认', None],
                              [None, 'empty', '', None, None, '待确认', None],
                              [None, 'Max', 'happy', 'Nothing relevant', None, '活动对白', None]])
        task = self.run_scan(source)
        self.assertEqual(task['state'], 'completed', task)
        self.assertEqual(task['counts'], {'命中术语': 0, '未命中': 1, '待确认': 2, '未使用': 0})
        self.assertEqual(task['records'][2]['speaker'], 'Max')
        self.assertEqual(task['records'][2]['terms'], '')
        book = load_workbook(task['result_file'])
        self.assertEqual((book.active['D2'].value, book.active['D2'].data_type), ('=1+1', 'f'))
        book.close()

    def test_rejects_wrong_stage_repeated_results_and_done_materials(self):
        for headers, expected in [(['game_text_en'], '第二步'), (['game_text_en', '分类', 'Context', '关联术语'], '已有')]:
            task = self.run_scan(self.source([['Max']], headers))
            self.assertEqual(task['state'], 'error')
            self.assertIn(expected, task['status'])
        task = self.run_scan(self.source(), 'source_done.xlsx')
        self.assertEqual(task['state'], 'error')
        self.assertIn('done', task['status'])
        with self.assertRaisesRegex(ValueError, 'done'):
            TermScanner(self.root / 'tb_done.xlsx')
        with self.assertRaisesRegex(ValueError, 'XLSX'):
            TermScanner(self.root / 'tb.txt')

    def test_reads_customer_csv_and_keeps_original_upload_name(self):
        path = self.root / 'stored.csv'
        path.write_text('\ufeffEN,DE\nMax,客户译名\nkey,"Schlüssel, Variante"\n', encoding='utf-8')
        scanner = TermScanner(path, source_name='客户词库.csv')
        self.assertEqual(scanner.scan_text('Max keys'), 'Max (客户译名) | key (Schlüssel, Variante)')
        self.assertEqual(scanner.filename, '客户词库.csv')
        self.assertIn('客户词库.csv', scanner.find_matches('Max')[0]['sources'][0])
        path.write_bytes(b'EN,DE\nkey,\xff\n')
        with self.assertRaisesRegex(ValueError, 'UTF-8'):
            TermScanner(path)
        path.write_text('EN,DE,EN\na,b,c\n', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '重复'):
            TermScanner(path)





if __name__ == '__main__':
    unittest.main()
