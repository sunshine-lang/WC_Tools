import base64
import hashlib
import io
import json
import shutil
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Border, Font, PatternFill, Side

from core_tools.context_processor import ContextProcessor
from core_tools.key_classifier import CATEGORIES, PENDING, KeyClassifier
from core_tools.layout_segmenter import segment_sheet
from pipeline import run_classify, run_normalize, run_context
from pipeline_service import PipelineError, PipelineService

ROOT = Path(__file__).resolve().parents[1]
PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')


class EvidenceWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.ctx = {'game': json.loads((ROOT / 'profiles/game/wc.json').read_text()),
                    'language': {'source_lang': 'en', 'target_lang': 'de'}, 'batch_dir': str(self.root)}

    def tearDown(self):
        self.tmp.cleanup()

    def source(self, rows, formatted=True, images=()):
        book = Workbook()
        sheet = book.active
        sheet.append(['key', 'game_text_en', 'game_text_de', 'note'])
        for row in rows:
            sheet.append([*row[:2], 'KEEP', row[2] if len(row) > 2 else None])
        if formatted:
            for row in sheet.iter_rows(min_row=2, max_col=3):
                for cell in row:
                    cell.fill = PatternFill('solid', fgColor='DDEEDD')
        for start, end in images:
            image = Image(io.BytesIO(PNG))
            image.anchor = TwoCellAnchor(_from=AnchorMarker(col=0, row=start-1), to=AnchorMarker(col=1, row=end))
            sheet.add_image(image)
        path = self.root / 'source.xlsx'
        book.save(path)
        book.close()
        return path

    def classify(self, source, overrides=None):
        if overrides is not None:
            (self.root / 'classification_overrides.json').write_text(json.dumps(overrides))
        task = run_classify(self.ctx, source, self.root / '01-分类.xlsx')
        self.assertEqual(task['state'], 'completed', task.get('status'))
        return task

    def context(self):
        task = {}
        ContextProcessor().process(str(self.root / '01-分类.xlsx'), str(self.root / '02-上下文.xlsx'),
            {'t': task}, 't', classification_path=self.root / '01-分类.classification.json',
            rules_path=ROOT / 'profiles/game/wc_key_rules.json')
        self.assertEqual(task['state'], 'completed', task.get('status'))
        return task

    def test_no_key_requires_layout_and_content_then_explicit_confirmation(self):
        path = self.source([['cat_multiboards_title', 'Cat event'], ['', 'Collect rewards!'],
                            ['cat_multiboards_pre_desc', 'Help the cats']])
        task = self.classify(path)
        row = task['records'][1]
        self.assertEqual((row['category'], row['classification_status']), (CATEGORIES[0], 'inferred'))
        self.assertEqual(self.context()['records'][1]['context'], '')
        self.classify(path, {row['id']: {'category': CATEGORIES[0], 'record_hash': row['record_hash']}})
        record = self.context()['records'][1]
        self.assertIn('Cat event', record['context'])
        self.assertIn('confirmed_layout_region', record['context_record']['context_evidence'])
        self.assertEqual(record['context_record']['activity_id'], 'cat')

    def test_colour_or_adjacency_alone_never_classifies_keyless_text(self):
        for formatted, text in [(False, 'Collect rewards!'), (True, 'This is an ambiguous sentence.')]:
            task = self.classify(self.source([['cat_multiboards_title', 'Cats'], ['', text],
                                              ['cat_multiboards_pre_desc', 'Play with cats']], formatted))
            self.assertEqual(task['records'][1]['category'], PENDING)

    def test_layout_key_conflict_and_manual_priority(self):
        path = self.source([['Ahalloween_set1_name', 'Autumn'], ['Ahalloween_set1_1_name', 'Pumpkin'],
                            ['mdnew_halloween_mdnew_decoration_name_1', 'Lantern']])
        task = self.classify(path)
        row = task['records'][2]
        self.assertEqual(row['category'], PENDING)
        self.assertIn('layout_conflict', [e['source'] for e in row['evidence']])
        task = self.classify(path, {row['id']: {'category': CATEGORIES[2], 'record_hash': row['record_hash']}})
        self.assertEqual(task['records'][2]['category'], CATEGORIES[2])
        self.assertEqual(task['records'][2]['classification_status'], 'confirmed')

    def test_activity_fences_and_reused_colours_are_not_one_region(self):
        task = self.classify(self.source([['cat_multiboards_title', 'Cats'],
            ['midautumn_multiboards_title', 'Moon'], ['cat_multiboards_pre_desc', 'Cats again']]))
        self.assertEqual(len({r['region_id'] for r in task['records']}), 3)

    def test_unrelated_cell_style_does_not_split_region_but_border_does(self):
        book = Workbook(); sheet = book.active
        for row in [['key','en','de','note'],['a','A','AA',''],['b','B','BB',''],['c','C','CC','']]: sheet.append(row)
        sheet['D3'].fill = PatternFill('solid', fgColor='FF0000')
        before = segment_sheet(sheet, [1,2,3,4])
        self.assertEqual(before['rows'][2]['region_id'], before['rows'][3]['region_id'])
        for cell in (sheet['A3'], sheet['B3']): cell.border = Border(bottom=Side(style='thin'))
        after = segment_sheet(sheet, [1,2,3,4])
        self.assertNotEqual(after['rows'][3]['region_id'], after['rows'][4]['region_id'])
        sheet.merge_cells('A3:B3')
        merged = segment_sheet(sheet, [1,2,3,4])
        self.assertIn('A3:B3', merged['rows'][3]['merged_ranges'])

    def test_csv_roundtrip_keeps_no_format_fallback(self):
        source = self.root / 'input.csv'
        source.write_text('key,en\ncat_multiboards_title,Cats\n,Collect rewards!\ncat_multiboards_pre_desc,Help cats\n')
        self.ctx['source_format'] = '.csv'
        normalized = self.root / '00.xlsx'
        self.assertEqual(run_normalize(self.ctx, source, normalized)['state'], 'completed')
        task = self.classify(normalized)
        self.assertFalse(next(iter(task['layouts'].values()))['available'])
        self.assertEqual(task['records'][1]['category'], PENDING)
        self.assertEqual(task['records'][0]['category'], CATEGORIES[0])

    def test_sets_do_not_cross_and_activity_theme_is_explicit(self):
        path = self.source([['album_halloween_title','Halloween'], ['album_halloween_pre_desc','Collect spooky cards'],
            ['Ahalloween_set1_name','Autumn'], ['Ahalloween_set1_1_name','Pumpkin'],
            ['Ahalloween_set2_name','Ocean'], ['Ahalloween_set2_1_name','Shark'],
            ['cat_multiboards_title','CAT_SENTINEL']], formatted=False)
        self.classify(path)
        row = self.context()['records'][3]
        self.assertIn('Autumn', row['context']); self.assertIn('Collect spooky cards', row['context'])
        self.assertNotIn('Shark', row['context']); self.assertNotIn('CAT_SENTINEL', row['context'])
        self.assertEqual(row['context_record']['group_id'], 'halloween/set/1')
        self.assertNotIn('Sheet:7', row['context_record']['related_record_ids'])

    def test_unique_decoration_image_and_dialogue_with_ambiguous_image_excluded(self):
        path = self.source([['mdnew_forestconcert_mdnew_decoration_name_1','Owl'],
            ['mdnew_forestconcert_mdnew_decoration_name_2','Fox'],
            ['mdnew_forestconcert_mdnew_decoration_1_1','The owl sings.'],
            ['mdnew_forestconcert_mdnew_decoration_2_1','FOX_SENTINEL']], formatted=False,
            images=[(2,2),(2,3)])
        self.classify(path)
        task = self.context(); row = task['records'][0]
        self.assertIn('The owl sings.', row['context']); self.assertNotIn('FOX_SENTINEL',row['context'])
        self.assertEqual([im['image_id'] for im in row['context_record']['image_evidence']], ['Sheet:image:1'])
        self.assertIn('归属不唯一', row['context'])
        with ZipFile(path) as a, ZipFile(task['result_file']) as b:
            media = lambda z: {n:z.read(n) for n in z.namelist() if n.startswith('xl/media/')}
            self.assertEqual(media(a), media(b))

    def test_same_keys_in_other_sheet_do_not_leak_context(self):
        path = self.source([['cat_multiboards_title','Cats'], ['cat_multiboards_pre_desc','Play here']])
        book = load_workbook(path); sheet = book.copy_worksheet(book.active); sheet['B3'] = 'OTHER_SHEET_SENTINEL'; book.save(path); book.close()
        self.classify(path)
        self.assertNotIn('OTHER_SHEET_SENTINEL', self.context()['records'][0]['context'])

    def test_classification_manifest_detects_workbook_or_rule_changes(self):
        path = self.source([['cat_multiboards_title','Cats']]); self.classify(path)
        output = self.root / '01-分类.xlsx'
        book = load_workbook(output); book.active['B2'] = 'Changed'; book.save(output); book.close()
        task = run_context(self.ctx, output, self.root / 'out.xlsx')
        self.assertEqual(task['state'], 'error'); self.assertFalse((self.root/'out.xlsx').exists())

    def test_manual_override_cannot_survive_changed_source(self):
        path = self.source([['', 'Collect rewards!']]); row = self.classify(path)['records'][0]
        (self.root/'classification_overrides.json').write_text(json.dumps({row['id']:{'category':CATEGORIES[0],'record_hash':row['record_hash']}}))
        self.source([['', 'Different source']])
        task = run_classify(self.ctx,path,self.root/'changed.xlsx')
        self.assertEqual(task['state'], 'error')

    def test_customer_fixture_and_non_target_values_styles_images_preserved(self):
        source = ROOT.parent / 'WC_UI部分翻译需求.xlsx'
        if not source.exists(): self.skipTest('客户材料不随代码仓库分发')
        fixture = json.loads((ROOT/'tests/fixtures/wc_classification_expected.json').read_text())
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), fixture['source_sha256'])
        normalized = self.root/'00.xlsx'; run_normalize(self.ctx,source,normalized)
        task = self.classify(normalized); by_row = {r['row']:r for r in task['records']}
        for category, ranges in fixture['ranges'].items():
            for start,end in ranges:
                for row in range(start,end+1): self.assertEqual(by_row[row]['category'],category,(row,by_row[row]['basis']))
        self.assertEqual(len(task['records']),346)
        context = self.context()
        before, after = load_workbook(source), load_workbook(context['result_file'])
        try:
            self.assertEqual(before.sheetnames,after.sheetnames)
            for sheet in before:
                for row in sheet:
                    for cell in row:
                        actual=after[sheet.title].cell(cell.row,cell.column)
                        self.assertEqual(cell.value,actual.value,(cell.coordinate,'value'))
                        self.assertEqual(cell._style,actual._style,(cell.coordinate,'style'))
                self.assertEqual(len(sheet._images),len(after[sheet.title]._images))
                self.assertEqual([image.anchor for image in sheet._images],
                                 [image.anchor for image in after[sheet.title]._images])
        finally:
            before.close(); after.close()
        with ZipFile(source) as original, ZipFile(context['result_file']) as result:
            media = lambda z: sorted(hashlib.sha256(z.read(name)).hexdigest() for name in z.namelist()
                                     if name.startswith('xl/media/') and not name.endswith('/'))
            self.assertEqual(media(original), media(result))


class ClassificationReviewTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.root=Path(self.tmp.name)
        self.service=PipelineService(workspace=self.root/'workspace')
        self.project=self.service.create_project('Evidence test')['project']['id']
        book=Workbook(); book.active.append(['key','en','de']); book.active.append(['cat_multiboards_title','Cats','KEEP']); book.active.append(['','Collect rewards!','KEEP'])
        self.source=self.root/'input.xlsx'; book.save(self.source); book.close()
        class Upload:
            filename='input.xlsx'
            def save(inner,path): shutil.copyfile(self.source,path)
        self.batch=self.service.create_batch(self.project,Upload())['batch']['id']

    def tearDown(self): self.tmp.cleanup()

    def wait(self):
        for _ in range(200):
            view=self.service.batch_state(self.project,self.batch)
            if not view['running']: return view
            time.sleep(.02)
        self.fail('Pipeline timeout')

    def test_review_blocks_bypass_and_override_rebuilds_context_and_invalidates_downstream(self):
        self.service.start(self.project,self.batch); view=self.wait()
        self.assertEqual(view['stages'][1]['state'],'review')
        with self.assertRaises(PipelineError): self.service.confirm(self.project,self.batch)
        with self.assertRaises(PipelineError): self.service.start(self.project,self.batch,stage='context',mode='rerun')
        row=view['stages'][1]['preview']['records'][0]
        folder=self.service._batch_dir(self.project,self.batch)
        (folder/'05-TM.xlsx').write_bytes(b'old')
        self.service.save_classifications(self.project,self.batch,[{'id':row['id'],'record_hash':row['record_hash'],'category':CATEGORIES[0]}])
        view=self.wait()
        self.assertEqual(view['stages'][1]['state'],'done'); self.assertEqual(view['stages'][2]['state'],'review')
        self.assertFalse((folder/'05-TM.xlsx').exists())
        self.assertTrue((folder/'02-上下文.context.json').exists())
        state=PipelineService(workspace=self.root/'workspace').batch_state(self.project,self.batch)
        overridden=[r for r in state['stages'][1]['preview']['records'] if r.get('manual_category')]
        self.assertEqual(len(overridden),1)
        self.service.save_classifications(self.project,self.batch,[{'id':row['id'],'record_hash':row['record_hash'],'category':None}])
        self.assertEqual(self.wait()['stages'][1]['state'],'review')
        self.assertFalse((folder/'02-上下文.xlsx').exists())

    def test_changed_rules_invalidate_finished_results(self):
        self.service.start(self.project,self.batch); view=self.wait()
        rules=self.root/'rules.json'; rules.write_text((ROOT/'profiles/game/wc_key_rules.json').read_text())
        with patch.object(self.service,'_key_rules_path',return_value=rules):
            state=self.service._load_batch(self.project,self.batch)
            self.assertEqual(state['stages']['classify']['state'],'pending')
            self.assertFalse((self.service._batch_dir(self.project,self.batch)/'01-分类.xlsx').exists())

    def test_invalid_override_request_does_not_mutate_batch(self):
        self.service.start(self.project,self.batch); self.wait()
        folder=self.service._batch_dir(self.project,self.batch)
        before=(folder/'01-分类.xlsx').read_bytes()
        with self.assertRaises(PipelineError):
            self.service.save_classifications(self.project,self.batch,[{'id':'alien:2','category':CATEGORIES[0]}])
        self.assertFalse((folder/'classification_overrides.json').exists())
        self.assertEqual(before,(folder/'01-分类.xlsx').read_bytes())

    def test_bulk_confirmation_is_atomic_and_only_reruns_classification_once(self):
        import pipeline_service
        from unittest.mock import Mock
        book=load_workbook(self.source);book.active.append(['','Help the cats!','KEEP']);book.save(self.source);book.close()
        class Upload:
            filename='two-pending.xlsx'
            def save(inner,target):shutil.copyfile(self.source,target)
        self.batch=self.service.create_batch(self.project,Upload())['batch']['id']
        self.service.start(self.project,self.batch);view=self.wait()
        pending=[r for r in view['stages'][1]['preview']['records'] if r['classification_status']!='confirmed']
        self.assertEqual(len(pending),2)
        changes=[{'id':r['id'],'record_hash':r['record_hash'],'category':CATEGORIES[0]} for r in pending]
        folder=self.service._batch_dir(self.project,self.batch)
        before=(folder/'01-分类.xlsx').read_bytes()
        with self.assertRaises(PipelineError):
            self.service.save_classifications(self.project,self.batch,[changes[0],{**changes[1],'record_hash':'stale'}])
        self.assertFalse((folder/'classification_overrides.json').exists())
        self.assertEqual(before,(folder/'01-分类.xlsx').read_bytes())
        classify=Mock(wraps=pipeline_service.RUNNERS['classify'])
        context=Mock(wraps=pipeline_service.RUNNERS['context'])
        with patch.dict(pipeline_service.RUNNERS,{'classify':classify,'context':context}):
            self.service.save_classifications(self.project,self.batch,changes)
            view=self.wait()
        self.assertEqual(classify.call_count,1)
        self.assertEqual(context.call_count,1)
        self.assertEqual(view['stages'][1]['state'],'done')
        self.assertEqual(view['stages'][2]['state'],'review')
        self.assertEqual(len(json.loads((folder/'classification_overrides.json').read_text())),2)

    def test_non_dialogue_pipeline_and_preview_do_not_require_speaker(self):
        book = Workbook(); sheet = book.active
        sheet.append(['key', 'en', 'de'])
        for key, source in [('cat_multiboards_title', 'Cats'), ('cat_multiboards_pre_desc', 'Help cats'),
                            ('Ahalloween_set1_name', 'Autumn'), ('Ahalloween_set1_1_name', 'Pumpkin'),
                            ('mdnew_forestconcert_mdnew_decoration_name_1', 'Owl')]:
            sheet.append([key, source, 'KEEP'])
        path = self.root / 'non-dialogue.xlsx'; book.save(path); book.close()
        class Upload:
            filename = 'non-dialogue.xlsx'
            def save(inner, target): shutil.copyfile(path, target)
        self.batch = self.service.create_batch(self.project, Upload())['batch']['id']
        self.service.start(self.project, self.batch); view = self.wait()
        self.assertEqual(view['stages'][2]['state'], 'review')
        self.assertNotIn('speaker', view['stages'][0]['preview']['head'])
        self.service.confirm(self.project, self.batch); view = self.wait()
        self.assertEqual(view['done'], 6)
        output, _ = self.service.final_artifact(self.project, self.batch)
        result = load_workbook(output)
        try:
            self.assertNotIn('speaker', [cell.value for cell in result.active[1]])
            self.assertEqual(result.active['C2'].value, 'KEEP')
        finally:
            result.close()


if __name__ == '__main__': unittest.main()
