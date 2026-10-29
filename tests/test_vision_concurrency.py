import hashlib
import io
import json
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import Mock, patch
from zipfile import ZipFile

import requests
from PIL import Image
from openpyxl import load_workbook

from core_tools.model_concurrency import model_slot, worker_count
from core_tools.vision_context import VisionContextAnalyzer
from core_tools.tm_semantic import SemanticTMRetriever
from core_tools.term_filter import TermFilter
from core_tools.dialogue_context import DialogueContextAnalyzer
from pipeline import run_context
import test_evidence_workflow as evidence_fixtures


def visual_result():
    return {'status': 'analyzed', 'summary': '白底上有一个红色方形。',
            'objects': [{'description': '红色方形', 'location': '中央'}],
            'visible_text': [], 'uncertainties': []}


class ModelConcurrencyTests(unittest.TestCase):
    def test_same_model_shares_thirty_slots_while_other_model_can_run(self):
        release, full = threading.Event(), threading.Event()
        lock = threading.Lock()
        active = peak = 0
        def request():
            nonlocal active, peak
            with model_slot('https://test.invalid/chat', 'shared-model'):
                with lock:
                    active += 1; peak = max(active, peak)
                    if active == 30: full.set()
                release.wait(5)
                with lock: active -= 1
        with ThreadPoolExecutor(max_workers=42) as pool:
            futures = [pool.submit(request) for _ in range(40)]
            try:
                self.assertTrue(full.wait(3), '30 concurrent calls were not admitted')
                def other():
                    with model_slot('https://test.invalid/chat', 'independent-model'): return True
                self.assertTrue(pool.submit(other).result(timeout=2))
                with lock: self.assertEqual(active, 30)
            finally:
                release.set()
            for future in futures: future.result()
        self.assertEqual(peak,30)

    def test_slots_release_on_failure_and_worker_cap_is_thirty(self):
        with self.assertRaises(ValueError):
            with model_slot('https://test.invalid/chat', 'failure-model'): raise ValueError('failure')
        self.assertEqual(worker_count({}),30)
        self.assertEqual(worker_count({'max_workers':50}),30)
        self.assertEqual(worker_count({'max_workers':0}),1)


class VisionContextTests(unittest.TestCase):
    def setUp(self):
        self.fixture = evidence_fixtures.EvidenceWorkflowTests(); self.fixture.setUp()
        self.root = self.fixture.root
        self.config = self.root/'config.json'
        base = {'api_url':'https://example.invalid/chat','api_key':'TEST_SECRET','model_name':'vision-test','max_workers':30}
        self.config.write_text(json.dumps({'llm':base,'embedding':{**base,'model_name':'embed'},'rerank':{**base,'model_name':'rank'}}))
        self.fixture.ctx['config_path'] = self.config
        self.client = VisionContextAnalyzer(self.config,self.root/'cache')
        buffer=io.BytesIO(); Image.new('RGB',(16,16),'red').save(buffer,format='PNG');self.data=buffer.getvalue()
        self.guard=patch('requests.sessions.Session.request',side_effect=AssertionError('NETWORK_DISABLED'));self.guard.start()

    def tearDown(self):
        self.guard.stop();self.fixture.tearDown()

    def test_all_workers_accept_thirty(self):
        tm=SemanticTMRetriever(self.config)
        self.assertEqual((tm.embed_workers,tm.rank_workers),(30,30))
        self.assertEqual(TermFilter(self.config).max_workers,30)
        self.assertEqual(DialogueContextAnalyzer(self.config).max_workers,30)
        self.assertEqual(self.client.max_workers,30)

    def test_request_contains_real_pixels_and_no_paths_or_target_text(self):
        response=Mock(status_code=200)
        response.json.return_value={'choices':[{'message':{'content':json.dumps(visual_result())}}]}
        with patch('core_tools.vision_context.requests.post',return_value=response) as post:
            self.client.analyze(self.data,'set')
        body=post.call_args.kwargs['json'];content=body['messages'][1]['content']
        self.assertEqual(content[0]['type'],'text')
        self.assertTrue(content[1]['image_url']['url'].startswith('data:image/png;base64,'))
        self.assertNotIn('TEST_SECRET',json.dumps(body));self.assertNotIn(str(self.root),json.dumps(body))
        self.assertEqual(json.loads(content[0]['text']),{'scope':'set'})
        response.close.assert_called_once()

    def test_cache_changes_with_pixels_scope_model_and_endpoint(self):
        with patch.object(self.client,'_request',return_value=visual_result()) as request:
            self.assertFalse(self.client.analyze(self.data,'set')['cache_hit'])
            self.assertTrue(self.client.analyze(self.data,'set')['cache_hit'])
            self.client.analyze(self.data+b'extra','set')
            self.client.analyze(self.data,'activity')
            self.client.model='other';self.client.analyze(self.data,'set')
            self.client.api_url='https://second.invalid/chat';self.client.analyze(self.data,'set')
            self.assertEqual(request.call_count,5)
        for path in (self.root/'cache').glob('*.json'):
            self.assertNotIn('TEST_SECRET',path.read_text())

    def test_bad_and_unclear_output_are_distinct(self):
        bad=visual_result();bad['objects'][0].pop('location')
        with patch.object(self.client,'_request',return_value=bad),self.assertRaises(ValueError):
            self.client.analyze(self.data,'activity')
        self.assertFalse(list((self.root/'cache').glob('*.json')))
        unclear={'status':'unclear','summary':'图像过小。','objects':[],'visible_text':[],'uncertainties':['无法判断物体']}
        with patch.object(self.client,'_request',return_value=unclear):
            self.assertEqual(self.client.analyze(self.data,'activity')['status'],'unclear')

    def test_transient_retry_and_permanent_failure(self):
        busy=Mock(status_code=429);ok=Mock(status_code=200)
        ok.json.return_value={'choices':[{'message':{'content':json.dumps(visual_result())}}]}
        with patch('core_tools.vision_context.requests.post',side_effect=[requests.Timeout(),busy,ok]) as post,patch('core_tools.vision_context.time.sleep'):
            self.client.analyze(self.data,'activity');self.assertEqual(post.call_count,3)
        with patch('core_tools.vision_context.requests.post',return_value=Mock(status_code=401)) as post,self.assertRaises(ValueError):
            self.client.analyze(self.data,'decoration')
        self.assertEqual(post.call_count,1)

    def test_pipeline_deduplicates_images_preserves_bytes_and_uses_cache(self):
        path=self.fixture.source([['cat_multiboards_title','Cat event'],['cat_multiboards_pre_desc','Help cats']],images=[(2,3),(2,3)])
        self.fixture.classify(path)
        with patch.object(VisionContextAnalyzer,'_request',return_value=visual_result()) as request:
            first=run_context(self.fixture.ctx,self.root/'01-分类.xlsx',self.root/'02.xlsx')
            self.assertEqual(first['state'],'completed',first.get('status'))
            self.assertEqual(request.call_count,1)
            self.assertIn('视觉观察（模型识别，需核对）',first['records'][0]['context'])
            second=run_context(self.fixture.ctx,self.root/'01-分类.xlsx',self.root/'retry.xlsx')
            self.assertEqual(second['vision_cached_images'],1);self.assertEqual(request.call_count,1)
        original,result=load_workbook(path),load_workbook(self.root/'02.xlsx')
        try:
            self.assertEqual([im.anchor for im in original.active._images],[im.anchor for im in result.active._images])
            self.assertEqual(result.active['C2'].value,'KEEP')
        finally: original.close();result.close()
        with ZipFile(path) as a,ZipFile(self.root/'02.xlsx') as b:
            media=lambda z:sorted(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('xl/media/'))
            self.assertEqual(media(a),media(b))

    def test_ambiguous_images_are_not_sent_and_failure_does_not_publish(self):
        path=self.fixture.source([['mdnew_forestconcert_mdnew_decoration_name_1','Owl'],
                                  ['mdnew_forestconcert_mdnew_decoration_name_2','Fox']],images=[(2,3)])
        self.fixture.classify(path)
        with patch.object(VisionContextAnalyzer,'_request') as request:
            task=run_context(self.fixture.ctx,self.root/'01-分类.xlsx',self.root/'ambiguous.xlsx')
            self.assertEqual(task['state'],'completed',task.get('status'));request.assert_not_called()
        path=self.fixture.source([['cat_multiboards_title','Cats'],['cat_multiboards_pre_desc','Help cats']],images=[(2,3)])
        self.fixture.classify(path)
        with patch.object(VisionContextAnalyzer,'_request',side_effect=ValueError('API failed')):
            task=run_context(self.fixture.ctx,self.root/'01-分类.xlsx',self.root/'failed.xlsx')
        self.assertEqual(task['state'],'error');self.assertFalse((self.root/'failed.xlsx').exists())

    def test_vision_uncertainty_reaches_context_and_metadata(self):
        path=self.fixture.source([['cat_multiboards_title','Cats'],['cat_multiboards_pre_desc','Help cats']],images=[(2,3)])
        self.fixture.classify(path)
        result=visual_result();result.update(status='unclear',uncertainties=['文字过小，不能确认'])
        with patch.object(VisionContextAnalyzer,'_request',return_value=result):
            task=run_context(self.fixture.ctx,self.root/'01-分类.xlsx',self.root/'unclear.xlsx')
        self.assertEqual(task['state'],'completed',task.get('status'))
        self.assertIn('文字过小',task['records'][0]['context'])
        self.assertEqual(task['records'][0]['context_record']['context_status'],'partial')

    def test_failed_image_retry_reuses_other_successful_image(self):
        path=self.fixture.source([['cat_multiboards_title','Cats'],['cat_multiboards_pre_desc','Help cats']],images=[(2,3),(2,3)])
        buffer=io.BytesIO();Image.new('RGB',(16,16),'blue').save(buffer,format='PNG');blue=buffer.getvalue()
        book=load_workbook(path);book.active._images[0].ref=io.BytesIO(self.data);book.active._images[1].ref=io.BytesIO(blue);book.save(path);book.close()
        self.fixture.classify(path)
        def fail_blue(data,scope):
            if data==blue:raise ValueError('temporary failure')
            return visual_result()
        with patch.object(VisionContextAnalyzer,'_request',side_effect=fail_blue):
            first=run_context(self.fixture.ctx,self.root/'01-分类.xlsx',self.root/'failure.xlsx')
        self.assertEqual(first['state'],'error');self.assertFalse((self.root/'failure.xlsx').exists())
        with patch.object(VisionContextAnalyzer,'_request',return_value=visual_result()) as request:
            second=run_context(self.fixture.ctx,self.root/'01-分类.xlsx',self.root/'retry.xlsx')
        self.assertEqual(second['state'],'completed',second.get('status'))
        self.assertEqual(request.call_count,1);self.assertEqual(second['vision_cached_images'],1)


if __name__ == '__main__': unittest.main()
