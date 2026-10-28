import json
import shutil
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from core_tools.language import LanguagePair, language_pair, source_tokens
from core_tools.dialogue_context import DialogueContextAnalyzer
from core_tools.term_filter import TermFilter
from core_tools.term_scanner import TermScanner
from core_tools.tm_matcher import TMMatcher, differences
from core_tools.tm_semantic import SemanticTMRetriever
from core_tools.preflight import preflight
from pipeline_service import PipelineService, PipelineError


ZH = LanguagePair('zh-CN', 'en-US')
PREFIX = 'mdnew_forestconcert_mdnew_decoration_3_'
ROOT = Path(__file__).resolve().parents[1]


class Upload:
    def __init__(self, path):
        self.path = Path(path)
        self.filename = self.path.name

    def save(self, destination):
        shutil.copyfile(self.path, destination)


class LanguageSupportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.config = self.root / 'config.json'
        self.config.write_text(json.dumps({
            k: {'api_url': 'https://example.invalid/' + k, 'api_key': 'TEST_ONLY', 'model_name': 'test'}
            for k in ('llm', 'embedding', 'rerank')
        }))
        self.terms = self.root / 'terms.csv'
        self.terms.write_text('ZH,EN\n城堡,castle\n城堡大门,castle gate\n钥匙,key\nkey,key\n', encoding='utf-8')
        self.tmx = self.memory()
        self.service = PipelineService(workspace=self.root / 'workspace')
        self.network = patch('requests.sessions.Session.request', side_effect=AssertionError('NO_NETWORK'))
        self.network.start()

    def tearDown(self):
        for runtime in list(self.service.running.values()):
            runtime['thread'].join(timeout=5)
        self.network.stop()
        self.temp.cleanup()

    def memory(self, target_lang='en-US', name='memory.tmx'):
        p = self.root / name
        p.write_text('<tmx version="1.4"><body><tu>'
                     '<tuv xml:lang="zh-CN"><seg>请修复城堡大门。</seg></tuv>'
                     f'<tuv xml:lang="{target_lang}"><seg>Please repair the castle gate.</seg></tuv>'
                     '</tu></body></tmx>')
        return p

    def source(self, custom=False):
        book = Workbook()
        s = book.active
        s.title = '对白'
        s.append(['note', 'key', '中文台词' if custom else 'game_text_cn',
                  '英文稿' if custom else 'game_text_en', '表情' if custom else 'animation', 'bubbleType'])
        s.append([PREFIX + '2', 'Chad', '请修复城堡大门。', 'EXISTING_ENGLISH_DO_NOT_READ', 'happy', 1])
        s.append([PREFIX + '1', 'Christine', '钥匙在哪里？', None, 'thinking', 2])
        s.freeze_panes = 'C2'
        book.create_sheet('说明')['A1'] = '保留说明'
        p = self.root / ('custom.xlsx' if custom else 'input.xlsx')
        book.save(p)
        book.close()
        return p

    @staticmethod
    def scene(payload):
        first = payload['dialogue'][0]
        return {'scene_structure': '两名角色连续交流，其中一条已标注为思考。',
                'situation_tone': '角色寻找钥匙并讨论修复城堡大门，注意各句的不同意图。',
                'evidence': [{'key': first['key'], 'quote': first['source_text']}],
                'material_evidence': [], 'warnings': []}

    def wait(self, project, batch):
        deadline = time.monotonic() + 8
        while time.monotonic() < deadline:
            s = self.service.batch_state(project, batch)
            if not s['running']:
                return s
            time.sleep(.02)
        self.fail('worker did not finish')

    def test_chinese_fields_reject_emotion_alias_collision_and_ambiguous_sources(self):
        get = lambda *names: next(({'game_text_cn': 'happy', 'emotion': '高兴'}[n] for n in names if n in {'game_text_cn', 'emotion'}), '')
        self.assertEqual(ZH.emotion(get), '高兴')
        self.assertEqual(ZH.emotion(lambda *names: 'happy' if 'game_text_cn' in names else ''), '')
        self.assertEqual(LanguagePair().emotion(lambda *names: 'happy' if 'game_text_cn' in names else ''), 'happy')
        with self.assertRaisesRegex(ValueError, '同一列'):
            language_pair(ZH, {'source': 'CN', 'emotion': 'cn'})
        with self.assertRaisesRegex(ValueError, '多个候选'):
            ZH.headers(['game_text_cn', '中文', 'game_text_en'])
        with self.assertRaisesRegex(ValueError, '不能同时'):
            language_pair(ZH, {'source': 'animation'}).headers(['animation', 'game_text_en'])
        with self.assertRaisesRegex(ValueError, '当前支持'):
            LanguagePair('zh-CN', 'en-GB')

    def test_chinese_terms_match_contiguous_longest_and_latin_boundaries(self):
        scanner = TermScanner(self.terms, language=ZH)
        matches = scanner.find_matches('请修复<i>城堡大门</i>，找回钥匙。不要匹配keyboard，key可以。')
        self.assertEqual([m['source_text'] for m in matches], ['城堡大门', '钥匙', 'key'])
        self.assertEqual(scanner.find_matches('我们回到城堡。')[0]['target_text'], ['castle'])
        with self.assertRaisesRegex(ValueError, '当前语向表头'):
            TermScanner(self.terms)

    def test_chinese_tm_pair_labels_and_wrong_locale_rejection(self):
        matcher = TMMatcher(self.tmx, language=ZH)
        match = matcher.find_matches('请修复城堡大门。')
        self.assertEqual(match[0]['match_type'], '原文一致')
        self.assertIn('【中】请修复城堡大门。', matcher.format_matches(match))
        self.assertIn('【英】Please repair', matcher.format_matches(match))
        self.assertTrue(matcher.find_matches('请修复城堡大门！'))
        self.assertIn('当前原文与 TM 原文的否定表达不同', differences('不要开门', '请开门', 'Open it', ZH))
        with self.assertRaisesRegex(ValueError, '美式英语'):
            TMMatcher(self.memory('en-GB', 'british.tmx'), language=ZH)
        with self.assertRaisesRegex(ValueError, '英语→德语'):
            TMMatcher(self.tmx)

    def test_chinese_tokens_and_language_scoped_embedding_cache(self):
        self.assertIn('城堡', source_tokens('修复城堡大门', 'han'))
        self.assertIn('vip', source_tokens('获得VIP奖励', 'han'))
        with patch.object(SemanticTMRetriever, 'embed', return_value=__import__('numpy').array([[1., .5]])) as embed:
            one = SemanticTMRetriever(self.config, self.root / 'cache', language=ZH)
            one.prepare(['同一句'], 'same-tmx')
            two = SemanticTMRetriever(self.config, self.root / 'cache')
            two.prepare(['同一句'], 'same-tmx')
            self.assertEqual(embed.call_count, 2)
            three = SemanticTMRetriever(self.config, self.root / 'cache', language=ZH)
            three.prepare(['同一句'], 'same-tmx')
            self.assertTrue(three.cache_hit)
            self.assertEqual(embed.call_count, 2)

    def test_chinese_semantic_retrieval_with_mocked_models(self):
        matcher = TMMatcher(self.tmx, language=ZH, semantic=True, config_path=self.config, cache_dir=self.root / 'vectors')
        def response(section, payload, timeout, label):
            if isinstance(payload.get('input'), list):
                return {'data': [{'index': i, 'embedding': [1., .1]} for i, _ in enumerate(payload['input'])]}
            return {'results': [{'index': i, 'relevance_score': .9} for i, _ in enumerate(payload['documents'])]}
        with patch.object(SemanticTMRetriever, '_post', side_effect=response):
            matches = matcher.find_matches('我们把城堡入口修好吧。')
        self.assertEqual(matches[0]['match_type'], '语义相似')
        self.assertEqual(matches[0]['source_text'], '请修复城堡大门。')

    def test_filter_uses_chinese_source_and_us_english_without_current_target(self):
        f = TermFilter(self.config, language=ZH, materials={'cast': '完整人物信息' * 500})
        record = {'source_text': '钥匙在哪里？', 'context': '寻找钥匙', 'category': '活动对白',
                  'key': PREFIX + '1', 'speaker': 'Christine', 'emotion': 'thinking', 'candidates': ['钥匙 (key)']}
        response = Mock(ok=True)
        response.json.return_value = {'choices': [{'message': {'content': json.dumps({'decisions': [{'id': 0, 'decision': 'keep', 'reason': '实体钥匙'}]})}}]}
        with patch('core_tools.term_filter.requests.post', return_value=response) as post:
            result = f.filter_terms_with_llm(record)
        messages = post.call_args.kwargs['json']['messages']
        self.assertIn('美式英语', messages[0]['content'])
        self.assertIn('完整人物信息' * 500, messages[1]['content'])
        self.assertEqual(json.loads(messages[-1]['content'])['language']['source_lang'], 'zh-CN')
        self.assertEqual(result[0]['decision'], 'keep')

    def test_dialogue_cache_and_evidence_use_current_source_language(self):
        line = {'key': PREFIX + '1', 'source_text': '钥匙在哪里？', 'speaker': 'Christine', 'emotion': '', 'bubble': '2'}
        with patch.object(DialogueContextAnalyzer, '_request', side_effect=self.scene) as req:
            one = DialogueContextAnalyzer(self.config, cache_dir=self.root / 'scenes', language=ZH)
            one.analyze('scene', [line], [], [])
            one.analyze('scene', [line], [], [])
            two = DialogueContextAnalyzer(self.config, cache_dir=self.root / 'scenes')
            two.analyze('scene', [line], [], [])
            self.assertEqual(req.call_count, 2)
        bad = self.scene({'dialogue': [line]})
        bad['evidence'][0]['quote'] = 'Where is the key?'
        with patch.object(one, '_request', return_value=bad):
            with self.assertRaisesRegex(ValueError, '引用了不存在'):
                one.analyze('other-scene', [line], [], [])

    def test_service_full_six_stages_custom_columns_preserves_source_and_target(self):
        project = self.service.create_project('中英验收', 'zh-CN', 'en-US', use_tm=False)['project']
        pid = project['id']
        self.assertFalse(next(s for s in project['library'] if s['key'] == 'tm')['req'])
        self.service.upload_library(pid, 'termbase', Upload(self.terms))
        source = self.source(custom=True)
        mapping = {'source': '中文台词', 'target': '英文稿', 'emotion': '表情'}
        batch = self.service.create_batch(pid, Upload(source), mapping)['batch']
        original = source.read_bytes()
        model_payloads = []
        def scene(payload):
            model_payloads.append(payload)
            return self.scene(payload)
        def decide(record):
            return [{'term': t, 'decision': 'keep', 'reason': '测试候选'} for t in record['candidates']]
        with patch.object(DialogueContextAnalyzer, '_request', side_effect=scene), patch.object(TermFilter, 'filter_terms_with_llm', side_effect=decide):
            self.service.start(pid, batch['id'])
            state = self.wait(pid, batch['id'])
            for _ in range(3):
                self.assertFalse(any(s['state'] == 'failed' for s in state['stages']), str(state))
                if not any(s['state'] == 'review' for s in state['stages']):
                    break
                self.service.confirm(pid, batch['id'])
                state = self.wait(pid, batch['id'])
        self.assertEqual(state['done'], 6)
        self.assertEqual(source.read_bytes(), original)
        payload = model_payloads[0]
        self.assertEqual(payload['language']['target_lang'], 'en-US')
        self.assertEqual([r['source_text'] for r in payload['dialogue']], ['钥匙在哪里？', '请修复城堡大门。'])
        self.assertEqual(payload['dialogue'][1]['emotion'], 'happy')
        self.assertNotIn('EXISTING_ENGLISH', json.dumps(payload))
        output, _ = self.service.final_artifact(pid, batch['id'])
        book = load_workbook(output)
        header = {c.value: c.column for c in book['对白'][1]}
        self.assertEqual(book['对白']['C2'].value, '请修复城堡大门。')
        self.assertEqual(book['对白']['D2'].value, 'EXISTING_ENGLISH_DO_NOT_READ')
        self.assertEqual(book['说明']['A1'].value, '保留说明')
        self.assertEqual(book['对白'].freeze_panes, 'C2')
        self.assertIn('城堡大门 (castle gate)', book['对白'].cell(2, header['生效术语']).value)
        self.assertTrue(book['对白'].cell(2, header['Context']).value.startswith('【场景结构】'))
        self.assertIsNone(book['对白'].cell(2, header['TM参考']).value)
        book.close()
        self.assertTrue(any('不使用 TM' in c for c in state['stages'][-1]['checks']))
        copy = self.service.clone_batch(pid, batch['id'])['batch']
        self.assertEqual(copy['language']['columns'], mapping)
        loaded = PipelineService(workspace=self.service.workspace).batch_state(pid, batch['id'])
        self.assertEqual(loaded['language'], batch['language'])

    def test_missing_tm_is_optional_and_wrong_library_rejected(self):
        project = self.service.create_project('需要 TM', 'zh', 'en')['project']
        pid = project['id']
        self.service.upload_library(pid, 'termbase', Upload(self.terms))
        batch = self.service.create_batch(pid, Upload(self.source()))['batch']
        self.assertIsNone(self.service._project_tm(self.service._read_store()['projects'][pid]))
        with self.assertRaisesRegex(PipelineError, '美式英语'):
            self.service.upload_library(pid, 'tm', Upload(self.memory('de', 'wrong.tmx')))
        self.service.upload_library(pid, 'tm', Upload(self.tmx))
        self.assertIsNotNone(self.service._project_tm(self.service._read_store()['projects'][pid]))

    def test_emotion_only_source_is_rejected_at_upload_and_no_partial_batch_saved(self):
        pid = self.service.create_project('中英', 'zh', 'en', use_tm=False)['project']['id']
        p = self.root / 'wrong.csv'
        p.write_text('key,game_text_cn,game_text_en\nChad,happy,Hello\nMax,thinking,Hi\n')
        with self.assertRaisesRegex(PipelineError, '情绪标记'):
            self.service.create_batch(pid, Upload(p))
        self.assertEqual(self.service._all_batches(pid), [])

    def test_chinese_preflight_keeps_existing_target_and_parses_tm(self):
        src = self.source()
        book = load_workbook(src)
        sheet = book.active
        sheet.append([])
        sheet.cell(1, 7, 'TM参考')
        sheet.cell(2, 7, '[TM参考 1｜原文一致]\n【中】请修复城堡大门。\n【英】Please repair the castle gate.\n提示：仅供参考')
        final = self.root / 'final.xlsx'
        book.save(final)
        book.close()
        checks = {c['name']: c for c in preflight(src, final, language=ZH)['checks']}
        self.assertTrue(checks['目标列未写入']['ok'])
        self.assertTrue(checks['源文未改写']['ok'])
        self.assertTrue(checks['无悬空引用']['ok'])

    def test_http_create_upload_without_prompt_management(self):
        import app
        import pipeline_api
        from io import BytesIO
        with patch.object(pipeline_api, 'service', self.service):
            client = app.app.test_client()
            response = client.post('/api/pipeline/projects', json={'name': 'API中英', 'src': 'zh-CN', 'tgt': 'en-US', 'use_tm': False})
            self.assertEqual(response.status_code, 200)
            pid = response.json['project']['id']
            self.assertNotIn('prompts', response.json['project'])
            self.assertNotIn('promptLinks', client.get('/pipeline').get_data(as_text=True))
            for scene in ('01', '02', '03', '04'):
                prompt = client.get(f'/api/pipeline/projects/{pid}/prompts/{scene}')
                self.assertEqual(prompt.status_code, 404)
                prompt.close()
            source = self.source(custom=True)
            mapping = {'source': '中文台词', 'target': '英文稿', 'emotion': '表情'}
            result = client.post(f'/api/pipeline/projects/{pid}/batches', data={
                'columns_map': json.dumps(mapping), 'file': (BytesIO(source.read_bytes()), source.name)})
            self.assertEqual(result.status_code, 200, result.json)
            self.assertEqual(result.json['batch']['language']['columns'], mapping)
            bad = client.post(f'/api/pipeline/projects/{pid}/batches', data={
                'columns_map': '[]', 'file': (BytesIO(source.read_bytes()), source.name)})
            self.assertEqual(bad.status_code, 400)
            unsupported = client.post('/api/pipeline/projects', json={'name': 'bad', 'src': 'zh', 'tgt': 'fr'})
            self.assertEqual(unsupported.status_code, 400)

    def test_adds_missing_target_and_preserves_custom_header_case(self):
        from core_tools.normalizer import Normalizer
        source = self.root / 'no-target.csv'
        source.write_text('key,中文\na,城堡\n')
        pair = language_pair(ZH, {'source': '中文', 'target': 'English_Output'})
        result = self.root / 'normalized.xlsx'
        task = {}
        Normalizer(language=pair).process(source, result, {'t': task}, 't')
        self.assertEqual(task['state'], 'completed')
        book = load_workbook(result)
        self.assertEqual([c.value for c in book.active[1]], ['key', '中文', 'English_Output'])
        self.assertIsNone(book.active['C2'].value)
        book.close()

    def test_multisheet_source_count_and_english_regression_inputs(self):
        p = self.source()
        book = load_workbook(p)
        s = book.create_sheet('另一张表')
        s.append(['key', 'game_text_cn', 'game_text_en'])
        s.append(['title', '城堡', 'Castle'])
        book.save(p)
        book.close()
        self.assertEqual(self.service._inspect_source(p, ZH)[0], 3)
        default = self.service._project_view(self.service._read_store()['projects']['wc'])
        self.assertEqual((default['src'], default['tgt']), ('en', 'de'))
        self.assertTrue(default['use_tm'])

    def test_missing_cli_tm_is_reported_as_skipped(self):
        from pipeline import run_tm
        book = Workbook()
        book.active.append(['game_text_cn', 'game_text_en', '分类', 'Context', '生效术语', '无效术语'])
        book.active.append(['城堡', None, '集卡名称', '', '', ''])
        source = self.root / 'no_tm.xlsx'
        book.save(source)
        book.close()
        result = run_tm({'language': {**ZH.as_dict(), 'tms': []}}, source, self.root / 'out.xlsx')
        self.assertEqual(result['state'], 'completed', result.get('status'))
        self.assertTrue(result['skipped'])
        self.assertIn('未提供记忆库', result['status'])

    def test_all_four_categories_prepare_chinese_and_keep_non_data_sheets(self):
        from pipeline import STAGES
        book = Workbook()
        sheet = book.active
        sheet.append(['key', 'game_text_cn', 'game_text_en', 'limit'])
        rows = [
            ['cat_multiboards_title', '城堡活动', None, 30],
            ['cat_multiboards_pre_desc', '修复城堡，获得奖励。', None, None],
            ['Ahalloween_set2_name', '奇妙城堡', None, None],
            ['Ahalloween_set2_1_name', '城堡大门', None, None],
            ['mdnew_forestconcert_mdnew_decoration_name_3', '城堡大门', None, None],
            [PREFIX + '1', '钥匙在哪里？', None, None],
        ]
        for row in rows:
            sheet.append(row)
        note = book.create_sheet('备注页')
        note.append(['说明', '说明'])
        path = self.root / 'all.xlsx'
        book.save(path)
        book.close()

        ctx = {'game': {'key_rules': str(ROOT / 'profiles/game/wc_key_rules.json')},
               'language': {**ZH.as_dict(), 'termbases': [str(self.terms)], 'tms': [], 'use_tm': False},
               'config_path': self.config}
        tasks = {}
        with patch.object(DialogueContextAnalyzer, '_request', side_effect=self.scene), patch.object(TermFilter, 'filter_terms_with_llm', side_effect=lambda record: [
            {'term': t, 'decision': 'keep', 'reason': '测试'} for t in record['candidates']]):
            for key, _, artifact, runner in STAGES:
                target = self.root / artifact
                tasks[key] = runner(ctx, path, target)
                self.assertEqual(tasks[key]['state'], 'completed', tasks[key].get('status'))
                path = target
        self.assertEqual(set(r['category'] for r in tasks['classify']['records']),
                         {'活动对白', '活动 UI 与商店文案', '集卡名称', '装饰物名称'})
        self.assertEqual(tasks['tm']['counts']['未使用'], 6)
        self.assertEqual(tasks['tm']['counts']['未命中'], 0)
        book = load_workbook(path)
        self.assertEqual([list(r)[:4] for r in book.active.iter_rows(min_row=2, values_only=True)], rows)
        self.assertEqual([c.value for c in book['备注页'][1]], ['说明', '说明'])
        book.close()

    def test_story_note_keys_group_by_chapter_task_and_stage(self):
        from pipeline import run_context
        book = Workbook()
        book.active.append(['key', 'note', 'speaker', 'game_text_cn', 'game_text_en', '分类'])
        keys = ['story3_1_begin_10', 'story3_1_begin_2', 'story3_1_end_1',
                'story4_1_begin_1', 'Fire1_1']
        for key in keys:
            book.active.append(['Christine', key, 'Christine', '找到钥匙了。', None, '活动对白'])
        source = self.root / 'story.xlsx'
        book.save(source)
        book.close()
        with patch.object(DialogueContextAnalyzer, '_request', side_effect=self.scene) as req:
            result = run_context({'game': {}, 'language': ZH.as_dict(), 'config_path': self.config}, source, self.root / 'story_context.xlsx')
        self.assertEqual(result['state'], 'completed', result.get('status'))
        self.assertEqual(req.call_count, 4)
        self.assertEqual([r['key'] for r in req.call_args_list[0].args[0]['dialogue']], keys[1::-1])


if __name__ == '__main__':
    unittest.main()
