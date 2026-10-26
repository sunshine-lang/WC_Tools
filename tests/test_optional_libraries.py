import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import test_language_support as fixtures
from core_tools.dialogue_context import DialogueContextAnalyzer
from core_tools.term_filter import TermFilter
from core_tools.tm_semantic import SemanticTMRetriever
from pipeline_service import PipelineError
from pipeline import run_term_filter, run_term


class OptionalLibraryTests(unittest.TestCase):
    def setUp(self):
        self.f = fixtures.LanguageSupportTests()
        self.f.setUp()
        contents = self.f.tmx.read_text().replace('</body>',
            '<tu><tuv xml:lang="zh-CN"><seg>钥匙在哪里？</seg></tuv>'
            '<tuv xml:lang="en-US"><seg>Where is the key?</seg></tuv></tu></body>')
        self.f.tmx.write_text(contents)

    def tearDown(self):
        self.f.tearDown()

    def test_all_library_combinations_finish_and_only_review_executed_work(self):
        f = self.f
        for term, tm in [(False, False), (True, False), (False, True), (True, True)]:
            with self.subTest(term=term, tm=tm):
                service = f.service
                pid = service.create_project(f'库组合 {term}/{tm}', 'zh-CN', 'en-US')['project']['id']
                if term:
                    service.upload_library(pid, 'termbase', fixtures.Upload(f.terms))
                if tm:
                    service.upload_library(pid, 'tm', fixtures.Upload(f.tmx))
                source = f.source()
                batch = service.create_batch(pid, fixtures.Upload(source))['batch']
                reviewed = []
                with patch.object(DialogueContextAnalyzer, '_request', side_effect=f.scene), patch.object(TermFilter, 'filter_terms_with_llm', side_effect=lambda record: [
                    {'term': t, 'decision': 'keep', 'reason': '本地测试'} for t in record['candidates']]) as filter_call, patch.object(SemanticTMRetriever, '_post', side_effect=AssertionError('exact/no-TM must not call model')):
                    service.start(pid, batch['id'])
                    state = f.wait(pid, batch['id'])
                    for _ in range(3):
                        self.assertFalse(any(s['state'] == 'failed' for s in state['stages']), str(state))
                        pending = next((s for s in state['stages'] if s['state'] == 'review'), None)
                        if not pending:
                            break
                        reviewed.append(pending['key'])
                        service.confirm(pid, batch['id'])
                        state = f.wait(pid, batch['id'])
                self.assertEqual(state['done'], 6)
                self.assertEqual(reviewed, ['context'] + (['term_filter'] if term else []) + (['tm'] if tm else []))
                self.assertEqual(filter_call.call_count > 0, term)
                stages = {s['key']: s for s in state['stages']}
                self.assertEqual(stages['term']['skipped'], not term)
                self.assertEqual(stages['term_filter']['skipped'], not term)
                self.assertEqual(stages['tm']['skipped'], not tm)
                output, _ = service.final_artifact(pid, batch['id'])
                before, after = load_workbook(source), load_workbook(output)
                try:
                    for name in before.sheetnames:
                        self.assertEqual([[c.value for c in r] for r in before[name]],
                                         [[c.value for c in r[:before[name].max_column]] for r in after[name]])
                    cols = {c.value: c.column for c in after.active[1]}
                    self.assertEqual(bool(after.active.cell(2, cols['生效术语']).value), term)
                    self.assertFalse(after.active.cell(2, cols['无效术语']).value)
                    self.assertEqual(bool(after.active.cell(2, cols['TM参考']).value), tm)
                    self.assertTrue(after.active.cell(2, cols['Context']).value)
                finally:
                    before.close()
                    after.close()

    def test_existing_project_becomes_optional_without_migration_and_upload_reenables_tm(self):
        f = self.f
        pid = f.service.create_project('旧项目', 'zh', 'en', use_tm=False)['project']['id']
        before = f.service.projects_file.read_bytes()
        view = f.service.bootstrap(pid)['project']
        self.assertEqual(f.service.projects_file.read_bytes(), before)
        self.assertTrue(all(not slot['req'] for slot in view['library']))
        self.assertTrue(all(not slot.get('disabled') for slot in view['library']))
        f.service.upload_library(pid, 'tm', fixtures.Upload(f.tmx))
        project = f.service._read_store()['projects'][pid]
        self.assertTrue(project['use_tm'])
        self.assertTrue(f.service._project_tm(project).is_file())

    def test_registered_missing_or_corrupt_library_does_not_become_skip(self):
        f = self.f
        pid = f.service.create_project('错误文件', 'zh', 'en')['project']['id']
        source = f.source()
        batch = f.service.create_batch(pid, fixtures.Upload(source))['batch']
        for slot, file in [('termbase', f.terms), ('tm', f.tmx)]:
            with self.subTest(slot=slot):
                f.service.upload_library(pid, slot, fixtures.Upload(file))
                project = f.service._read_store()['projects'][pid]
                stored = f.service._resolve(project['library'][slot]['path'])
                stored.unlink()
                with self.assertRaisesRegex(PipelineError, '已登记'):
                    f.service.start(pid, batch['id'])
                f.service.remove_library(pid, slot)
        f.service.upload_library(pid, 'termbase', fixtures.Upload(f.terms))
        project = f.service._read_store()['projects'][pid]
        f.service._project_term(project).write_text('bad,headers\n1,2\n')
        with self.assertRaisesRegex(ValueError, '表头'):
            run_term(f.service._ctx(project, batch), source, f.root / 'out.xlsx')

    def test_no_termbase_drops_stale_candidates_without_model_or_config(self):
        f = self.f
        book = Workbook()
        book.active.append(['game_text_cn', 'game_text_en', '分类', 'Context', '关联术语'])
        book.active.append(['城堡', 'KEEP', '集卡名称', '', '城堡 (castle)'])
        source, output = f.root / 'old_terms.xlsx', f.root / 'filtered.xlsx'
        book.save(source)
        book.close()
        with patch.object(TermFilter, 'filter_terms_with_llm', side_effect=AssertionError('must skip model')):
            task = run_term_filter({'language': {**fixtures.ZH.as_dict(), 'termbases': []},
                                    'config_path': f.root / 'absent-config.json'}, source, output)
        self.assertEqual(task['state'], 'completed', task.get('status'))
        self.assertTrue(task['skipped'])
        self.assertEqual(task['model_requests'], 0)
        result = load_workbook(output)
        self.assertEqual([c.value for c in result.active[1]][-2:], ['生效术语', '无效术语'])
        self.assertEqual([c.value for c in result.active[2]][-2:], [None, None])
        result.close()


if __name__ == '__main__':
    unittest.main()
