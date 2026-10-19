import io
import time
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core_tools.tm_matcher import TMMatcher, context_text, differences
from core_tools.term_filter import FILTER_COLUMNS


class TMMatcherTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.tm = self.root / 'customer.tmx'
        self.database = self.root / 'database'
        self.database.mkdir()
        self.save_tm([('Collect {0} stars!', 'Sammle {0} Sterne!')])

    def tearDown(self):
        self.temp.cleanup()

    def save_tm(self, pairs, languages=('en-US', 'de-DE')):
        root = ET.Element('tmx', version='1.4')
        ET.SubElement(root, 'header', srclang=languages[0])
        body = ET.SubElement(root, 'body')
        for index, pair in enumerate(pairs, 1):
            unit = ET.SubElement(body, 'tu', tuid=str(index), changedate=f'202610{index:02d}T000000Z')
            ET.SubElement(unit, 'prop', type='x-document').text = f'original{index}.xlsx'
            for lang, value in zip(languages, pair):
                tuv = ET.SubElement(unit, 'tuv', {'{http://www.w3.org/XML/1998/namespace}lang': lang})
                ET.SubElement(tuv, 'prop', type='x-context-pre').text = 'Previous context'
                ET.SubElement(tuv, 'seg').text = value
        ET.ElementTree(root).write(self.tm, encoding='utf-8', xml_declaration=True)
        return root

    def source(self, rows=None, headers=None):
        w = Workbook()
        s = w.active
        s.title = '活动ui'
        s.append(headers or ['limit', 'key', 'game_text_cn', 'game_text_en', 'game_text_de', '分类', 'Context', *FILTER_COLUMNS])
        for row in rows or [[30, 'event', '中文', 'Collect {0} stars!', 'KEEP_TARGET', '活动 UI 与商店文案', 'KEEP_CONTEXT', 'KEEP_TERMS']]:
            s.append(row[:7] + ['KEEP_ACTIVE', ''] if headers is None else row)
        s['D2'].font = Font(bold=True, color='112233')
        s.freeze_panes = 'B2'
        s.row_dimensions[2].height = 42
        w.create_sheet('说明')['A1'] = 'KEEP_OTHER_SHEET'
        path = self.root / 'step4.xlsx'
        w.save(path)
        w.close()
        return path

    def run_match(self, path, input_name='step4.xlsx'):
        task = {'state': 'processing', 'input_name': input_name}
        TMMatcher(self.tm).process(path, self.root / 'result.xlsx', {'t': task}, 't')
        return task

    def test_keeps_all_translation_variants_and_duplicate_sources(self):
        self.save_tm([('Same source', 'Variante A'), ('Same source', 'Variante A'), ('Same source', 'Variante B')])
        tm = TMMatcher(self.tm, source_name='客户历史.tmx')
        self.assertEqual(tm.status()['valid_count'], 3)
        self.assertEqual(tm.status()['source_count'], 1)
        self.assertEqual(tm.status()['pair_count'], 2)
        self.assertEqual(tm.status()['conflict_count'], 1)
        matches = tm.find_matches('Same source')
        self.assertEqual([m['german'] for m in matches], ['Variante A', 'Variante B'])
        self.assertEqual(len(matches[0]['origins']), 2)
        self.assertEqual(matches[0]['origins'][0]['document'], 'original1.xlsx')
        self.assertEqual(matches[0]['origins'][0]['previous'], 'Previous context')
        self.assertTrue(all(m['warnings'] for m in matches))
        self.assertIn('客户历史.tmx', tm.format_matches(matches))

    def test_restores_memoq_tags_parameters_and_literal_newline_from_utf16(self):
        root = self.save_tm([('', '')])
        for seg in root.findall('.//seg'):
            ET.SubElement(seg, 'bpt').text = '<rpr id="0">'
            seg[-1].tail = 'Collect '
            ET.SubElement(seg, 'ph').text = '<mq:rxt val="&lt;color=#fff&gt;" />'
            ET.SubElement(seg, 'ph').text = '<mq:rxt val="{0}" />'
            ET.SubElement(seg, 'ph').text = '<mq:rxt val="&lt;/color&gt;" />'
            ET.SubElement(seg, 'ph').text = '<mq:ch val="\n" />'
            ET.SubElement(seg, 'ept').text = '</rpr>'
            seg[-1].tail = 'stars!'
        xml = '<?xml version="1.0" encoding="utf-16"?><!DOCTYPE tmx SYSTEM "tmx14.dtd">' + ET.tostring(root, encoding='unicode')
        self.tm.write_bytes(xml.encode('utf-16'))
        tm = TMMatcher(self.tm)
        self.assertEqual(tm.groups[0]['english'], 'Collect <color=#fff>{0}</color>\nstars!')
        self.assertEqual(tm.groups[0]['variants'][0]['german'], tm.groups[0]['english'])
        self.assertNotIn('mq:', tm.format_matches(tm.find_matches(tm.groups[0]['english'])))

    def test_wrong_language_empty_and_ambiguous_units_do_not_silently_pass(self):
        self.save_tm([('中文', 'English')], ('zh-CN', 'en-US'))
        with self.assertRaisesRegex(ValueError, '英语→德语'):
            TMMatcher(self.tm)
        self.save_tm([('Good', 'Gut'), ('Empty', '')])
        tm = TMMatcher(self.tm)
        self.assertEqual((tm.valid_count, tm.skipped_count), (1, 1))
        root = self.save_tm([('One', 'Eins')])
        tuv = ET.SubElement(root.find('.//tu'), 'tuv', {'lang': 'en-GB'})
        ET.SubElement(tuv, 'seg').text = 'Another source'
        ET.ElementTree(root).write(self.tm)
        with self.assertRaisesRegex(ValueError, '多个英语'):
            TMMatcher(self.tm)

    def test_decodes_historical_context_and_preserves_unknown_markup(self):
        root = self.save_tm([('One', 'Eins')])
        prop = root.find('.//tuv/prop')
        seg = ET.Element('seg')
        seg.text = 'Current score: '
        ET.SubElement(seg, 'ph').text = '<mq:rxt val="{0}" />'
        prop.text = ET.tostring(seg, encoding='unicode')
        ET.ElementTree(root).write(self.tm, encoding='utf-8')
        origin = TMMatcher(self.tm).find_matches('One')[0]['origins'][0]
        self.assertEqual(origin['previous'], 'Current score: {0}')
        self.assertEqual(context_text('ordinary context'), 'ordinary context')
        self.assertEqual(context_text('<seg>broken'), '原始上下文（格式未识别）：<seg>broken')

    def test_rejects_entities_unsupported_inline_and_done_sources(self):
        self.tm.write_text('<!DOCTYPE tmx [<!ENTITY x "bad">]><tmx/>')
        with self.assertRaisesRegex(ValueError, '实体'):
            TMMatcher(self.tm)
        root = self.save_tm([('One', 'Eins')])
        ET.SubElement(root.find('.//seg'), 'unknown').text = 'hidden content'
        ET.ElementTree(root).write(self.tm)
        with self.assertRaisesRegex(ValueError, '行内元素'):
            TMMatcher(self.tm)
        with self.assertRaisesRegex(ValueError, 'done'):
            TMMatcher(self.tm, source_name='answer_done.tmx')

    def test_exact_first_fuzzy_changes_and_short_unrelated_texts(self):
        self.save_tm([('Collect 5 golden stars to unlock the special gift.', 'Sammle 5 goldene Sterne.'),
                      ('Tap', 'Tippen'), ('Forest Friends', 'Waldfreunde')])
        tm = TMMatcher(self.tm)
        exact = tm.find_matches('Tap')
        self.assertEqual(exact[0]['match_type'], '原文一致')
        self.assertEqual(tm.find_matches('Tadpole'), [])
        self.assertEqual(tm.find_matches('Forest Concert'), [])
        self.assertEqual(tm.find_matches('Maximum power!'), [])
        fuzzy = tm.find_matches('Collect 10 golden stars to unlock the special gift.')
        self.assertEqual(len(fuzzy), 1)
        self.assertEqual(fuzzy[0]['match_type'], '文字相似')
        self.assertIn('当前原文与 TM 原文的数值不同', fuzzy[0]['warnings'])
        self.assertTrue(any('10' in change['current'] and '5' in change['history'] for change in fuzzy[0]['changes']))

    def test_protected_differences_are_visible(self):
        warnings = differences('Do not get {1} <sprite name="new">\n', 'Get {0} <sprite name="old">', 'Hol {2}')
        for word in ['参数不同', '参数不一致', '标签不同', '换行不同', '否定表达不同']:
            self.assertTrue(any(word in message for message in warnings), warnings)
        self.assertEqual(differences('Hello', 'Hello', 'Hallo'), [])

    def test_workbook_preservation_no_match_and_formula_pending(self):
        path = self.source([[30, 'event', '中文', 'Collect {0} stars!', 'KEEP_TARGET', '活动 UI 与商店文案', 'KEEP_CONTEXT', 'KEEP_TERMS'],
                            [None, 'new', '', 'Brand new unique name', '', '集卡名称', '', ''],
                            [None, 'bad', '', '=1+1', '', '待确认', '', '']])
        task = self.run_match(path)
        self.assertEqual(task['state'], 'completed', task)
        self.assertEqual(task['counts'], {'原文一致': 1, '相似参考': 0, '未命中': 1, '待确认': 1})
        self.assertEqual(task['matched_count'], 1)
        before, after = load_workbook(path), load_workbook(task['result_file'])
        try:
            self.assertEqual(after.sheetnames, before.sheetnames)
            for sheet in before:
                for row in sheet:
                    for cell in row:
                        new = after[sheet.title].cell(cell.row, cell.column)
                        self.assertEqual((cell.value, cell.data_type, cell._style), (new.value, new.data_type, new._style))
                self.assertEqual(sheet.freeze_panes, after[sheet.title].freeze_panes)
                self.assertEqual(sheet.row_dimensions[2].height, after[sheet.title].row_dimensions[2].height)
            self.assertEqual(after.active['J1'].value, 'TM参考')
            self.assertIn('Sammle {0} Sterne!', after.active['J2'].value)
            self.assertIsNone(after.active['J3'].value)
            self.assertEqual(after.active['D4'].data_type, 'f')
        finally:
            before.close(); after.close()

    def test_wrong_stage_repeated_output_and_bad_file_have_error_status(self):
        for headers in [['game_text_en'], ['game_text_en', '分类', 'Context', '关联术语'], ['game_text_en', '分类', 'Context', '关联术语', *FILTER_COLUMNS, 'TM参考']]:
            task = self.run_match(self.source([['Something']], headers))
            self.assertEqual(task['state'], 'error')
        self.assertEqual(self.run_match(self.source(), 'file_done.xlsx')['state'], 'error')
        bad = self.root / 'bad.xlsx'; bad.write_bytes(b'bad')
        self.assertEqual(self.run_match(bad)['state'], 'error')

    def test_accepts_only_active_and_invalid_terms_without_auxiliary_columns(self):
        path = self.source()
        book = load_workbook(path)
        self.assertEqual([c.value for c in book.active[1]][-2:], ['生效术语', '无效术语'])
        self.assertNotIn('关联术语', [c.value for c in book.active[1]])
        book.close()
        task = self.run_match(path)
        self.assertEqual(task['state'], 'completed', task)
        self.assertEqual(task['records'][0]['status'], '原文一致')
        self.assertNotIn('filter_status', task['records'][0])
        book = load_workbook(path)
        book.active['H2'] = None
        book.save(path)
        book.close()
        task = self.run_match(path)
        self.assertEqual(task['state'], 'completed', task)

    def test_http_database_replacement_failed_upload_and_snapshot_reload(self):
        import app as server
        with patch.dict(server.app.config, {'DB_FOLDER': str(self.database)}):
            client = server.app.test_client()
            self.assertFalse(server.get_tm_db_status()['ready'])
            response = client.post('/update/tm_db', data={'tm_tmx_file': (io.BytesIO(self.tm.read_bytes()), '客户TM.tmx')})
            self.assertEqual(response.status_code, 200, response.json)
            old, _ = server.load_tm_db()
            status = server.get_tm_db_status()
            files = set(self.database.iterdir())
            for data, name in [(b'bad', 'broken.tmx'), (b'<tmx><body/></tmx>', 'empty.tmx'), (self.tm.read_bytes(), 'file_done.tmx')]:
                response = client.post('/update/tm_db', data={'tm_tmx_file': (io.BytesIO(data), name)})
                self.assertEqual(response.status_code, 400)
                self.assertEqual(server.get_tm_db_status(), status)
                self.assertEqual(set(self.database.iterdir()), files)
            with patch.object(server, 'write_metadata', side_effect=OSError('write failure')):
                response = client.post('/update/tm_db', data={'tm_tmx_file': (io.BytesIO(self.tm.read_bytes()), 'new.tmx')})
                self.assertEqual(response.status_code, 400)
            self.assertEqual(server.get_tm_db_status(), status)
            self.assertEqual(set(self.database.iterdir()), files)
            self.save_tm([('Collect {0} stars!', 'CUSTOMER TARGET {0}')])
            client.post('/update/tm_db', data={'tm_tmx_file': (io.BytesIO(self.tm.read_bytes()), 'new.tmx')})
            new, _ = server.load_tm_db()
            self.assertEqual(new.find_matches('Collect {0} stars!')[0]['german'], 'CUSTOMER TARGET {0}')
            self.assertEqual(old.find_matches('Collect {0} stars!')[0]['german'], 'Sammle {0} Sterne!')
            self.assertNotEqual(old.sha256, new.sha256)

    def test_http_fourth_stage_upload_preview_and_download(self):
        import app as server
        task_id = None
        path = self.source()
        with patch.dict(server.app.config, {'DB_FOLDER': str(self.database), 'UPLOAD_FOLDER': str(self.root)}):
            client = server.app.test_client()
            self.assertEqual(client.get('/tm').status_code, 200)
            self.assertEqual(client.post('/upload/tm', data={'trans_file': (io.BytesIO(path.read_bytes()), 'step4.xlsx')}).status_code, 400)
            client.post('/update/tm_db', data={'tm_tmx_file': (io.BytesIO(self.tm.read_bytes()), 'customer.tmx')})
            self.assertEqual(client.post('/upload/tm', data={'trans_file': (io.BytesIO(b'bad'), 'bad.csv')}).status_code, 400)
            try:
                response = client.post('/upload/tm', data={'trans_file': (io.BytesIO(path.read_bytes()), 'step4.xlsx')})
                task_id = response.json['task_id']
                for _ in range(100):
                    progress = client.get('/progress/' + task_id).json
                    if progress['state'] != 'processing': break
                    time.sleep(.02)
                self.assertEqual(progress['state'], 'completed', progress)
                self.assertNotIn('records', progress)
                preview = client.get('/api/tm/' + task_id)
                self.assertEqual(preview.json['tm_db']['filename'], 'customer.tmx')
                self.assertNotIn('KEEP_TARGET', preview.get_data(as_text=True))
                self.assertIn('Sammle {0} Sterne!', preview.json['records'][0]['reference'])
                response = client.get('/download/' + task_id)
                book = load_workbook(io.BytesIO(response.data))
                self.assertEqual(book.active['J1'].value, 'TM参考')
                self.assertEqual(book.active['E2'].value, 'KEEP_TARGET')
                book.close(); response.close()
                self.assertEqual(client.get('/api/tm/missing').status_code, 404)
                self.assertEqual(client.post('/generate/tm_db').status_code, 404)
            finally:
                if task_id: server.tasks.pop(task_id, None)


if __name__ == '__main__':
    unittest.main()
