import base64
import io
from pathlib import Path
from tempfile import TemporaryDirectory
import time
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Font

from core_tools.context_processor import ContextProcessor, validate_material_name
from core_tools.key_classifier import CATEGORIES
from pipeline import run_normalize, run_classify


class ContextProcessorTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.reference = self.root / 'reference.xlsx'
        self.guide = self.root / 'guide.xlsx'
        guide = Workbook()
        sheet = guide.active
        sheet.title = '人设表'
        sheet.cell(3, 4, 'Max')
        sheet.cell(3, 5, 8)
        sheet.cell(3, 7, '喜欢绘画。喜欢解谜。后续剧情不应进入。')
        sheet.cell(3, 10, 'FOREIGN_RULE_DO_NOT_INCLUDE')
        guide.save(self.guide)
        self.processor = ContextProcessor(self.reference, self.guide, materials_dir=self.root)

    def tearDown(self):
        self.temp.cleanup()

    def save_reference(self, rows, images=False):
        book = Workbook()
        sheet = book.active
        sheet.title = '参考'
        sheet.append(['key', 'Speaker', 'Animation', 'BubbleType', 'game_text_cn', 'game_text_en', 'game_text_de', 'note'])
        for row in rows:
            sheet.append(row)
        if images:
            png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')
            for start in [1, 3]:
                image = Image(io.BytesIO(png))
                image.anchor = TwoCellAnchor(_from=AnchorMarker(col=0, row=start), to=AnchorMarker(col=1, row=start + 2))
                sheet.add_image(image)
        book.save(self.reference)

    def main(self, rows, classified=True):
        book = Workbook()
        sheet = book.active
        sheet.append(['limit', 'key', 'game_text_cn', 'game_text_en', 'game_text_de'] + (['分类'] if classified else []))
        for row in rows:
            sheet.append(row)
        sheet['D2'].font = Font(bold=True, color='123456')
        path = self.root / 'classified.xlsx'
        book.save(path)
        return path

    def run_context(self, path, enhance=True):
        task = {'state': 'processing', 'dl_name': 'result.xlsx'}
        self.processor.process(str(path), str(self.root / 'result.xlsx'), {'test': task}, 'test', enhance=enhance)
        self.assertEqual(task['state'], 'completed', task.get('status'))
        return task

    def test_restores_dialogue_orders_numerically_and_keeps_source(self):
        prefix = 'mdnew_forestconcert_mdnew_decoration_3_'
        self.save_reference([
            [prefix + '10', 'Max', 'happy', 1, '第十句', 'Ten', 'DE_SENTINEL'],
            [prefix + '2', 'Max', 'happy', 2, '第二句', '  <i>Two</i> {0}  ', 'DE_SENTINEL'],
            ['mdnew_forestconcert_mdnew_decoration_5_1', 'Max', 'neutral', 1, '', 'Unrelated group', 'DE_SENTINEL'],
        ])
        path = self.main([[None, 'Max', 'happy', '  <i>Two</i> {0}  ', 'EXISTING_TARGET', CATEGORIES[3]],
                          [None, 'Max', 'happy', 'Ten', None, CATEGORIES[3]]])
        before = path.read_bytes()
        task = self.run_context(path)
        record = task['records'][0]
        self.assertEqual(record['key'], prefix + '2')
        self.assertEqual(record['status'], '已关联')
        context = record['context']
        self.assertLess(context.index('Two'), context.index('Ten'))
        self.assertIn('(Max / happy / 思考)', context)
        self.assertEqual(context, task['records'][1]['context'])
        for excluded in ['DE_SENTINEL', 'FOREIGN_RULE_DO_NOT_INCLUDE', 'Unrelated group', '后续剧情不应进入']:
            self.assertNotIn(excluded, context)
        sheet = load_workbook(task['result_file']).active
        self.assertEqual([c.value for c in sheet[1]], ['limit', 'key', 'game_text_cn', 'game_text_en', 'game_text_de', '分类', 'Context'])
        self.assertEqual(sheet['D2'].value, '  <i>Two</i> {0}  ')
        self.assertEqual(sheet['E2'].value, 'EXISTING_TARGET')
        self.assertEqual(sheet['D2'].font.color.rgb, '00123456')
        self.assertEqual(path.read_bytes(), before)

    def test_source_changes_and_classification_conflicts_are_not_silent(self):
        self.save_reference([['cat_multiboards_pre_desc', '', '', '', '', 'Old source', 'DE_SENTINEL']])
        path = self.main([[None, 'cat_multiboards_pre_desc', '', 'New source', None, CATEGORIES[0]]])
        record = self.run_context(path)['records'][0]
        self.assertEqual(record['status'], '待确认')
        self.assertEqual(record['context'], '')
        self.assertIn('参考英文', record['warning'])
        path = self.main([[None, 'cat_multiboards_pre_desc', '', 'Old source', None, CATEGORIES[1]]])
        self.assertIn('分类与参考', self.run_context(path)['records'][0]['warning'])

    def test_same_english_uses_key_and_ambiguous_dialogue_is_flagged(self):
        keys = ['event_forestconcert_mdnew_prediction_title', 'event_forestconcert_mdnew_decorations_title']
        description = 'event_forestconcert_mdnew_prediction_desc'
        self.save_reference([[k, '', '', '', '', 'Forest Friends', 'DE_SENTINEL'] for k in keys] +
                            [[description, '', '', '', '', 'Solve the woodland challenge.', 'DE_SENTINEL']])
        path = self.main([[25, k, '', 'Forest Friends', None, CATEGORIES[0]] for k in keys] +
                         [[None, description, '', 'Solve the woodland challenge.', None, CATEGORIES[0]]])
        records = self.run_context(path)['records']
        self.assertEqual([r['key'] for r in records[:2]], keys)
        for record in records:
            self.assertEqual(record['status'], '已关联')
            self.assertIn(record['key'], record['context'])
            self.assertNotIn('← 当前句', record['context'])
        self.assertNotEqual(records[0]['context'], records[1]['context'])
        self.save_reference([[f'mdnew_forestconcert_mdnew_decoration_{n}_1', 'Max', 'happy', 1, '', 'Same', 'DE_SENTINEL'] for n in [3, 5]])
        path = self.main([[None, 'Max', 'happy', 'Same', None, CATEGORIES[3]]])
        self.assertEqual(self.run_context(path)['records'][0]['status'], '待确认')

    def test_missing_role_is_visible_and_classified_input_is_required(self):
        self.save_reference([['mdnew_forestconcert_mdnew_decoration_3_1', 'Unknown', 'happy', 1, '', 'Hello', 'DE_SENTINEL']])
        path = self.main([[None, 'Unknown', 'happy', 'Hello', None, CATEGORIES[3]]])
        self.assertEqual(self.run_context(path)['records'][0]['status'], '缺资料')
        path = self.main([[None, 'Unknown', 'happy', 'Hello', None]], classified=False)
        task = {'state': 'processing'}
        self.processor.process(str(path), str(self.root / 'bad.xlsx'), {'t': task}, 't')
        self.assertEqual(task['state'], 'error')
        self.assertIn('第一步', task['status'])

    def test_cross_segment_reference_is_checked_against_source(self):
        self.save_reference([
            ['mdnew_forestconcert_mdnew_decoration_5_10', 'Max', 'happy', 1, '', 'They are treble clefs.', 'DE_SENTINEL'],
            ['mdnew_forestconcert_mdnew_decoration_12_3', 'Max', 'happy', 1, '', 'We saw a treble clef earlier.', 'DE_SENTINEL'],
        ])
        path = self.main([[None, 'Max', 'happy', 'They are treble clefs.', None, CATEGORIES[3]],
                          [None, 'Max', 'happy', 'We saw a treble clef earlier.', None, CATEGORIES[3]]])
        self.assertIn('mdnew_forestconcert_mdnew_decoration_5_10 ', self.run_context(path)['records'][1]['context'])
        book = load_workbook(self.reference)
        book.active['F2'] = 'Changed clue'
        book.save(self.reference)
        record = self.run_context(path)['records'][1]
        self.assertEqual(record['status'], '缺资料')
        self.assertNotIn('mdnew_forestconcert_mdnew_decoration_5_10 ', record['context'])

    def test_card_images_and_bundle_do_not_cross_sets(self):
        self.save_reference([
            ['Ahalloween_set1_name', '', '', '', '', 'Set one', 'DE_SENTINEL'],
            ['Ahalloween_set1_1_name', '', '', '', '', 'First card', 'DE_SENTINEL'],
            ['Ahalloween_set2_name', '', '', '', '', 'Set two', 'DE_SENTINEL'],
            ['Ahalloween_set2_1_name', '', '', '', '', 'Second card', 'DE_SENTINEL'],
        ], images=True)
        path = self.main([[24, 'Ahalloween_set1_1_name', '', 'First card', None, CATEGORIES[1]]])
        task = self.run_context(path)
        record = task['records'][0]
        self.assertEqual(record['status'], '已关联')
        self.assertEqual(len(record['images']), 1)
        self.assertEqual(record['images'][0]['group'], 'halloween/set/1')
        self.assertNotIn('Second card', record['context'])
        with ZipFile(task['bundle_file']) as archive:
            self.assertIn('result.xlsx', archive.namelist())
            self.assertIn(record['images'][0]['path'], archive.namelist())


    def test_core_mode_reads_only_uploaded_file_and_uses_its_peers(self):
        path = self.main([[30, 'cat_multiboards_pre_desc', '', 'Current intro', None, CATEGORIES[0]],
                          [30, 'cat_multiboards_guide_1_1', '', 'Current guide', None, CATEGORIES[0]],
                          [None, 'Max', 'happy', 'Cannot infer its segment', None, CATEGORIES[3]]])
        self.reference = self.root / 'missing_reference.xlsx'
        self.processor = ContextProcessor(self.reference, self.root / 'missing_guide.xlsx', materials_dir=self.root)
        original_loader = load_workbook
        def only_input(file, *args, **kwargs):
            self.assertEqual(Path(file), path)
            return original_loader(file, *args, **kwargs)
        with patch('core_tools.context_processor.load_workbook', side_effect=only_input) as loader:
            task = self.run_context(path, enhance=False)
            self.assertEqual(loader.call_count, 1)
        self.assertFalse(task['enhanced'])
        self.assertEqual(task['counts'], {'已关联': 2, '缺资料': 0, '待确认': 1})
        self.assertIn('Current guide', task['records'][0]['context'])
        self.assertEqual(task['records'][2]['context'], '')
        self.assertIn('缺少完整对白 Key', task['records'][2]['warning'])
        self.assertEqual(task['asset_paths'], {})

    def test_customer_note_keys_survive_pipeline_and_group_dialogue_without_emotions(self):
        prefix = 'mdnew_forestconcert_mdnew_decoration_'
        book = Workbook()
        sheet = book.active
        sheet.append(['note', 'limit', 'key', 'game_text_en', 'game_text_de'])
        for row in [
                [prefix + '3_10', None, 'Max', 'Tenth line', 'KEEP_TARGET'],
                [prefix + '3_2', None, 'Christine', 'Second line', None],
                [None, None, None, None, None],
                [prefix + '5_1', None, 'Max', 'Other scene first', None],
                [prefix + '5_2', None, 'Max', 'Other scene second', None],
                [None, 30, prefix + 'name_3', 'Scene three prop', None]]:
            sheet.append(row)
        original = [[cell.value for cell in row] for row in sheet]
        path = self.root / 'customer.xlsx'
        book.save(path)
        book.close()
        before = path.read_bytes()
        ctx = {'game': {'speakers': ['Max', 'Christine'], 'key_rules': 'profiles/game/wc_key_rules.json'}}
        normalized, classified = self.root / 'normalized.xlsx', self.root / 'classified.xlsx'
        for runner, source, target in [(run_normalize, path, normalized), (run_classify, normalized, classified)]:
            task = runner(ctx, source, target)
            self.assertEqual(task['state'], 'completed', task.get('status'))
        task = self.run_context(classified, enhance=False)
        records = task['records']
        self.assertEqual([record['key'] for record in records[:2]], [prefix + '3_10', prefix + '3_2'])
        self.assertEqual([record['speaker'] for record in records[:2]], ['Max', 'Christine'])
        context = records[0]['context']
        self.assertLess(context.index('Second line'), context.index('Tenth line'))
        self.assertIn('Scene three prop', context)
        self.assertNotIn('Other scene', context)
        self.assertEqual(context, records[1]['context'])
        self.assertNotIn('Tenth line', records[2]['context'])
        self.assertIn('Tenth line', records[-1]['context'])
        result = load_workbook(task['result_file'])
        self.assertEqual([[cell.value for cell in row[:5]] for row in result.active], original)
        result.close()
        self.assertEqual(path.read_bytes(), before)

    def test_note_identity_requires_full_dialogue_key_and_preserves_existing_key(self):
        prefix = 'mdnew_forestconcert_mdnew_decoration_'
        for note_header in ['note', '备注']:
            with self.subTest(note_header=note_header):
                book = Workbook()
                book.active.append(['key', 'speaker', note_header, 'game_text_en'])
                for row in [
                        ['Max', None, prefix + 'intro_1', 'Without normalized speaker'],
                        ['Max', 'Max', 'ordinary note', 'No guessed scene'],
                        [prefix + '3_1', 'Max', prefix + '5_1', 'Existing key wins'],
                        ['cat_multiboards_title', None, prefix + '5_1', 'Keep UI key'],
                        [None, 'Max', prefix + '5_1', 'Missing key'],
                        ['Max', 'Max', prefix + 'name_3', 'Decoration is not dialogue']]:
                    book.active.append(row)
                records = self.processor.read_records(book, 'customer.xlsx')
                self.assertEqual([record['key'] for record in records], [
                    prefix + 'intro_1', 'Max', prefix + '3_1', 'cat_multiboards_title', prefix + '5_1', 'Max'])
                self.assertEqual(records[0]['speaker'], 'Max')
                book.close()

    def test_ui_context_contains_only_related_source_and_no_processing_metadata(self):
        path = self.main([
            [None, 'cat_multiboards_pre_desc', 'CN_SENTINEL', 'Event intro', None, CATEGORIES[0]],
            [None, 'cat_multiboards_guide_1_1', '', 'Merge for keys', None, CATEGORIES[0]],
            [30, 'cat_multiboards_complete_title', '', 'Complete title', None, CATEGORIES[0]],
            [None, 'cat_multiboards_complete_desc', '', 'Complete description', None, CATEGORIES[0]],
            [30, 'cat_multiboards_fail_title', '', 'Failure title', None, CATEGORIES[0]],
            [None, 'cat_multiboards_fail_desc', '', 'Failure description', None, CATEGORIES[0]],
            [30, 'catMultiboards_book_title', '', 'Collection title', None, CATEGORIES[0]],
            [30, 'cat_multiboards_book_subtitle1', '', 'Collection section', None, CATEGORIES[0]],
        ])
        records = self.run_context(path, enhance=False)['records']
        self.assertEqual(records[0]['context'], '[关联语境参考]\ncat_multiboards_pre_desc: Event intro\ncat_multiboards_guide_1_1: Merge for keys')
        self.assertEqual(records[0]['context'], records[1]['context'])
        self.assertEqual(records[2]['context'], records[3]['context'])
        self.assertIn('Complete description', records[2]['context'])
        self.assertNotIn('Failure', records[2]['context'])
        self.assertNotIn('Collection', records[2]['context'])
        self.assertEqual(records[6]['context'], records[7]['context'])
        for record in records:
            for excluded in ['[关联状态]', '[源位置]', '[分类]', '[参考来源]', '[辅助中文]', 'CN_SENTINEL', '← 当前句', 'classified.xlsx']:
                self.assertNotIn(excluded, record['context'])
            self.assertIn('classified.xlsx', record['source'])

    def test_rank_entry_and_results_do_not_mix_states(self):
        path = self.main([[None, key, '', english, None, CATEGORIES[0]] for key, english in [
            ('forestconcert_rank_title', 'Sprint'),
            ('forestconcert_rank_score_need', 'Collect {0} tokens'),
            ('forestconcert_rank_start_desc', 'Join the race'),
            ('forestconcert_rank_final_desc', 'You placed <b>{0}</b>'),
        ]])
        records = self.run_context(path, enhance=False)['records']
        self.assertIn('Join the race', records[1]['context'])
        self.assertNotIn('You placed', records[1]['context'])
        self.assertIn('You placed <b>{0}</b>', records[3]['context'])
        self.assertNotIn('Join the race', records[3]['context'])
        self.assertEqual(records[3]['related_count'], 2)

    def test_done_materials_and_external_or_cache_paths_are_rejected(self):
        for name in ['source_done.xlsx', 'SOURCE_DONE.xlsx', r'C:\temp\source_done (1).xlsx']:
            with self.assertRaises(ValueError):
                validate_material_name(name)
        for name in ['../outside.xlsx', 'WC_Tools/uploads/cache.xlsx', '/another/project/source.xlsx']:
            with self.assertRaises(ValueError):
                self.processor.resolve_material(name)
        path = self.main([[30, 'cat_multiboards_pre_desc', '', 'Text', None, CATEGORIES[0]]])
        task = {'state': 'processing'}
        with patch('core_tools.context_processor.load_workbook') as loader:
            self.processor.process(str(path), str(self.root / 'blocked.xlsx'), {'test': task}, 'test',
                                   reference_path=self.root / 'source_done.xlsx', enhance=True)
            loader.assert_not_called()
        self.assertEqual(task['state'], 'error')
        self.assertIn('done', task['status'])



if __name__ == '__main__':
    unittest.main()
