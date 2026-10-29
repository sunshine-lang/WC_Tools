import unittest
from tempfile import TemporaryDirectory
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core_tools.normalizer import Normalizer


class NormalizerTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.dir = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def build(self, rows, headers=('limit', 'key', 'game_text_en', 'game_text_de')):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(list(headers))
        for row in rows:
            sheet.append(list(row))
        return workbook

    def save(self, workbook, name='in.xlsx'):
        path = self.dir / name
        workbook.save(path)
        return path

    def run_normalizer(self, source, **kwargs):
        out = self.dir / 'out.xlsx'
        tasks, task = {'t': {}}, None
        tasks['t'] = {}
        Normalizer(**kwargs).process(str(source), str(out), tasks, 't')
        self.assertEqual(tasks['t'].get('state'), 'completed', tasks['t'].get('status'))
        return load_workbook(out).worksheets[0]

    def test_trims_trailing_empty_columns(self):
        workbook = self.build([(30, 'cat_title', 'Feline Frenzy!', None)])
        sheet = workbook.active
        for column in range(6, 20):          # 模拟导出残留的空列
            sheet.cell(1, column).number_format = '@'
        source = self.save(workbook)

        result = self.run_normalizer(source)
        headers = [c.value for c in result[1] if c.value is not None]
        self.assertEqual(headers, ['limit', 'key', 'game_text_en', 'game_text_de'])

    def test_adds_speaker_for_known_names(self):
        workbook = self.build([
            (None, 'Chad', 'Being rich is the main thing.', None),
            (30, 'cat_title', 'Feline Frenzy!', None),
        ])
        source = self.save(workbook)

        result = self.run_normalizer(source, speakers=['Chad', 'Max'])
        headers = {c.value: c.column for c in result[1] if c.value is not None}
        self.assertIn('speaker', headers)
        self.assertEqual(result.cell(2, headers['speaker']).value, 'Chad')
        self.assertIsNone(result.cell(3, headers['speaker']).value)

    def test_keeps_existing_speaker_column(self):
        workbook = self.build(
            [(None, 'Chad', 'line', None, 'Chad')],
            headers=('limit', 'key', 'game_text_en', 'game_text_de', 'speaker'),
        )
        source = self.save(workbook)

        result = self.run_normalizer(source, speakers=['Chad'])
        headers = {c.value: c.column for c in result[1] if c.value is not None}
        self.assertEqual(headers['speaker'], 5)
        self.assertEqual(result.cell(2, 5).value, 'Chad')

    def test_without_speakers_leaves_sheet_alone(self):
        workbook = self.build([(None, 'Chad', 'line', None)])
        source = self.save(workbook)

        result = self.run_normalizer(source)
        headers = [c.value for c in result[1] if c.value is not None]
        self.assertNotIn('speaker', headers)

    def test_non_dialogue_does_not_add_empty_speaker_with_configured_names(self):
        source = self.save(self.build([(30, 'cat_multiboards_title', 'Cats', 'Katzen'),
                                       (30, 'Ahalloween_set1_name', 'Autumn', 'Herbst')]))
        result = self.run_normalizer(source, speakers=['Chad', 'Max'])
        self.assertEqual([cell.value for cell in result[1]],
                         ['limit', 'key', 'game_text_en', 'game_text_de'])

    def test_existing_speaker_alias_is_used_without_adding_another_column(self):
        source = self.save(self.build([(None, 'Chad', 'Hello', None, 'Provided Name')],
            headers=('limit', 'key', 'game_text_en', 'game_text_de', '说话人')))
        result = self.run_normalizer(source, speakers=['Chad'])
        self.assertEqual(result.max_column, 5)
        self.assertEqual(result['E2'].value, 'Provided Name')

    def test_existing_blank_speaker_column_is_preserved_for_non_dialogue(self):
        source = self.save(self.build([(30, 'cat_multiboards_title', 'Cats', None, None)],
            headers=('limit', 'key', 'game_text_en', 'game_text_de', 'speaker')))
        result = self.run_normalizer(source, speakers=['Chad'])
        self.assertEqual(result['E1'].value, 'speaker')
        self.assertIsNone(result['E2'].value)


if __name__ == '__main__':
    unittest.main()
