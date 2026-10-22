import json
import tempfile
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
from unittest.mock import Mock, patch

import numpy as np
import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core_tools.tm_matcher import TMMatcher
from core_tools.tm_semantic import SemanticTMRetriever
from pipeline import run_tm


class SemanticTMTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.config = self.root / 'models.json'
        self.settings = {
            'embedding': {'api_url': 'https://example.invalid/embed', 'api_key': 'TEST_SECRET',
                          'model_name': 'embedding-test', 'batch_size': 2},
            'rerank': {'api_url': 'https://example.invalid/rerank', 'api_key': 'TEST_SECRET',
                       'model_name': 'rerank-test', 'recall_k': 30, 'final_k': 5,
                       'threshold': 0.65, 'request_retries': 3, 'retry_backoff': 0.1},
        }
        self.config.write_text(json.dumps(self.settings))
        self.cache = self.root / 'cache'
        self.client = SemanticTMRetriever(self.config, self.cache)
        self.guard = patch('requests.sessions.Session.request', side_effect=AssertionError('NETWORK_DISABLED'))
        self.guard.start()

    def tearDown(self):
        self.guard.stop()
        self.temp.cleanup()

    @staticmethod
    def response(section, payload, timeout, label):
        if 'input' in payload and isinstance(payload['input'], list):
            return {'data': [{'index': i, 'embedding': [1.0, 0.1] if any(
                word in text.lower() for word in ('cat', 'feline')) else [0.1, 1.0]}
                for i, text in reversed(list(enumerate(payload['input'])))]}
        return {'results': [{'index': i, 'relevance_score': 0.96 if 'feline' in text else 0.1}
                            for i, text in enumerate(payload['documents'])]}

    def save_tm(self, pairs):
        root = ET.Element('tmx', version='1.4')
        body = ET.SubElement(root, 'body')
        for en, de in pairs:
            unit = ET.SubElement(body, 'tu')
            for lang, value in [('en-US', en), ('de-DE', de)]:
                ET.SubElement(ET.SubElement(unit, 'tuv', {'lang': lang}), 'seg').text = value
        path = self.root / 'customer.tmx'
        ET.ElementTree(root).write(path, encoding='utf-8')
        return path

    def matcher(self, pairs):
        return TMMatcher(self.save_tm(pairs), semantic=True, config_path=self.config, cache_dir=self.cache)

    def source(self):
        book = Workbook()
        book.active.append(['game_text_en', 'game_text_de', '分类', 'Context', '生效术语', '无效术语'])
        book.active.append(['The cat naps.', 'KEEP_TARGET', '活动对白', 'CONTEXT_SENTINEL', 'TERM_SENTINEL', ''])
        book.active['A2'].font = Font(bold=True)
        path = self.root / 'input.xlsx'
        book.save(path)
        book.close()
        return path

    def test_semantic_synonym_with_no_shared_words_is_retrieved(self):
        matcher = self.matcher([('Feline slumber', 'Katzenschlaf'), ('Buy a train ticket', 'Kaufe ein Zugticket')])
        with patch.object(SemanticTMRetriever, '_post', side_effect=self.response) as post:
            matches = matcher.find_matches('The cat naps.')
        self.assertEqual([(m['english'], m['match_type']) for m in matches], [('Feline slumber', '语义相似')])
        self.assertEqual(post.call_count, 3)  # library, query, rerank

    def test_exact_and_normalized_matches_bypass_models_and_keep_variants(self):
        matcher = self.matcher([('Tap', 'Tippen'), ('Tap', 'Antippen')])
        with patch.object(SemanticTMRetriever, '_post', side_effect=AssertionError('no API for exact matches')):
            matches = matcher.find_matches('Tap')
            normalized = matcher.find_matches('<i>TAP</i>')
        self.assertEqual(len(matches), 2)
        self.assertTrue(all(m['match_type'] == '原文一致' for m in matches))
        self.assertTrue(all(m['warnings'] for m in normalized))

    def test_cache_reuses_only_matching_tm_source_order_and_model(self):
        with patch.object(SemanticTMRetriever, '_post', side_effect=self.response) as post:
            self.client.prepare(['one', 'two'], 'hash-a')
            initial = post.call_count
            another = SemanticTMRetriever(self.config, self.cache)
            another.prepare(['one', 'two'], 'hash-a')
            self.assertTrue(another.cache_hit)
            self.assertEqual(post.call_count, initial)
            another.prepare(['two', 'one'], 'hash-a')
            self.assertFalse(another.cache_hit)
            another.prepare(['two', 'one'], 'hash-b')
            another.embedding = {**another.embedding, 'model_name': 'new-model'}
            another.prepare(['two', 'one'], 'hash-b')
            self.assertEqual(post.call_count, initial + 3)
        self.assertEqual(len(list(self.cache.glob('*.npz'))), 4)
        for path in self.cache.glob('*.npz'):
            with np.load(path, allow_pickle=False) as cache:
                self.assertNotIn('TEST_SECRET', cache['manifest'].item())

    def test_embedding_indices_are_honored_and_zero_or_missing_vectors_fail(self):
        with patch.object(self.client, '_post', side_effect=self.response):
            vectors = self.client.embed(['cat', 'train'])
        np.testing.assert_allclose(vectors, [[1, .1], [.1, 1]])
        for data in [[], [{'index': 0, 'embedding': [0., 0.]}],
                     [{'index': 0, 'embedding': [float('nan'), 1.]}],
                     [{'index': True, 'embedding': [1., 1.]}]]:
            with self.subTest(data=data), patch.object(self.client, '_post', return_value={'data': data}):
                with self.assertRaisesRegex(ValueError, '向量响应'):
                    self.client.embed(['cat'])

    def test_failed_embedding_does_not_create_a_cache(self):
        with patch.object(self.client, '_post', side_effect=ValueError('模拟接口失败')):
            with self.assertRaises(ValueError):
                self.client.prepare(['cat'], 'hash')
        self.assertFalse(list(self.cache.glob('*.npz')))

    def test_rerank_incomplete_duplicate_invalid_scores_fail_not_miss(self):
        invalid = [[], [{'index': 0, 'relevance_score': .9}],
                   [{'index': 0, 'relevance_score': .9}, {'index': 0, 'relevance_score': .8}],
                   [{'index': 0, 'relevance_score': .9}, {'index': 2, 'relevance_score': .8}],
                   [{'index': 0, 'relevance_score': .9}, {'index': 1, 'relevance_score': float('inf')}]]
        for items in invalid:
            with self.subTest(items=items), patch.object(self.client, '_post', return_value={'results': items}):
                with self.assertRaisesRegex(ValueError, '未当作'):
                    self.client.rerank('query', ['a', 'b'])

    def test_transport_retry_and_auth_failure_do_not_leak_credentials(self):
        for response in [requests.Timeout('TEST_SECRET'), Mock(status_code=503)]:
            with self.subTest(response=response), patch('requests.Session.post', side_effect=
                    response if isinstance(response, Exception) else None, return_value=response) as post, \
                    patch('core_tools.tm_semantic.time.sleep'):
                with self.assertRaisesRegex(ValueError, '3 次') as caught:
                    self.client._post(self.client.rerank_config, {}, 1, '重排模型')
                self.assertEqual(post.call_count, 3)
                self.assertNotIn('TEST_SECRET', str(caught.exception))
        with patch('requests.Session.post', return_value=Mock(status_code=401)) as post:
            with self.assertRaisesRegex(ValueError, '401'):
                self.client._post(self.client.rerank_config, {}, 1, '重排模型')
            self.assertEqual(post.call_count, 1)

    def test_recall_then_threshold_and_final_k(self):
        sources = [f'candidate {i}' for i in range(8)]
        with patch.object(self.client, '_post', side_effect=self.response):
            self.client.prepare(sources, 'hash')
        self.client.recall_k = 6
        self.client.final_k = 3
        with patch.object(self.client, 'rerank', return_value=[.9, .2, .8, .3, .7, .6]) as rerank:
            result = self.client._search('query', np.array([.1, 1.]))
        self.assertEqual(len(rerank.call_args.args[1]), 6)
        self.assertEqual([i for i, _, _ in result], [0, 2, 4])

    def test_semantic_result_preserves_tags_parameters_and_difference_warnings(self):
        matcher = self.matcher([('Collect <b>{0}</b> feline tokens!', 'Sammle <b>{0}</b> Katzenmarken!')])
        with patch.object(SemanticTMRetriever, '_post', side_effect=self.response):
            match = matcher.find_matches('Gather {1} cat coins!')[0]
        self.assertIn('<b>{0}</b>', match['english'])
        self.assertIn('<b>{0}</b>', match['german'])
        self.assertTrue(any('参数不同' in w for w in match['warnings']))

    def test_pipeline_enables_semantics_and_only_appends_reference_column(self):
        tm = self.save_tm([('Feline slumber', 'Katzenschlaf')])
        source = self.source()
        output = self.root / 'output.xlsx'
        before_bytes = source.read_bytes()
        with patch.object(SemanticTMRetriever, '_post', side_effect=self.response) as post:
            result = run_tm({'language': {'tms': [str(tm)]}, 'config_path': self.config,
                             'tm_cache_dir': self.cache}, source, output)
        self.assertEqual(result['state'], 'completed', result)
        self.assertEqual(result['counts']['相似参考'], 1)
        self.assertEqual(result['tm_policy'], 'wc-en-de-semantic-v1')
        self.assertEqual(source.read_bytes(), before_bytes)
        before, after = load_workbook(source), load_workbook(output)
        try:
            self.assertEqual([c.value for c in after.active[1]], [c.value for c in before.active[1]] + ['TM参考'])
            for row in before.active:
                for old in row:
                    new = after.active.cell(old.row, old.column)
                    self.assertEqual((old.value, old.data_type, old._style), (new.value, new.data_type, new._style))
            self.assertIn('Katzenschlaf', after.active['G2'].value)
        finally:
            before.close()
            after.close()
        payloads = json.dumps([c.args[1] for c in post.call_args_list])
        for sentinel in ['KEEP_TARGET', 'CONTEXT_SENTINEL', 'TERM_SENTINEL']:
            self.assertNotIn(sentinel, payloads)

    def test_model_failure_stops_workbook_instead_of_publishing_empty_references(self):
        matcher = self.matcher([('Feline slumber', 'Katzenschlaf')])
        task = {}
        output = self.root / 'result.xlsx'
        with patch.object(SemanticTMRetriever, '_post', side_effect=ValueError('TM 重排接口失败')):
            matcher.process(self.source(), output, {'t': task}, 't')
        self.assertEqual(task['state'], 'error')
        self.assertFalse(output.exists())
        self.assertNotIn('result_file', task)

    def test_valid_low_scores_are_a_real_miss(self):
        matcher = self.matcher([('Buy a train ticket', 'Kaufe ein Zugticket')])
        with patch.object(SemanticTMRetriever, '_post', side_effect=self.response):
            self.assertEqual(matcher.find_matches('The cat naps.'), [])

    def test_dimension_mismatch_is_an_error(self):
        with patch.object(self.client, '_post', side_effect=self.response):
            self.client.prepare(['cat'], 'hash')
        with self.assertRaisesRegex(ValueError, '维度不一致'):
            self.client._search('query', np.array([1., 2., 3.]))

    def test_tag_only_entries_are_not_embedded_and_indices_still_map_to_original(self):
        matcher = self.matcher([('<sprite name="coin">', '<sprite name="coin">'),
                                ('Feline slumber', 'Katzenschlaf')])
        with patch.object(SemanticTMRetriever, '_post', side_effect=self.response) as post:
            matches = matcher.find_matches('The cat naps.')
        self.assertEqual(matches[0]['german'], 'Katzenschlaf')
        for call in post.call_args_list:
            if isinstance(call.args[1].get('input'), list):
                self.assertNotIn('', call.args[1]['input'])

    def test_corrupt_cache_is_rebuilt(self):
        with patch.object(self.client, '_post', side_effect=self.response):
            self.client.prepare(['cat'], 'hash')
        next(self.cache.glob('*.npz')).write_bytes(b'broken cache')
        replacement = SemanticTMRetriever(self.config, self.cache)
        with patch.object(replacement, '_post', side_effect=self.response) as post:
            replacement.prepare(['cat'], 'hash')
        self.assertEqual(post.call_count, 1)
        self.assertFalse(replacement.cache_hit)


if __name__ == '__main__':
    unittest.main()
