"""交付前的准备自检：只读产物、不改数据，输出一份可追溯的 Markdown 报告。

检查的是“准备阶段产物的完整性”，不是译后质量：
行数、列、源文、标签、TM 引用、长度余量和覆盖率。
"""
from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .text_metrics import PARAM_PATTERN, TAG_PATTERN, display_length


PREFLIGHT_POLICY = 'wc-prep-preflight-v1'
NEAR_LIMIT_RATIO = 0.85
BLOCK = '阻断'
WARN = '告警'

CHECK_ORDER = (
    ('rows', '行数守恒', BLOCK),
    ('columns', '列可逆', BLOCK),
    ('target', '目标列未写入', BLOCK),
    ('source', '源文未改写', BLOCK),
    ('tags', '标签序列一致', BLOCK),
    ('slots', '无重复槽位', BLOCK),
    ('refs', '无悬空引用', BLOCK),
    ('length', '长度余量', WARN),
    ('coverage', '覆盖率', WARN),
)


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_sheet(path: Path) -> tuple[list[str], list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [text(cell).strip() for cell in next(rows, ())]
        body = []
        for values in rows:
            if not any(value not in (None, '') for value in values):
                continue
            body.append({header[i]: (values[i] if i < len(values) else None)
                         for i in range(len(header)) if header[i]})
        return header, body
    finally:
        workbook.close()


def get(row: dict, *names, default=''):
    for name in names:
        if name in row and row[name] is not None:
            return text(row[name])
    for key, value in row.items():
        if key.casefold() in {n.casefold() for n in names} and value is not None:
            return text(value)
    return default


def tags(value: str) -> Counter:
    found = TAG_PATTERN.findall(value) + PARAM_PATTERN.findall(value)
    return Counter(found)


def tm_blocks(value: str) -> list[dict]:
    """解析 `[TM参考 n｜类型]` 块，返回每块的英文与德文。"""
    blocks = []
    header_pattern = re.compile(r'\[TM参考\s+(\d+)\s*[｜|]\s*([^\]]+)\]')
    matches = list(header_pattern.finditer(value))
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(value)
        body = value[start:end]
        english = ''
        german = ''
        english_match = re.search(r'【英】\s*(.*?)(?=\n【德】|\Z)', body, re.S)
        german_match = re.search(r'【德】\s*(.*?)(?=\n\[TM参考|\Z)', body, re.S)
        if english_match:
            english = english_match.group(1).strip()
        if german_match:
            german = german_match.group(1).strip()
        blocks.append({'number': int(match.group(1)), 'type': match.group(2).strip(),
                       'english': english, 'german': german})
    return blocks


def check_row(ok: bool, name: str, detail: str, gate: str) -> dict:
    return {'key': name, 'ok': bool(ok), 'name': name, 'detail': detail, 'gate': gate}


def preflight(source_path: Path, final_path: Path, batch_name: str = '') -> dict:
    """执行九条断言，返回 checks / stats / preview / report 文本。"""
    source_header, source_rows = read_sheet(Path(source_path))
    final_header, final_rows = read_sheet(Path(final_path))

    checks = []
    details = {}

    # 1. 行数守恒
    rows_ok = len(source_rows) == len(final_rows)
    checks.append(check_row(rows_ok, '行数守恒', f'{len(source_rows)} → {len(final_rows)}', BLOCK))
    details['rows'] = (len(source_rows), len(final_rows))

    # 2. 列可逆：原始列按原顺序保留，新增列只能在右侧
    prefix_ok = len(final_header) >= len(source_header) and final_header[:len(source_header)] == source_header
    checks.append(check_row(prefix_ok, '列可逆',
                            f'{len(source_header)} 列保留，新增 {max(0, len(final_header) - len(source_header))} 列在右侧',
                            BLOCK))
    details['columns'] = (source_header, final_header)

    # 3. 目标列未写入（准备阶段不应替人翻译）
    target_rows = [get(row, 'game_text_de') for row in final_rows]
    target_ok = bool(target_rows) and all(not value.strip() for value in target_rows)
    checks.append(check_row(target_ok, '目标列未写入',
                            'game_text_de 逐行为空' if target_ok else '发现非空译文，准备阶段不应写目标列',
                            BLOCK))

    # 4. 源文未改写
    source_pairs = [(get(row, 'key'), get(row, 'game_text_en')) for row in source_rows]
    final_pairs = [(get(row, 'key'), get(row, 'game_text_en')) for row in final_rows]
    source_ok = source_pairs == final_pairs
    changed = sum(1 for a, b in zip(source_pairs, final_pairs) if a != b)
    checks.append(check_row(source_ok, '源文未改写',
                            'key 与 game_text_en 未变化' if source_ok else f'{changed} 行源文或 key 发生变化',
                            BLOCK))

    # 5. 标签序列一致
    tag_bad = 0
    for before, after in zip(source_rows, final_rows):
        if tags(get(before, 'game_text_en')) != tags(get(after, 'game_text_en')):
            tag_bad += 1
    checks.append(check_row(tag_bad == 0, '标签序列一致',
                            '标签与占位符计数一致' if tag_bad == 0 else f'{tag_bad} 行标签或占位符变化',
                            BLOCK))

    # 6. 无重复槽位
    duplicate = 0
    for row in final_rows:
        blocks = tm_blocks(get(row, 'TM参考'))
        if len({(b['type'], b['english']) for b in blocks}) != len(blocks):
            duplicate += 1
    checks.append(check_row(duplicate == 0, '无重复槽位',
                            'TM参考 单元格无重复块' if duplicate == 0 else f'{duplicate} 行出现重复 TM 块',
                            BLOCK))

    # 7. 无悬空引用（每个块有英德；原文一致的英文必须等于当前源文）
    dangling = 0
    exact = 0
    refs = 0
    for row in final_rows:
        current = get(row, 'game_text_en')
        for block in tm_blocks(get(row, 'TM参考')):
            refs += 1
            if not block['english'] or not block['german']:
                dangling += 1
            if block['type'] == '原文一致':
                exact += 1
                if block['english'] != current:
                    dangling += 1
    checks.append(check_row(dangling == 0, '无悬空引用',
                            f'{refs} 条参考均可追溯，其中原文一致 {exact} 条' if dangling == 0
                            else f'{dangling} 条参考缺英/德文或原文对不上',
                            BLOCK))

    # 8. 长度余量：源文已达到 limit 的 85%，交付前最好人工确认
    near = []
    for row in final_rows:
        raw_limit = get(row, 'limit')
        if not raw_limit:
            continue
        try:
            limit = float(raw_limit)
        except ValueError:
            continue
        if limit <= 0:
            continue
        length = display_length(get(row, 'game_text_en'))
        if length / limit >= NEAR_LIMIT_RATIO:
            near.append({'row': get(row, 'key') or '—', 'length': length, 'limit': limit,
                         'source': get(row, 'game_text_en')})
    checks.append(check_row(not near, '长度余量',
                            f'{len(near)} 条接近 limit（≥{int(NEAR_LIMIT_RATIO * 100)}%），需人工确认'
                            if near else '没有接近 limit 的源文',
                            WARN))
    details['near'] = near

    # 9. 覆盖率：每条记录都至少完成分类
    covered = sum(1 for row in final_rows if get(row, '分类').strip())
    coverage_ok = covered == len(final_rows) and bool(final_rows)
    checks.append(check_row(coverage_ok, '覆盖率',
                            f'{covered} / {len(final_rows)} 有分类产出', WARN))

    passed = sum(1 for check in checks if check['ok'])
    warnings = sum(1 for check in checks if not check['ok'] and check['gate'] == WARN)
    blocked = sum(1 for check in checks if not check['ok'] and check['gate'] == BLOCK)
    stats = [['断言', len(checks), ' 条'], ['通过', passed, ''], ['告警', warnings, ''], ['阻断', blocked, '']]

    preview_rows = []
    for check in checks:
        level = 'ok' if check['ok'] else ('warn' if check['gate'] == WARN else 'bad')
        preview_rows.append([level, check['name'], check['detail'],
                             '告警' if check['gate'] == WARN else '阻断'])
    preview = {'type': 'checks', 'cap': '断言结果', 'rows': preview_rows}

    review = [
        f'阻断项 {blocked}，' + ('可以交付' if blocked == 0 else '不要交付，先回头修上游'),
        f'长度告警 {len(near)} 条，需人工确认是否改写',
        f'覆盖率 {covered} / {len(final_rows)}',
    ]

    report = render_report(batch_name, checks, details, len(final_rows))
    return {'checks': checks, 'review': review, 'stats': stats, 'preview': preview,
            'near': near, 'blocked': blocked, 'warnings': warnings,
            'passed': passed, 'report': report, 'rows': len(final_rows),
            'exact': exact, 'refs': refs}


def render_report(batch_name: str, checks: list[dict], details: dict, rows: int) -> str:
    lines = [
        '# 准备自检报告',
        '',
        f'- 批次：{batch_name or "未命名"}',
        f'- 时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}',
        f'- 规则：{PREFLIGHT_POLICY}',
        f'- 记录数：{rows}',
        '',
        '## 断言结果',
        '',
        '| 结果 | 检查项 | 说明 | 级别 |',
        '| --- | --- | --- | --- |',
    ]
    for check in checks:
        mark = '通过' if check['ok'] else ('告警' if check['gate'] == WARN else '阻断')
        lines.append(f"| {mark} | {check['name']} | {check['detail']} | {check['gate']} |")

    near = details.get('near') or []
    lines += ['', '## 长度余量明细', '']
    if near:
        lines += ['| Key | 源文长度 | limit | 源文 |', '| --- | ---: | ---: | --- |']
        for item in near:
            source = item['source'].replace('|', '\\|')
            lines.append(f"| {item['row']} | {item['length']} | {int(item['limit'])} | {source} |")
    else:
        lines.append('没有接近 limit 的源文。')

    lines += ['', '## 说明', '',
              '本报告只检查准备阶段产物的完整性，不评价译文质量。',
              f'长度规则：源文达到 limit 的 {int(NEAR_LIMIT_RATIO * 100)}% 即列为告警。',
              '阻断项失败时不应进入交付；告警项只记录，由人工确认。', '']
    return '\n'.join(lines)


def process(source_path, final_path, task_dict, task_id, batch_name=''):
    """模块适配：与其余 core_tools 一样接收 task_dict/task_id。"""
    task = task_dict[task_id]
    try:
        task.update(progress=15, status='正在读取规范化与 TM 产物...')
        source_path = Path(source_path)
        final_path = Path(final_path)
        if not source_path.is_file():
            raise ValueError('找不到规范化产物，无法自检')
        if not final_path.is_file():
            raise ValueError('找不到 TM 产物，无法自检')
        task.update(progress=45, status='正在核对行数、列、源文与标签...')
        result = preflight(source_path, final_path, batch_name)
        task.update(progress=85, status='正在生成自检报告...')
        final_path = Path(final_path)
        report_path = final_path.with_name('06-自检报告.md')
        report_path.write_text(result['report'], encoding='utf-8')
        blocked = result['blocked']
        task.update(
            progress=100,
            state='error' if blocked else 'completed',
            result_file=str(report_path),
            checks=result['checks'],
            review=result['review'],
            stats=result['stats'],
            preview=result['preview'],
            near_limit=result['near'],
            counts={'passed': result['passed'], 'warnings': result['warnings'], 'blocked': blocked},
            preflight_policy=PREFLIGHT_POLICY,
            status=(f'自检发现 {blocked} 个阻断项，报告已生成' if blocked
                    else f'自检通过：{result["passed"]} / {len(result["checks"])} 条，'
                         f'{result["warnings"]} 条告警'),
        )
    except Exception as error:  # noqa: BLE001
        task.update(state='error', status=f'准备自检失败：{error}')
