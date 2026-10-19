"""译后质量检查：只读取外部翻译回填后的表格，生成独立问题报告，不修改交付文件。

规则与游戏、语种无关，源列／译文列等全部由调用方指定，可用于 WC 英德、RO 中英等项目。
"""
import csv
import hashlib
import os
import re
import uuid
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


QA_POLICY = 'qa-generic-v1'
CHECKS = ('行对应', '漏译', '源文异常', '参数', '标签', '标签顺序', '换行', '数字',
          '字数限制', '长度比例', '未翻译', '首尾空格', '术语', '译法不一致')
# 常见占位符：{0}、{name}、%s、%1$s、%(name)s。
# %i、%f 要求后面不紧跟字母，否则德语 100%ige、法语等写法会被误判成参数。
PARAM_PATTERN = re.compile(r'\{\w*\}|%\([^)]+\)[sdif]|%(?:\d+\$)?(?:[sd]|[if](?![^\W\d_]))')
TAG_PATTERN = re.compile(r'</?[A-Za-z][^<>]*>')
SPRITE_PATTERN = re.compile(r'<sprite\b', re.IGNORECASE)
# 千分位、小数点写法因语种而异（1,000 / 1.000 / 1 000），比较时只看数字本身。
NUMBER_PATTERN = re.compile(r"\d+(?:[.,\u00a0\u202f']\d+)*")
TERM_ENTRY = re.compile(r'^(.*)\(([^()]*)\)\s*$', re.S)
MAX_SHIFT_REPORTS = 20
LENGTH_RULE = '长度按去除 <…> 标签和换行符后的字符数计算，空格和参数按原样计入'


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def display_length(value):
    return len(re.sub(r'\\n|\n', '', TAG_PATTERN.sub('', value)))


def numbers(value):
    plain = PARAM_PATTERN.sub(' ', TAG_PATTERN.sub(' ', value))
    return Counter(re.sub(r"[.,\u00a0\u202f']", '', n) for n in NUMBER_PATTERN.findall(plain))


def line_breaks(value):
    return value.count('\\n'), value.count('\n')


def describe(source, target):
    missing, extra = source - target, target - source
    parts = []
    if missing:
        parts.append('缺少 ' + '、'.join(sorted(missing.elements())))
    if extra:
        parts.append('多出 ' + '、'.join(sorted(extra.elements())))
    return '；'.join(parts)
def parse_terms(value):
    """解析第三步的 `原词 (译法1 / 译法2) | …` 格式；无法解析的片段跳过。"""
    terms = []
    for entry in value.split(' | '):
        match = TERM_ENTRY.match(entry.strip())
        if match and match.group(1).strip():
            variants = [v.strip() for v in match.group(2).split(' / ') if v.strip()]
            if variants:
                terms.append((match.group(1).strip(), variants))
    return terms


def read_sheets(path, display_name):
    """返回 [(表名, 表头行, [(行号, [(文本, 是否公式), …]), …])]。"""
    path = Path(path)
    if path.suffix.lower() == '.csv':
        try:
            with open(path, encoding='utf-8-sig', newline='') as file:
                rows = list(csv.reader(file, strict=True))
        except UnicodeDecodeError as error:
            raise ValueError(f'{display_name} 须为 UTF-8 编码的 CSV') from error
        if not rows:
            raise ValueError(f'{display_name} 为空')
        body = [(number, [(cell, False) for cell in row]) for number, row in enumerate(rows[1:], 2)]
        return [(Path(display_name).stem, rows[0], body)]
    workbook = load_workbook(path, data_only=False)
    try:
        sheets = []
        for sheet in workbook:
            rows = list(sheet.iter_rows())
            if not rows:
                continue
            header = [text(cell.value) for cell in rows[0]]
            body = [(row[0].row, [(text(cell.value), cell.data_type == 'f') for cell in row]) for row in rows[1:]]
            sheets.append((sheet.title, header, body))
        return sheets
    finally:
        workbook.close()


def find_column(header, name, sheet, required=True):
    wanted = name.strip().casefold()
    matches = [i for i, value in enumerate(header) if value.strip().casefold() == wanted]
    if len(matches) > 1:
        raise ValueError(f'“{sheet}”有多个“{name}”列，请保留一列')
    if not matches and required:
        raise ValueError(f'“{sheet}”缺少“{name}”列')
    return matches[0] if matches else None


def read_records(path, display_name, columns, require_target=True):
    records, used_sheets = [], []
    for sheet, header, body in read_sheets(path, display_name):
        source_col = find_column(header, columns['source'], sheet, required=False)
        if source_col is None:
            continue
        target_col = find_column(header, columns['target'], sheet, required=require_target)
        optional = {name: find_column(header, columns[name], sheet, required=False)
                    for name in ('key', 'limit') if columns.get(name)}
        # 优先使用降噪后的“生效术语”，否则用第三步的候选“关联术语”。
        term_col = None
        for name in columns.get('terms', ()):
            term_col = find_column(header, name, sheet, required=False)
            if term_col is not None:
                break
        used_sheets.append(sheet)

        def cell(row, index):
            return row[index] if index is not None and index < len(row) else ('', False)

        for number, row in body:
            source, source_formula = cell(row, source_col)
            target, target_formula = cell(row, target_col)
            key = cell(row, optional.get('key'))[0].strip()
            if not (source.strip() or target.strip() or key):
                continue
            records.append({'sheet': sheet, 'row': number, 'key': key,
                            'source': source, 'source_formula': source_formula,
                            'target': target, 'target_formula': target_formula,
                            'limit': cell(row, optional.get('limit'))[0].strip(),
                            'terms': cell(row, term_col)[0], 'issues': []})
    if not used_sheets:
        raise ValueError(f'{display_name} 未找到“{columns["source"]}”源文列，请核对列名')
    return records, used_sheets
def add(record, level, check, message):
    record['issues'].append({'level': level, 'check': check, 'message': message})


def parse_limit(value):
    if not value:
        return None
    try:
        number = float(value)
    except ValueError:
        return False
    return int(number) if number.is_integer() and number > 0 else False


def check_record(record, ratio):
    source, target = record['source'], record['target']
    if record['source_formula']:
        add(record, '提示', '源文异常', '源文单元格为公式，请核对实际源文')
    if record['target_formula']:
        add(record, '错误', '漏译', '译文单元格为公式，不是文本')
        return
    if not source.strip():
        if target.strip():
            add(record, '提示', '源文异常', '源文为空但有译文，请确认是否错行')
        return
    if not target.strip():
        add(record, '错误', '漏译', '有源文但译文为空')
        return

    for check, finder in (('参数', PARAM_PATTERN), ('标签', TAG_PATTERN)):
        expected, actual = Counter(finder.findall(source)), Counter(finder.findall(target))
        if expected != actual:
            add(record, '错误', check, describe(expected, actual))
    # 图片 <sprite> 允许随语序移动，其余标签的先后顺序应与源文一致。
    def order(value):
        return [tag for tag in TAG_PATTERN.findall(value) if not SPRITE_PATTERN.match(tag)]
    if all(issue['check'] != '标签' for issue in record['issues']) and order(source) != order(target):
        add(record, '提示', '标签顺序', '标签数量一致但先后顺序不同（图片标签除外）')
    if line_breaks(source) != line_breaks(target):
        (literal_s, real_s), (literal_t, real_t) = line_breaks(source), line_breaks(target)
        add(record, '错误', '换行', f'源文字面 \\n {literal_s} 个、实际换行 {real_s} 个；译文分别为 {literal_t}、{real_t} 个')
    expected, actual = numbers(source), numbers(target)
    if expected != actual:
        add(record, '错误', '数字', describe(expected, actual))

    length, limit = display_length(target), parse_limit(record['limit'])
    record['length'] = length
    if limit is False:
        add(record, '提示', '字数限制', f'limit 值“{record["limit"]}”无法识别，未检查字限')
    elif limit and length > limit:
        add(record, '错误', '字数限制', f'译文 {length} 字符，超过限制 {limit}')
    elif not limit and ratio:
        source_length = display_length(source)
        if source_length and length > source_length * ratio and length - source_length >= 5:
            add(record, '提示', '长度比例', f'无字限；译文 {length} 字符，为源文 {source_length} 字符的 {length / source_length:.0%}')

    if target.strip() == source.strip() and len(re.findall(r'[^\W\d_]', TAG_PATTERN.sub('', source))) >= 2:
        add(record, '提示', '未翻译', '译文与源文完全相同，请确认是否需要翻译')
    if (source[:1].isspace(), source[-1:].isspace()) != (target[:1].isspace(), target[-1:].isspace()):
        add(record, '提示', '首尾空格', '译文首尾空格与源文不一致')
    folded = target.casefold()
    for term, variants in parse_terms(record['terms']):
        if not any(v.casefold() in folded for v in variants):
            add(record, '提示', '术语', f'未在译文中找到“{term}”的术语译法：{" / ".join(variants)}（词形变化时可忽略）')


def check_consistency(records):
    groups = defaultdict(list)
    for record in records:
        if record['source'].strip() and record['target'].strip():
            groups[record['source'].strip()].append(record)
    for group in groups.values():
        targets = {r['target'].strip() for r in group}
        if len(targets) > 1:
            places = '、'.join(f'{r["sheet"]} 第{r["row"]}行' for r in group)
            for record in group:
                add(record, '提示', '译法不一致', f'相同源文有 {len(targets)} 种译法：{places}')


def compare_original(records, originals):
    """按顺序逐条对照原始待译表，发现增删行、错行和覆盖已有译文。"""
    issues = []
    if len(records) != len(originals):
        issues.append({'level': '错误', 'check': '行对应',
                       'message': f'记录数不一致：原始文件 {len(originals)} 条，译后文件 {len(records)} 条'})
    shifted = 0
    for record, original in zip(records, originals):
        if (record['source'], record['key']) != (original['source'], original['key']):
            shifted += 1
            if shifted <= MAX_SHIFT_REPORTS:
                add(record, '错误', '行对应', f'与原始文件第 {original["row"]} 行的 Key 或源文不一致，可能错行或源文被修改')
        elif original['target'].strip() and original['target'] != record['target']:
            add(record, '提示', '行对应', '原始文件已有译文，本次被修改')
    if shifted > MAX_SHIFT_REPORTS:
        issues.append({'level': '错误', 'check': '行对应',
                       'message': f'共 {shifted} 条与原始文件不对应，仅逐条列出前 {MAX_SHIFT_REPORTS} 条'})
    return issues
def status(record):
    levels = {issue['level'] for issue in record['issues']}
    return '错误' if '错误' in levels else '提示' if levels else '通过'


def write_cell(sheet, row, column, value):
    value = text(value)
    if len(value) > 32767:
        value = value[:32700] + '…［超过 Excel 单元格上限，完整内容见译后文件］'
    cell = sheet.cell(row, column, value)
    cell.data_type = 's'  # 以 = 开头的源文也按文本保存，避免被 Excel 当成公式
    cell.alignment = Alignment(vertical='top', wrap_text=True)
    return cell


def style_header(sheet, widths):
    for index, width in enumerate(widths, 1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='334155')
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = 'A2'


class QAChecker:
    def __init__(self, source_column, target_column, key_column='key', limit_column='limit',
                 term_columns=('生效术语', '关联术语'), length_ratio=1.5):
        self.columns = {'source': source_column, 'target': target_column, 'key': key_column,
                        'limit': limit_column, 'terms': tuple(term_columns)}
        self.length_ratio = length_ratio

    def write_report(self, path, records, file_issues, summary):
        book = Workbook()
        sheet = book.active
        sheet.title = '问题清单'
        headers = ('工作表', '原行号', 'Key', '源文', '译文', '字限', '译文长度', '级别', '检查项', '说明')
        sheet.append(headers)
        style_header(sheet, (14, 9, 28, 50, 50, 8, 10, 8, 12, 60))
        row = 2
        for issue in file_issues:
            for column, value in enumerate(('全文件', '', '', '', '', '', '', issue['level'], issue['check'], issue['message']), 1):
                write_cell(sheet, row, column, value)
            row += 1
        for record in records:
            for issue in record['issues']:
                values = (record['sheet'], record['row'], record['key'], record['source'], record['target'],
                          record['limit'], record.get('length', ''), issue['level'], issue['check'], issue['message'])
                for column, value in enumerate(values, 1):
                    write_cell(sheet, row, column, value)
                sheet.cell(row, 2).value = record['row']
                row += 1
        sheet.auto_filter.ref = f'A1:J{max(row - 1, 1)}'
        overview = book.create_sheet('汇总')
        overview.append(('项目', '内容'))
        style_header(overview, (22, 90))
        for index, (name, value) in enumerate(summary, 2):
            write_cell(overview, index, 1, name)
            write_cell(overview, index, 2, value)
        temp = f'{path}.{uuid.uuid4().hex}.tmp.xlsx'
        try:
            book.save(temp)
            os.replace(temp, path)
        finally:
            book.close()
            if os.path.exists(temp):
                os.remove(temp)
    def process(self, translated_path, output_path, task_dict, task_id, original_path=None):
        task = task_dict[task_id]
        try:
            name = task.get('input_name') or Path(translated_path).name
            task.update(progress=10, status='正在读取译后文件...')
            records, sheets = read_records(translated_path, name, self.columns)
            if not records:
                raise ValueError('译后文件没有可检查的记录')
            file_issues = []
            original_name = task.get('original_name') or ''
            if original_path:
                task.update(progress=30, status='正在对照原始待译文件...')
                originals, _ = read_records(original_path, original_name or '原始文件', self.columns, require_target=False)
                file_issues = compare_original(records, originals)
            task.update(progress=50, status=f'正在逐条检查 {len(records)} 条记录...')
            for record in records:
                check_record(record, self.length_ratio)
            check_consistency(records)

            for record in records:
                record['status'] = status(record)
            counts = Counter(record['status'] for record in records)
            issues = [issue for record in records for issue in record['issues']] + file_issues
            by_check = Counter(issue['check'] for issue in issues)
            level_counts = Counter(issue['level'] for issue in issues)
            translated = sum(bool(r['target'].strip()) and not r['target_formula'] for r in records)
            summary = [
                ('检查规则', QA_POLICY),
                ('译后文件', f'{name}（sha256 {hashlib.sha256(Path(translated_path).read_bytes()).hexdigest()[:12]}）'),
                ('原始文件', original_name if original_path else '未上传，未做行对应检查'),
                ('检查工作表', '、'.join(sheets)),
                ('列设置', f'源文 {self.columns["source"]}，译文 {self.columns["target"]}，'
                          f'Key {self.columns["key"] or "无"}，字限 {self.columns["limit"] or "无"}'),
                ('长度规则', LENGTH_RULE + (f'；无字限时译文超过源文 {self.length_ratio:g} 倍给出提示'
                                         if self.length_ratio else '；未启用长度比例检查')),
                ('记录数', f'共 {len(records)} 条，已有译文 {translated} 条'),
                ('记录结果', f'通过 {counts["通过"]} 条，有错误 {counts["错误"]} 条，仅有提示 {counts["提示"]} 条'),
                ('问题数', f'错误 {level_counts["错误"]} 个，提示 {level_counts["提示"]} 个'),
                *((f'检查项：{check}', f'{by_check[check]} 个') for check in CHECKS),
                ('说明', '本报告只检查完整性、格式、数字、长度与一致性，不评价译文质量；'
                       '术语提示按字面查找，德语等词形变化可能误报。语言质量仍需人工审校。'),
            ]
            task.update(progress=85, status='正在生成检查报告...')
            self.write_report(output_path, records, file_issues, summary)
            task.update(progress=100, state='completed', result_file=str(output_path), qa_policy=QA_POLICY,
                        records=[r for r in records if r['issues']], file_issues=file_issues,
                        counts={name: counts[name] for name in ('错误', '提示', '通过')},
                        check_counts={check: by_check[check] for check in CHECKS if by_check[check]},
                        total=len(records), translated=translated,
                        status=f'检查完成，共 {len(records)} 条：错误 {counts["错误"]} 条，提示 {counts["提示"]} 条，通过 {counts["通过"]} 条')
        except Exception as error:
            task.update(state='error', status=f'译后检查失败：{error}')