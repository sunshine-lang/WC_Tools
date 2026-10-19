import hashlib
import html
import re
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .context_processor import validate_material_name
from .key_classifier import EMOTIONS, KeyClassifier
from .term_filter import FILTER_COLUMNS


TM_POLICY = 'wc-en-de-local-v1'
TM_STATUSES = ('原文一致', '相似参考', '未命中', '待确认')
STOP_WORDS = set("a an the and or of to in on at for from with by as is are was were be been being it its this that these those i me my we our you your he him his she her they them their do does did have has had will would can could should shall may might must all some any just then than so very".split())


def local_name(tag):
    return tag.rsplit('}', 1)[-1]


def memo_payload(value):
    if re.match(r'^\s*<mq:(?:rxt|ch)\b', value):
        match = re.search(r'\bval\s*=\s*([\"\x27])(.*?)\1', value, re.S)
        if not match:
            raise ValueError('MemoQ 标记缺少 val，无法可靠还原')
        # Parsing this inner XML again would normalize literal newlines in attributes to spaces.
        return html.unescape(match.group(2))
    if re.fullmatch(r'\s*</?rpr\b[^>]*>\s*', value):
        return ''
    if re.match(r'^\s*<mq:', value):
        raise ValueError('不支持的 MemoQ 标记，请核对后重新导出')
    return value


def segment_text(node):
    parts = [node.text or '']
    for child in node:
        tag = local_name(child.tag)
        if tag in ('ph', 'bpt', 'ept', 'it', 'ut'):
            parts.append(memo_payload(''.join(child.itertext())))
        elif tag in ('hi', 'sub'):
            parts.append(segment_text(child))
        else:
            raise ValueError(f'不支持的 TMX 行内元素 {tag}，请核对后重新导出')
        parts.append(child.tail or '')
    return ''.join(parts)


def properties(node):
    return {p.get('type', ''): p.text or '' for p in node if local_name(p.tag) == 'prop'}


def context_text(value):
    if not value.lstrip().startswith('<seg'):
        return value
    try:
        node = ET.fromstring(value)
        if local_name(node.tag) == 'seg':
            return segment_text(node)
    except (ET.ParseError, ValueError):
        # Preserve unrecognized historical context; the original TMX remains the source of truth.
        return '原始上下文（格式未识别）：' + value
    return value


def visible_text(text):
    text = re.sub(r'<sprite\b[^>]*>', ' ', text, flags=re.I)
    return re.sub(r'</?[A-Za-z][^>]*>', '', text)


def normalized(text):
    return re.sub(r'\s+', ' ', visible_text(text)).strip().casefold().replace('’', "'")


def tokens(text):
    return set(re.findall(r"[^\W_]+(?:'[^\W_]+)?", normalized(text)))


def protected(text):
    return Counter(re.findall(r'\{\d+\}|%(?:\d+\$)?[sdf](?:%%)?|%\([^)]+\)[sdf]', text))


def changed_fragments(current, candidate):
    a, b = visible_text(current).split(), visible_text(candidate).split()
    return [{'current': ' '.join(a[i:j]), 'history': ' '.join(b[k:l])}
            for op, i, j, k, l in SequenceMatcher(None, a, b, autojunk=False).get_opcodes()
            if op != 'equal']


def differences(current, candidate, target):
    warnings = []
    if protected(current) != protected(candidate):
        warnings.append('当前原文与 TM 原文的参数不同')
    if protected(candidate) != protected(target):
        warnings.append('历史英德译文的参数不一致')
    if Counter(re.findall(r'\d+(?:[.,]\d+)?', visible_text(current))) != Counter(re.findall(r'\d+(?:[.,]\d+)?', visible_text(candidate))):
        warnings.append('当前原文与 TM 原文的数值不同')
    tag_pattern = r'</?[A-Za-z][^>]*>'
    if Counter(re.findall(tag_pattern, current)) != Counter(re.findall(tag_pattern, candidate)):
        warnings.append('当前原文与 TM 原文的标签不同')
    if Counter(re.findall(tag_pattern, candidate)) != Counter(re.findall(tag_pattern, target)):
        warnings.append('历史英德译文的标签不同')
    if (current.count('\\n'), current.count('\n')) != (candidate.count('\\n'), candidate.count('\n')):
        warnings.append('当前原文与 TM 原文的换行不同')
    negative = lambda s: bool(re.search(r"\b(?:no|not|never|without)\b|n't\b", normalized(s)))
    if negative(current) != negative(candidate):
        warnings.append('当前原文与 TM 原文的否定表达不同')
    return warnings


class TMMatcher:
    def __init__(self, tmx_path, source_name=None):
        path = Path(tmx_path)
        self.filename = source_name or path.name
        for name in (path.name, path.resolve().name, self.filename):
            validate_material_name(name)
        if path.suffix.lower() != '.tmx':
            raise ValueError('请上传英德 TMX 文件')
        content = path.read_bytes()
        if re.search(br'<!\s*ENTITY\b|<!\s*DOCTYPE[^>]*\[', content.replace(b'\x00', b''), re.I):
            raise ValueError('TMX 不支持自定义实体或内部 DTD，请重新导出标准 TMX')
        self.sha256 = hashlib.sha256(content).hexdigest()
        root = ET.fromstring(content)
        if local_name(root.tag) != 'tmx':
            raise ValueError('文件不是 TMX 文档')
        self.tu_count = self.valid_count = self.skipped_count = 0
        groups = {}
        for tu in root.iter():
            if local_name(tu.tag) != 'tu':
                continue
            self.tu_count += 1
            languages = defaultdict(list)
            for tuv in tu:
                if local_name(tuv.tag) != 'tuv':
                    continue
                language = (tuv.get('{http://www.w3.org/XML/1998/namespace}lang') or tuv.get('lang') or '').lower().replace('_', '-')
                base = language.split('-')[0]
                if base not in ('en', 'de'):
                    continue
                seg = next((s for s in tuv if local_name(s.tag) == 'seg'), None)
                if seg is not None:
                    languages[base].append((segment_text(seg), language, properties(tuv)))
            if len(languages['en']) > 1 or len(languages['de']) > 1:
                raise ValueError(f'TU {self.tu_count} 含多个英语或德语版本，请提供唯一英德语言对')
            if not languages['en'] or not languages['de']:
                self.skipped_count += 1
                continue
            english, en_lang, en_props = languages['en'][0]
            german, de_lang, _ = languages['de'][0]
            if not english.strip() or not german.strip():
                self.skipped_count += 1
                continue
            self.valid_count += 1
            props = properties(tu)
            origin = {'tu': self.tu_count, 'tuid': tu.get('tuid', ''),
                      'document': props.get('x-document', '').strip(),
                      'date': tu.get('changedate') or tu.get('creationdate') or '',
                      'source_language': en_lang, 'target_language': de_lang,
                      'previous': context_text(en_props.get('x-context-pre', '')),
                      'next': context_text(en_props.get('x-context-post', ''))}
            variants = groups.setdefault(english, {})
            variants.setdefault(german, []).append(origin)
        if not groups:
            raise ValueError('TMX 未找到有效英语→德语句对，请核对 en/de 语言与空白译文')
        self.groups = [{'english': en, 'variants': [{'german': de, 'origins': origins} for de, origins in variants.items()]}
                       for en, variants in groups.items()]
        self.exact = {g['english']: index for index, g in enumerate(self.groups)}
        self.text_index = defaultdict(list)
        self.word_index = defaultdict(set)
        self.clean = []
        self.words = []
        for index, group in enumerate(self.groups):
            clean = normalized(group['english'])
            words = tokens(group['english']) - STOP_WORDS
            self.clean.append(clean)
            self.words.append(words)
            self.text_index[clean].append(index)
            for word in words:
                self.word_index[word].add(index)

    def status(self):
        return {'ready': True, 'filename': self.filename, 'sha256': self.sha256,
                'tu_count': self.tu_count, 'valid_count': self.valid_count,
                'source_count': len(self.groups), 'pair_count': sum(len(g['variants']) for g in self.groups),
                'conflict_count': sum(len(g['variants']) > 1 for g in self.groups),
                'skipped_count': self.skipped_count, 'method': '英文完全匹配与文字相似匹配'}

    def find_matches(self, text):
        if text in self.exact:
            selected = [(self.exact[text], 1.0, '原文一致')]
        elif normalized(text) in self.text_index:
            selected = [(i, 1.0, '正文一致，格式或大小写有差异') for i in self.text_index[normalized(text)]]
        else:
            query_words = tokens(text) - STOP_WORDS
            if len(query_words) < 2:
                return []
            overlap = Counter(i for word in query_words for i in self.word_index.get(word, ()))
            ranked = []
            for index, count in overlap.items():
                if count < 2:
                    continue
                dice = 2 * count / (len(query_words) + len(self.words[index]))
                ranked.append((dice, index))
            selected = []
            # shortcut: lexical candidates only; add semantic retrieval after measuring real misses.
            for dice, index in sorted(ranked, reverse=True)[:60]:
                ratio = SequenceMatcher(None, normalized(text), self.clean[index], autojunk=False).ratio()
                score = .6 * ratio + .4 * dice
                if ratio >= .6 and dice >= .5 and score >= .72:
                    selected.append((index, score, '文字相似'))
            selected.sort(key=lambda item: (-item[1], item[0]))
            selected = selected[:3]
        matches = []
        for index, score, kind in selected:
            group = self.groups[index]
            for variant in group['variants']:
                warnings = differences(text, group['english'], variant['german'])
                if len(group['variants']) > 1:
                    warnings.append(f"同一英文有 {len(group['variants'])} 种德语译法，需结合语境选择")
                matches.append({'english': group['english'], 'german': variant['german'],
                                'match_type': kind, 'score': round(score, 4),
                                'origins': variant['origins'], 'warnings': warnings,
                                'changes': changed_fragments(text, group['english'])})
        return matches

    def format_matches(self, matches):
        chunks = []
        for index, match in enumerate(matches, 1):
            origin = match['origins'][0]
            source = f"{self.filename} / {origin['document'] or '未提供原文档名'} / TU {origin['tu']}"
            lines = [f"[TM参考 {index}｜{match['match_type']}]", f"【英】{match['english']}",
                     f"【德】{match['german']}", f"来源：{source}"]
            if len(match['origins']) > 1:
                lines.append(f"同句同译共 {len(match['origins'])} 处来源，完整记录可在页面查看")
            if match['warnings']:
                lines.append('提示：' + '；'.join(match['warnings']))
            for change in match['changes']:
                lines.append(f"差异：当前〔{change['current'] or '无'}〕／历史〔{change['history'] or '无'}〕")
            chunks.append('\n'.join(lines))
        return '\n\n'.join(chunks)

    def process(self, trans_path, output_path, task_dict, task_id):
        task = task_dict[task_id]
        book = None
        try:
            path = Path(trans_path)
            for name in (path.name, path.resolve().name, task.get('input_name', '')):
                validate_material_name(name)
            if path.suffix.lower() != '.xlsx':
                raise ValueError('请上传第四步术语筛选下载的 XLSX 文件')
            task.update(progress=5, status='正在读取第四步术语筛选结果...')
            book = load_workbook(path)
            records = []
            for sheet in book:
                headers = {str(c.value).strip().casefold(): c.column for c in sheet[1] if c.value is not None}
                en_col = headers.get('game_text_en') or headers.get('en')
                if not en_col:
                    continue
                required = {'分类', 'context', *FILTER_COLUMNS}
                if not required <= headers.keys():
                    raise ValueError(f'“{sheet.title}”缺少术语筛选结果，请先完成第四步术语语境筛选')
                if any(name.startswith('tm参考') for name in headers):
                    raise ValueError('文件已有 TM 参考列，请上传第四步结果，避免覆盖旧数据')
                width = max(c.column for row in sheet for c in row if c.value is not None)
                col = width + 1
                header = sheet.cell(1, col, 'TM参考')
                header.font = Font(bold=True, color='FFFFFF')
                header.fill = PatternFill('solid', fgColor='334155')
                sheet.column_dimensions[header.column_letter].width = 90

                def value(row, name):
                    index = headers.get(name)
                    raw = sheet.cell(row, index).value if index else None
                    return '' if raw is None else str(raw)

                for row in range(2, sheet.max_row + 1):
                    if not any(sheet.cell(row, c).value is not None for c in range(1, width + 1)):
                        continue
                    cell = sheet.cell(row, en_col)
                    english = '' if cell.value is None else str(cell.value)
                    matches, warning = [], ''
                    if cell.data_type == 'f' or not isinstance(cell.value, str) or not english.strip():
                        status, warning = '待确认', '英文源文为空、非文本或为公式，请核对源文'
                    else:
                        matches = self.find_matches(english)
                        status = '未命中' if not matches else '原文一致' if any(m['match_type'] == '原文一致' for m in matches) else '相似参考'
                        if any(m['warnings'] for m in matches):
                            status, warning = '待确认', '候选有差异或历史译法冲突，请展开核对'
                    reference = self.format_matches(matches)
                    if len(reference) > 32767:
                        raise ValueError(f'{sheet.title} 第{row}行 TM 参考超过单元格上限，未截断，请缩小候选范围')
                    added = sheet.cell(row, col, reference)
                    added.data_type = 's'
                    added.alignment = Alignment(vertical='top', wrap_text=True)
                    key, speaker = value(row, 'key'), value(row, 'speaker')
                    emotion = value(row, 'game_text_cn') or value(row, 'cn')
                    if not speaker and emotion.strip().casefold() in EMOTIONS:
                        speaker, key = key, ''
                    records.append({'row': row, 'sheet': sheet.title, 'key': key, 'speaker': speaker,
                                    'english': english, 'category': value(row, '分类'), 'status': status,
                                    'warning': warning, 'reference': reference, 'matches': matches})
                    if len(records) % 20 == 0:
                        task.update(progress=min(90, 10 + int(75 * row / max(sheet.max_row, 1))),
                                    status=f'正在匹配，已处理 {len(records)} 条...')
            if not records:
                raise ValueError('未找到待处理记录，须包含英文源文和第四步术语筛选结果列')
            # openpyxl reads empty inline strings as None; restore them before round-tripping.
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        if cell.data_type == 'inlineStr' and cell.value is None:
                            cell.value = ''
            KeyClassifier.save_workbook(book, output_path)
            counts = Counter(r['status'] for r in records)
            task.update(progress=100, state='completed', records=records,
                        counts={name: counts[name] for name in TM_STATUSES},
                        matched_count=sum(bool(r['matches']) for r in records),
                        result_file=str(output_path), tm_db=self.status(), tm_policy=TM_POLICY,
                        status=f'匹配完成，共 {len(records)} 条；有参考 {sum(bool(r["matches"]) for r in records)} 条，待确认 {counts["待确认"]} 条')
        except Exception as error:
            task.update(state='error', status=f'TM 匹配失败：{error}')
        finally:
            if book:
                book.close()
