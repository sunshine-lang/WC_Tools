import csv
import hashlib
import html
import re
from collections import Counter
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .context_processor import validate_material_name
from .key_classifier import EMOTIONS, KeyClassifier


from .language import language_pair, HAN


TERM_POLICY = 'wc-pair-terms-v2'
TERM_STATUSES = ('命中术语', '未命中', '待确认', '未使用')

# Only known noun forms are expanded; names and unrelated endings are not stemmed.
NOUN_PLURALS = {
    'board': 'boards', 'item': 'items', 'reward': 'rewards', 'review': 'reviews',
    'event': 'events', 'coin': 'coins', 'buddy': 'buddies', 'safe': 'safes',
    'butterfly': 'butterflies', 'key': 'keys', 'chain': 'chains', 'bird': 'birds',
    'castle': 'castles', 'fireman': 'firemen', 'donut': 'donuts', 'bowl': 'bowls',
    'case': 'cases', 'box': 'boxes', 'hive': 'hives', 'chest': 'chests',
    'pack': 'packs', 'truck': 'trucks', 'gift': 'gifts', 'gem': 'gems',
    'frog': 'frogs', 'ghost': 'ghosts', 'jacket': 'jackets', 'scarf': 'scarves',
    'babysitter': 'babysitters', 'station': 'stations', 'lane': 'lanes',
    'clue': 'clues', 'token': 'tokens', 'investigation': 'investigations',
    'bookshelf': 'bookshelves', 'passage': 'passages', 'ring': 'rings',
    'room': 'rooms', 'suitcase': 'suitcases', 'ball': 'balls', 'trap': 'traps',
    'tripwire': 'tripwires', 'stock': 'stocks', 'lodge': 'lodges',
    'claim': 'claims', 'apiary': 'apiaries', 'company': 'companies', 'date': 'dates',
}


def normalized(text):
    return str(text).replace('’', "'").replace('‘', "'")


class TermScanner:
    def __init__(self, tb_file_path, source_name=None, language=None):
        self.language = language_pair(language)
        self.filename, self.sha256 = '未提供术语库', ''
        self.row_count = self.conflict_count = 0
        self.terms, self.patterns = [], []
        if tb_file_path is None:
            return
        path = Path(tb_file_path)
        self.filename = source_name or path.name
        validate_material_name(path.name)
        validate_material_name(path.resolve().name)
        validate_material_name(self.filename)
        if path.suffix.lower() not in ('.xlsx', '.csv'):
            raise ValueError('术语库须为包含 当前语向的源语、目标语两列的 XLSX 或 UTF-8 CSV 文件')
        contents = path.read_bytes()
        self.sha256 = hashlib.sha256(contents).hexdigest()
        self.row_count = 0
        by_english = {}

        def add_term(source_text, target_text, source, formula=False):
            if source_text in (None, '') and target_text in (None, ''):
                return
            if formula or any(not isinstance(v, str) or not v.strip() for v in (source_text, target_text)):
                raise ValueError(f'{source} 的 语言对 为空、非文本或为公式，请核对术语库')
            source_text, target_text = source_text.strip(), target_text.strip()
            key = ' '.join(normalized(source_text).casefold().split())
            entry = by_english.setdefault(key, {'source_text': source_text, 'target_text': [], 'sources': []})
            if target_text not in entry['target_text']:
                entry['target_text'].append(target_text)
            entry['sources'].append(f'{self.filename} / {source}')
            self.row_count += 1

        if path.suffix.lower() == '.xlsx':
            workbook = load_workbook(BytesIO(contents), data_only=False)
            try:
                for sheet in workbook:
                    columns = self.language.term_columns(c.value for c in sheet[1])
                    if None in columns:
                        continue
                    for row in range(2, sheet.max_row + 1):
                        en, de = (sheet.cell(row, index + 1) for index in columns)
                        if en.font.strike or de.font.strike:
                            continue
                        add_term(en.value, de.value, f'{sheet.title} 第{row}行（{en.coordinate}:{de.coordinate}）',
                                 formula=en.data_type == 'f' or de.data_type == 'f')
            finally:
                workbook.close()
        else:
            try:
                table = csv.reader(StringIO(contents.decode('utf-8-sig')), strict=True)
                columns = self.language.term_columns(next(table, []))
                if None in columns:
                    raise ValueError('术语 CSV 须在首行提供当前语向表头（EN/DE 或 ZH/EN）')
                for number, row in enumerate(table, 2):
                    if not row:
                        continue
                    values = [row[index] if index < len(row) else None for index in columns]
                    add_term(*values, f'CSV 第{number}行')
            except UnicodeDecodeError as error:
                raise ValueError('请将术语 CSV 保存为 UTF-8 编码后上传') from error
        self.terms = list(by_english.values())
        if not self.terms:
            raise ValueError('术语库未找到有效 ' + ('ZH／EN' if self.language.chinese else 'EN／DE') + ' 词条，请检查首行表头')

        self.conflict_count = sum(len(term['target_text']) > 1 for term in self.terms)
        self.patterns = []
        for index, term in enumerate(self.terms):
            words = normalized(term['source_text']).split()
            forms = [(' '.join(words), '原词')]
            last = words[-1].casefold()
            for singular, plural in ({} if self.language.chinese else NOUN_PLURALS).items():
                if last in (singular, plural):
                    alternate = plural if last == singular else singular
                    forms.append((' '.join([*words[:-1], alternate]), '词形'))
                    break
            for form, kind in forms:
                body = r'\s+'.join(re.escape(word) for word in form.split())
                left = '' if self.language.chinese and re.match(HAN, form[0]) else (r'(?<![A-Za-z0-9_])' if self.language.chinese else r'(?<!\w)')
                right = '' if self.language.chinese and re.match(HAN, form[-1]) else (r'(?![A-Za-z0-9_])' if self.language.chinese else r'(?!\w)')
                self.patterns.append((index, kind, re.compile(left + body + right, re.IGNORECASE)))

    def status(self):
        return {'ready': bool(self.terms), 'enabled': bool(self.terms), 'filename': self.filename, 'row_count': self.row_count,
                'term_count': len(self.terms), 'conflict_count': self.conflict_count,
                'sha256': self.sha256}

    def find_matches(self, text):
        if not isinstance(text, str) or not text.strip():
            return []
        visible = re.sub(r'<!--[\s\S]*?-->', '', html.unescape(text))
        # Inline formatting joins visible text; sprite objects separate it.
        visible = re.sub(r'<sprite\b[^>]*>', ' \ufffc ', visible, flags=re.IGNORECASE)
        visible = re.sub(r'</?[A-Za-z][^>]*>', '', visible)
        visible = normalized(visible)
        candidates = []
        for index, kind, pattern in self.patterns:
            for match in pattern.finditer(visible):
                candidates.append((match.start(), match.end(), index, kind, match.group()))
        candidates.sort(key=lambda m: (-(m[1] - m[0]), m[3] != '原词', m[2], m[0]))
        occupied, selected = [], []
        for start, end, index, kind, matched in candidates:
            if any(start < right and end > left for left, right in occupied):
                continue
            occupied.append((start, end))
            selected.append((start, index, kind, matched))
        result, seen = [], set()
        for _, index, kind, matched in sorted(selected):
            if index in seen:
                continue
            seen.add(index)
            result.append({**self.terms[index], 'matched_text': matched, 'match_type': kind})
        return result

    @staticmethod
    def format_matches(matches):
        return ' | '.join(f"{m['source_text']} ({' / '.join(m['target_text'])})" for m in matches)

    def scan_text(self, text):
        return self.format_matches(self.find_matches(text))

    def process(self, trans_file_path, output_file, task_dict, task_id):
        task = task_dict[task_id]
        workbook = None
        try:
            path = Path(trans_file_path)
            validate_material_name(path.name)
            validate_material_name(path.resolve().name)
            validate_material_name(task.get('input_name', ''))
            if path.suffix.lower() != '.xlsx':
                raise ValueError('请上传第二步下载的 XLSX 文件')
            task.update(progress=5, status='正在读取源文和当前语向术语库...')
            workbook = load_workbook(path)
            records = []
            for sheet in workbook:
                headers = self.language.headers(c.value for c in sheet[1])
                en_col = self.language.column(headers, 'source')
                if not en_col:
                    continue
                if not {'分类', 'context'} <= headers.keys():
                    raise ValueError(f'“{sheet.title}”缺少分类或 Context 列，请上传第二步下载的结果')
                if '关联术语' in headers:
                    raise ValueError('文件已有“关联术语”列，请上传第二步结果，避免覆盖旧数据')
                width = max(c.column for row in sheet for c in row if c.value is not None)
                column = width + 1
                cell = sheet.cell(1, column, '关联术语')
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='334155')
                sheet.column_dimensions[cell.column_letter].width = 65

                def value(row, *names):
                    col = next((headers[n] for n in names if n in headers), None)
                    raw = sheet.cell(row, col).value if col else None
                    return '' if raw is None else str(raw)

                for row in range(2, sheet.max_row + 1):
                    if not any(sheet.cell(row, col).value is not None for col in range(1, width + 1)):
                        continue
                    source = sheet.cell(row, en_col)
                    source_text = '' if source.value is None else str(source.value)
                    warning = ''
                    matches = []
                    if source.data_type == 'f' or not isinstance(source.value, str) or not source_text.strip():
                        warning = '源文为空、非文本或为公式，请先核对源文'
                    else:
                        matches = self.find_matches(source_text)
                        if any(len(m['target_text']) > 1 for m in matches):
                            warning = '同一源词存在多个目标语候选，请结合语境确认'
                    result = self.format_matches(matches)
                    if len(result) > 32767:
                        raise ValueError(f'{sheet.title} 第{row}行术语参考超过 Excel 单元格上限')
                    cell = sheet.cell(row, column, result)
                    cell.data_type = 's'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    get = lambda *names: value(row, *names)
                    key = self.language.value(get, 'key')
                    speaker = self.language.value(get, 'speaker')
                    emotion = self.language.emotion(get)
                    if not speaker and emotion.strip().casefold() in EMOTIONS:
                        speaker, key = key, ''
                    records.append({'sheet': sheet.title, 'row': row, 'key': key, 'speaker': speaker,
                                    'source_text': source_text, 'category': value(row, '分类'),
                                    'status': '待确认' if warning else '命中术语' if matches else '未命中' if self.terms else '未使用',
                                    'terms': result, 'matches': matches, 'warning': warning})
                task.update(progress=80, status=f'已扫描 {len(records)} 条，正在保存 Excel...')
            if not records:
                raise ValueError('未找到待处理记录，须包含源文、分类和 Context 列')
            KeyClassifier.save_workbook(workbook, output_file)
            counts = Counter(r['status'] for r in records)
            task.update(progress=100, state='completed', result_file=str(output_file),
                        records=records, counts={s: counts[s] for s in TERM_STATUSES},
                        term_db=self.status(), term_policy=TERM_POLICY,
                        status=f'匹配完成，共 {len(records)} 条；命中 {counts["命中术语"]} 条，未命中 {counts["未命中"]} 条，待确认 {counts["待确认"]} 条')
            if not self.terms:
                reason = '未提供术语库，已跳过术语匹配，候选列留空'
                task.update(skipped=True, skip_reason=reason, status=reason)
        except Exception as error:
            task.update(status=f'术语关联失败：{error}', state='error')
        finally:
            if workbook:
                workbook.close()
