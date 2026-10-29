"""Read layout evidence without changing cells or treating a colour as a category."""
from openpyxl.utils.cell import coordinate_to_tuple

LAYOUT_POLICY = 'wc-layout-v1'


def color_signature(color):
    return (color.type, str(color.index), color.tint) if color else None


def cell_signature(cell):
    fill = cell.fill
    return {
        'fill': (fill.patternType, color_signature(fill.fgColor), color_signature(fill.bgColor)),
        'font': (bool(cell.font.bold), cell.font.underline),
        'top': getattr(cell.border.top, 'style', None), 'bottom': getattr(cell.border.bottom, 'style', None),
    }


def image_anchors(sheet):
    result = []
    for index, image in enumerate(sheet._images):
        anchor = image.anchor
        start, end = getattr(anchor, '_from', None), getattr(anchor, 'to', None)
        if isinstance(anchor, str):
            row, col = coordinate_to_tuple(anchor)
            last = None
        elif start is not None:
            row, col = start.row + 1, start.col + 1
            last = max(row, end.row + (1 if end.rowOff else 0)) if end else None
        else:
            continue
        result.append({'image_id': f'{sheet.title}:image:{index + 1}', 'sheet': sheet.title,
                       'start_row': row, 'end_row': last, 'start_column': col,
                       'end_column': end.col + 1 if end else None,
                       'extent_known': last is not None})
    return result


def segment_sheet(sheet, columns, fences=None, enabled=True):
    """Require multi-column changes; a lone coloured cell or image cannot split a region."""
    cols = sorted(set(c for c in columns if c))
    anchors = image_anchors(sheet) if enabled else []
    fences = fences or {}
    merges = [m for m in sheet.merged_cells.ranges if any(m.min_col <= c <= m.max_col for c in cols)]
    rows, regions, previous, current = {}, {}, None, None
    for row in range(2, sheet.max_row + 1):
        populated = any(sheet.cell(row, c).value not in (None, '') for c in cols)
        if not populated:
            previous, current = None, None
            continue
        signatures = [cell_signature(sheet.cell(row, c)) for c in cols]
        reasons = []
        if current is None:
            reasons.append('sheet_start_or_blank_row')
        elif row in fences:
            reasons.append(fences[row])
        if previous and enabled:
            fills = sum(a['fill'] != b['fill'] for a, b in zip(previous, signatures))
            fonts = sum(a['font'] != b['font'] for a, b in zip(previous, signatures))
            borders = sum(bool(a['bottom'] or b['top']) for a, b in zip(previous, signatures))
            if fills >= 2:
                reasons.append('multi_column_fill_change')
            if borders >= 2:
                reasons.append('multi_column_border')
            if fonts >= 2 and (fills or borders):
                reasons.append('font_and_layout_change')
            if any(m.min_row == row or m.max_row + 1 == row for m in merges):
                reasons.append('merged_region_boundary')
        if reasons:
            current = f'{sheet.title}:region:{row}'
            regions[current] = {'region_id': current, 'rows': [], 'boundary_evidence': reasons,
                                'formatted': False, 'image_ids': []}
        region = regions[current]
        region['rows'].append(row)
        region['formatted'] |= enabled and any(s['fill'][0] or s['top'] or s['bottom'] for s in signatures)
        images = [a['image_id'] for a in anchors if a['start_row'] <= row <= (a['end_row'] or a['start_row'])]
        region['image_ids'] = list(dict.fromkeys([*region['image_ids'], *images]))
        rows[row] = {'region_id': current, 'style_signature': signatures if enabled else [],
                     'image_ids': images, 'merged_ranges': [str(m) for m in merges if m.min_row <= row <= m.max_row]}
        previous = signatures
    return {'policy': LAYOUT_POLICY, 'available': enabled, 'rows': rows,
            'regions': regions, 'images': anchors}
