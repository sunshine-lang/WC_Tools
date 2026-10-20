"""输入规范化。

客户导出的表格与模块期望的形状之间常有落差，本模块负责抹平，不改动任何业务判断：

- 去掉整列为空的多余列（导出常带几百个空列，会让每个产物后面挂一长串 None）
- 补出 `speaker` 列：说话人名写进 key 列时，分类器无法识别对白

产出的文件与输入在语义上等价，只是形状更规矩。
"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook

POLICY = 'normalize-v1'


class Normalizer:
    def __init__(self, speakers=(), key_column='key', speaker_column='speaker',
                 trim_empty_columns=True):
        self.speakers = {str(s).strip().casefold() for s in speakers if str(s).strip()}
        self.key_column = key_column.casefold()
        self.speaker_column = speaker_column
        self.trim_empty_columns = trim_empty_columns

    # ---------------------------------------------------------------- 读入

    @staticmethod
    def read_workbook(path):
        """xlsx 直接读；csv 转成单表 workbook，避免模块各自处理两种格式。"""
        path = Path(path)
        if path.suffix.lower() == '.csv':
            workbook = Workbook()
            sheet = workbook.active
            with open(path, encoding='utf-8-sig', newline='') as handle:
                for row in csv.reader(handle):
                    sheet.append(row)
            for row in sheet:
                for cell in row:
                    cell.data_type = 's'
            return workbook
        return load_workbook(path)

    # ---------------------------------------------------------------- 处理

    def process(self, input_file, output_file, task_dict, task_id):
        task = task_dict[task_id]
        try:
            task['status'] = '正在读取文件...'
            workbook = self.read_workbook(input_file)

            trimmed = 0
            added_speaker = 0
            for sheet in workbook.worksheets:
                if self.trim_empty_columns:
                    trimmed += self.trim(sheet)
                added_speaker += self.fill_speaker(sheet)

            task['status'] = '正在写出规范化文件...'
            workbook.save(output_file)
            task.update(
                progress=100, state='completed', result_file=output_file,
                status=f'规范化完成：删除空列 {trimmed} 个，补出 speaker {added_speaker} 行',
                trimmed_columns=trimmed, speaker_rows=added_speaker,
                policy=POLICY,
            )
        except Exception as error:  # noqa: BLE001 — 与其余模块保持一致的错误形状
            task.update(status=f'处理失败：{error}', state='error')

    # ---------------------------------------------------------------- 步骤

    @staticmethod
    def trim(sheet):
        """删除正文全空、且表头也空的尾随列，返回删除数量。"""
        last = 0
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value not in (None, ''):
                    last = max(last, cell.column)
        if last == 0 or last >= sheet.max_column:
            return 0
        removed = sheet.max_column - last
        sheet.delete_cols(last + 1, removed)
        return removed

    def fill_speaker(self, sheet):
        """key 列是说话人名时，补出 speaker 列；已有该列则不覆盖。"""
        if sheet.max_row < 2:
            return 0
        headers = {
            str(cell.value).strip().casefold(): cell.column
            for cell in sheet[1] if cell.value is not None
        }
        key_column = headers.get(self.key_column)
        if not key_column or not self.speakers:
            return 0

        speaker_column = headers.get(self.speaker_column.casefold())
        if speaker_column is None:
            speaker_column = max(
                (cell.column for row in sheet.iter_rows() for cell in row if cell.value not in (None, '')),
                default=0,
            ) + 1
            sheet.cell(1, speaker_column, self.speaker_column)

        filled = 0
        for row in range(2, sheet.max_row + 1):
            key = sheet.cell(row, key_column).value
            name = str(key).strip() if key is not None else ''
            if not name or name.casefold() not in self.speakers:
                continue
            if sheet.cell(row, speaker_column).value not in (None, ''):
                continue
            sheet.cell(row, speaker_column, name)
            filled += 1
        return filled
