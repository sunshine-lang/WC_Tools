import csv
import fnmatch
import hashlib
import json
import os
import re
import uuid
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


from .language import EMOTIONS, language_pair
from .layout_segmenter import segment_sheet


CATEGORIES = ('活动 UI 与商店文案', '集卡名称', '装饰物名称', '活动对白')
PENDING = '待确认'
META_HEADERS = ('分类',)
CLASSIFICATION_POLICY = 'wc-multi-evidence-v1'


class KeyClassifier:
    def __init__(self, config_path, language=None):
        self.config_path = config_path
        self.language = language_pair(language)

    def load_rules(self):
        with open(self.config_path, encoding='utf-8') as file:
            rules = json.load(file)
        self.validate_rules(rules)
        return rules

    @staticmethod
    def validate_rules(rules):
        if not isinstance(rules, dict) or set(rules) != set(CATEGORIES):
            raise ValueError('请保留 WC 的四种分类')
        for patterns in rules.values():
            if not isinstance(patterns, list) or len(patterns) > 50:
                raise ValueError('每类最多配置 50 条特征规则')
            if any(not isinstance(p, str) or not p.strip() or len(p) > 120 for p in patterns):
                raise ValueError('特征规则须为 1–120 个字符的非空文本')

    def save_rules(self, rules):
        self.validate_rules(rules)
        rules = {category: list(dict.fromkeys(p.strip() for p in rules[category])) for category in CATEGORIES}
        temp = f'{self.config_path}.{uuid.uuid4().hex}.tmp'
        try:
            with open(temp, 'w', encoding='utf-8') as file:
                json.dump(rules, file, ensure_ascii=False, indent=2)
            os.replace(temp, self.config_path)
        finally:
            if os.path.exists(temp):
                os.remove(temp)

    @staticmethod
    def text(value):
        return '' if value is None else str(value)

    def classify(self, key, source_text, emotion, speaker, rules):
        """Compatibility interface; workflow callers use the evidence record below."""
        result = self.classify_evidence(key, source_text, emotion, speaker, rules)
        return result['category'], result['basis']

    def classify_evidence(self, key, source_text, emotion, speaker, rules):
        evidence = []
        def result(category, state, confidence, basis):
            return dict(category=category, classification_status=state, confidence=confidence,
                        basis=basis, evidence=evidence, manual_category=None)
        if not source_text.strip():
            return result(PENDING, 'pending', 0, '缺少源文')
        hits = []
        for category, patterns in rules.items():
            matched = [p for p in patterns if key and (
                fnmatch.fnmatchcase(key.casefold(), p.casefold()) if any(c in p for c in '*?[')
                else p.casefold() in key.casefold())]
            if matched:
                hits.append(category)
                evidence.append({'source': 'key_pattern', 'category': category, 'patterns': matched})
        dialogue = bool(speaker.strip() or (emotion.casefold() in EMOTIONS and
                        re.fullmatch(r"[A-Za-z][A-Za-z .'-]{0,60}", key)))
        if dialogue:
            evidence.append({'source': 'speaker_field', 'category': CATEGORIES[3]})
        if len(hits) > 1:
            return result(PENDING, 'pending', 0, '同时命中多类规则：' + '、'.join(hits))
        if hits:
            if dialogue and hits[0] != CATEGORIES[3]:
                return result(PENDING, 'pending', 0, 'Key 与说话人证据冲突')
            return result(hits[0], 'confirmed', .95, 'Key 命中规则：' + '、'.join(evidence[0]['patterns']))
        if dialogue:
            return result(CATEGORIES[3], 'inferred', .7, '说话人字段提示对白，完整 Key 尚未证实')
        return result(PENDING, 'pending', 0, '无 Key，等待区块与文本证据或人工指定' if not key else '未命中分类规则')

    @staticmethod
    def content_kind(value):
        visible = re.sub(r'<[^>]*>', '', value).strip()
        if re.search(r'\b(?:tap|click|collect|unlock|join|help|claim|complete|earn|celebrate|merge|reward)\b|点击|收集|解锁|参与|帮助|领取|完成|奖励|合成', visible, re.I):
            return 'instruction'
        if visible and len(visible) <= 70 and len(visible.split()) <= 8 and not re.search(r'[.!?。！？\n]', visible):
            return 'short_name'
        return 'unknown'

    def resolve_evidence(self, sheet, records, headers, overrides, layout_available):
        # Reuse the same project identity parser as Context, including note-based dialogue IDs.
        from .context_processor import ContextProcessor
        fences, previous = {}, ''
        for record in records:
            identity = ContextProcessor.relation(record)
            record['activity_hint'], record['group_hint'] = identity[:2]
            if identity[0]:
                if previous and identity[0] != previous:
                    fences[record['row']] = 'explicit_activity_boundary'
                previous = identity[0]
        cols = [self.language.column(headers, field) for field in ('key', 'source', 'target', 'note')]
        layout = segment_sheet(sheet, cols, fences, enabled=layout_available)
        groups = defaultdict(list)
        for r in records:
            info = layout['rows'].get(r['row'], {})
            r['region_id'] = info.get('region_id', '')
            r['layout'] = info
            r['content_kind'] = self.content_kind(r['source_text'])
            r['evidence'].append({'source': 'text_content', 'feature': r['content_kind']})
            override = overrides.get(r['id'])
            if override:
                if override.get('record_hash') != r['record_hash'] or override.get('category') not in CATEGORIES:
                    raise ValueError('人工分类与当前源文不一致，请重新确认')
                if r['formula'] or not r['source_text'].strip():
                    raise ValueError('异常源文不能仅通过指定分类放行')
                r.update(category=override['category'], classification_status='confirmed', confidence=1,
                         manual_category=override['category'], basis='当前批次人工指定分类')
                r['evidence'].append({'source': 'manual_override', 'category': r['category']})
            groups[r['region_id']].append(r)
        for region_id, peers in groups.items():
            region = layout['regions'].get(region_id, {})
            anchors = [dict(r) for r in peers if r['classification_status'] == 'confirmed']
            counts = Counter(r['category'] for r in anchors)
            dominant, count = counts.most_common(1)[0] if counts else ('', 0)
            activities = {r['activity_hint'] for r in anchors if r['activity_hint']}
            for r in peers:
                r['region_activity'] = next(iter(activities)) if len(activities) == 1 else ''
                if region.get('formatted'):
                    r['evidence'].append({'source': 'layout_region', 'region_id': region_id,
                                          'boundary': region['boundary_evidence'], 'image_ids': region['image_ids']})
                others = [a for a in anchors if a['id'] != r['id']]
                if (not r['manual_category'] and region.get('formatted') and count >= 2
                        and len(counts) > 1 and r['category'] != dominant and r['classification_status'] == 'confirmed'
                        and len({a['category'] for a in others}) == 1):
                    r.update(category=PENDING, classification_status='pending', confidence=0,
                             basis='Key 与同一格式区块内多条明确分类冲突，请人工确认')
                    r['evidence'].append({'source': 'layout_conflict', 'record_ids': [a['id'] for a in others]})
                elif (r['category'] == PENDING and not r['key'] and not r['formula']
                      and region.get('formatted') and len(counts) == 1 and count >= 2 and len(activities) == 1):
                    compatible = ((dominant == CATEGORIES[0] and r['content_kind'] == 'instruction') or
                                  (dominant in CATEGORIES[1:3] and r['content_kind'] == 'short_name'))
                    if compatible:
                        r.update(category=dominant, classification_status='inferred', confidence=.75,
                                 basis='连续格式区块、明确分类邻项与文本特征一致，需人工确认')
                        r['evidence'].append({'source': 'same_region_peers', 'record_ids': [a['id'] for a in anchors]})
        return layout

    @staticmethod
    def counts(records):
        counts = Counter(record['category'] for record in records)
        return {category: counts[category] for category in (*CATEGORIES, PENDING)}

    @staticmethod
    def save_workbook(workbook, path):
        temp = f'{path}.{uuid.uuid4().hex}.tmp.xlsx'
        try:
            workbook.save(temp)
            os.replace(temp, path)
        finally:
            if os.path.exists(temp):
                os.remove(temp)

    def process(self, input_file, output_file, task_dict, task_id, overrides=None, layout_available=True):
        task = task_dict[task_id]
        workbook = None
        try:
            task['status'] = '正在读取文件...'
            rules = self.load_rules()
            if Path(input_file).suffix.lower() == '.csv':
                with open(input_file, encoding='utf-8-sig', newline='') as file:
                    data = list(csv.reader(file))
                workbook = Workbook()
                sheet = workbook.active
                for row in data:
                    sheet.append(row)
                for row in sheet:
                    for cell in row:
                        cell.data_type = 's'
            else:
                workbook = load_workbook(input_file)
            records, layouts = [], {}
            for sheet in workbook.worksheets:
                headers = self.language.headers(c.value for c in sheet[1])
                en_col = self.language.column(headers, 'source')
                if not en_col:
                    continue
                if any(name.casefold() in headers for name in (*META_HEADERS, '记录ID', '文本类型', '分类依据', '检查提示')):
                    raise ValueError('文件已有分类结果列，请上传原始待译文件')
                width = max(c.column for row in sheet for c in row if c.value is not None)
                columns = {name: width + i for i, name in enumerate(META_HEADERS, 1)}
                for name, col in columns.items():
                    cell = sheet.cell(1, col, name)
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='334155')
                    sheet.column_dimensions[cell.column_letter].width = 30

                def value(row, *names):
                    col = next((headers[name] for name in names if name in headers), None)
                    return self.text(sheet.cell(row, col).value).strip() if col else ''

                seen_keys = {}
                sheet_records = []
                for row in range(2, sheet.max_row + 1):
                    if not any(sheet.cell(row, c).value is not None for c in range(1, width + 1)):
                        continue
                    get = lambda *names: value(row, *names)
                    raw_key = self.language.value(get, 'key')
                    emotion = self.language.emotion(get)
                    speaker = self.language.value(get, 'speaker')
                    source_cell = sheet.cell(row, en_col)
                    source_text = self.text(source_cell.value)
                    from .context_processor import dialogue_identity
                    effective_key, speaker = dialogue_identity(raw_key, speaker, self.language.value(get, 'note'))
                    decision = self.classify_evidence(effective_key, source_text, emotion, speaker, rules)
                    category, basis = decision['category'], decision['basis']
                    warnings = []
                    if source_cell.data_type == 'f':
                        category, basis = PENDING, '源文单元格为公式，请核对实际源文'
                    if not source_text.strip():
                        warnings.append('缺少源文')
                    inferred_dialogue = not speaker and emotion.casefold() in EMOTIONS and bool(
                        re.fullmatch(r"[A-Za-z][A-Za-z .'-]{0,60}", raw_key)
                    )
                    if not raw_key or inferred_dialogue:
                        warnings.append('缺少文案 Key；使用内部记录 ID 定位')
                    if not raw_key:
                        warnings.append('请确认无 Key 文案的用途')
                    if raw_key and not inferred_dialogue:
                        if raw_key in seen_keys:
                            warnings.append(f'Key 与第 {seen_keys[raw_key]} 行重复')
                        seen_keys[raw_key] = row
                    limit = self.language.value(get, 'limit')
                    if not limit:
                        warnings.append('未提供长度限制')
                    record = {
                        'id': f'{sheet.title}:{row}', 'sheet': sheet.title,
                        'row': row, 'key': '' if inferred_dialogue and effective_key == raw_key else effective_key,
                        'speaker': raw_key if inferred_dialogue else speaker,
                        'emotion': emotion if speaker or inferred_dialogue else '',
                        'source_text': source_text, 'limit': limit, 'category': category,
                        **decision, 'category': category, 'basis': basis, 'warning': '；'.join(warnings),
                        'formula': source_cell.data_type == 'f',
                        'record_hash': hashlib.sha256(json.dumps([raw_key, source_text, speaker, emotion,
                            self.language.value(get, 'note')], ensure_ascii=False).encode()).hexdigest(),
                    }
                    if record['formula']:
                        record.update(classification_status='pending', confidence=0)
                    records.append(record)
                    sheet_records.append(record)
                layouts[sheet.title] = self.resolve_evidence(sheet, sheet_records, headers, overrides or {},
                    layout_available and Path(input_file).suffix.lower() != '.csv')
                for record in sheet_records:
                    self.write_record(sheet, columns, record)
                task['status'] = f'已识别 {len(records)} 条记录，正在生成分类文件...'
                task['progress'] = 80
            if not records:
                raise ValueError('未找到待分类记录；请核对源文列：' + ' / '.join(self.language.aliases('source')))
            self.save_workbook(workbook, output_file)
            task.update(progress=100, status=f'分类完成，共 {len(records)} 条',
                        state='completed', result_file=output_file, records=records,
                        counts=self.counts(records), layouts=layouts, classification_policy=CLASSIFICATION_POLICY,
                        review_required=any(r['classification_status'] != 'confirmed' for r in records))
        except Exception as error:
            task.update(status=f'处理失败：{error}', state='error')
        finally:
            if workbook:
                workbook.close()

    @staticmethod
    def write_record(sheet, columns, record):
        for name, value in zip(META_HEADERS, (record['category'],)):
            cell = sheet.cell(record['row'], columns[name], value)
            cell.alignment = Alignment(vertical='top', wrap_text=True)

