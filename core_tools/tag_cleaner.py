import csv
import os
import re
import uuid
from pathlib import Path

from openpyxl import load_workbook

from .context_processor import validate_material_name
from .key_classifier import KeyClassifier


TAG_CLEAN_POLICY = 'wc-tag-space-v1'


class TagCleaner:
    def __init__(self, column):
        self.column = column.strip()
        if not self.column:
            raise ValueError('请选择要清洗的列')

    def clean_text(self, value):
        if self.column.casefold() in ('cn', 'game_text_cn'):
            # RO 原规则仅用于中文列；英德列必须保留词间空格和换行。
            return re.sub(r'\s*(<[^>]+>)\s*', r'\1', value)
        return re.sub(r'( *)(<[^>]+>)( *)',
                      lambda m: (' ' if m[1] else '') + m[2] + (' ' if m[3] else ''), value)

    def process(self, input_file, output_file, task_dict, task_id):
        task = task_dict[task_id]
        book = None
        temp = None
        try:
            path = Path(input_file)
            validate_material_name(task.get('input_name', path.name))
            if path.suffix.lower() not in ('.xlsx', '.csv'):
                raise ValueError('请上传 XLSX 或 UTF-8 CSV 文件')
            if path.resolve() == Path(output_file).resolve():
                raise ValueError('清洗结果须另存文件，不能覆盖输入文件')
            task.update(progress=10, status='正在读取文件并清洗标签空格...')
            records, total, formulas, found = [], 0, 0, False

            def column_index(headers, sheet):
                nonlocal found
                matches = [index for index, name in enumerate(headers)
                           if name is not None and str(name).strip().casefold() == self.column.casefold()]
                if len(matches) > 1:
                    raise ValueError(f'“{sheet}”存在重复的“{self.column}”列')
                if not matches:
                    return None
                found = True
                return matches[0]

            def clean(value, sheet, row, key):
                nonlocal total
                if not isinstance(value, str) or not value:
                    return value
                total += 1
                result = self.clean_text(value)
                if result != value:
                    if len(result) > 32767 and path.suffix.lower() == '.xlsx':
                        raise ValueError(f'{sheet} 第{row}行超过单元格长度上限')
                    records.append({'sheet': sheet, 'row': row, 'key': '' if key is None else str(key),
                                    'before': value, 'after': result})
                return result

            if path.suffix.lower() == '.xlsx':
                book = load_workbook(path)
                for sheet in book:
                    headers = [cell.value for cell in sheet[1]]
                    index = column_index(headers, sheet.title)
                    if index is None:
                        continue
                    key_index = next((i for i, name in enumerate(headers)
                                      if str(name).strip().casefold() == 'key'), None)
                    for row in range(2, sheet.max_row + 1):
                        cell = sheet.cell(row, index + 1)
                        if cell.data_type == 'f':
                            formulas += 1
                            continue
                        key = sheet.cell(row, key_index + 1).value if key_index is not None else ''
                        result = clean(cell.value, sheet.title, row, key)
                        if result != cell.value:
                            cell.value = result
                            cell.data_type = 's'
                if not found:
                    raise ValueError(f'文件中没有“{self.column}”列，请核对清洗列')
                for sheet in book:
                    for row in sheet:
                        for cell in row:
                            if cell.data_type == 'inlineStr' and cell.value is None:
                                cell.value = ''
                KeyClassifier.save_workbook(book, output_file)
            else:
                with path.open(encoding='utf-8-sig', newline='') as file:
                    rows = list(csv.reader(file, strict=True))
                if not rows:
                    raise ValueError('CSV 文件为空')
                index = column_index(rows[0], path.stem)
                if index is None:
                    raise ValueError(f'文件中没有“{self.column}”列，请核对清洗列')
                key_index = next((i for i, name in enumerate(rows[0]) if name.strip().casefold() == 'key'), None)
                for number, row in enumerate(rows[1:], 2):
                    if index < len(row):
                        key = row[key_index] if key_index is not None and key_index < len(row) else ''
                        row[index] = clean(row[index], path.stem, number, key)
                temp = f'{output_file}.{uuid.uuid4().hex}.tmp'
                with open(temp, 'w', encoding='utf-8-sig', newline='') as file:
                    csv.writer(file).writerows(rows)
                os.replace(temp, output_file)
            task.update(progress=100, state='completed', result_file=str(output_file),
                        records=records, column=self.column, total=total, changed=len(records),
                        skipped_formulas=formulas, tag_clean_policy=TAG_CLEAN_POLICY,
                        status=f'清洗完成：检查 {total} 个文本单元格，修改 {len(records)} 个，跳过 {formulas} 个公式单元格')
        except UnicodeDecodeError:
            task.update(state='error', status='清洗失败：CSV 须为 UTF-8 编码')
        except Exception as error:
            task.update(state='error', status=f'清洗失败：{error}')
        finally:
            if book:
                book.close()
            if temp and os.path.exists(temp):
                os.remove(temp)
