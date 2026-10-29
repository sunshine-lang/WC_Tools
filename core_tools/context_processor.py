import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from io import BytesIO
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image as PillowImage
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .key_classifier import CATEGORIES, EMOTIONS, PENDING, KeyClassifier, CLASSIFICATION_POLICY
from .layout_segmenter import LAYOUT_POLICY, image_anchors, segment_sheet
from .dialogue_context import DIALOGUE_CONTEXT_POLICY, DialogueContextAnalyzer
from .language import language_pair
from .vision_context import VisionContextAnalyzer, VISION_POLICY


MATERIAL_DIR = Path(__file__).resolve().parents[2]
DEFAULT_REFERENCE = MATERIAL_DIR / 'WC_UI文案翻译-1016.xlsx'
DEFAULT_GUIDE = MATERIAL_DIR / '柠檬微趣_WC指导材料.xlsx'
CONTEXT_POLICY = 'wc-evidence-vision-context-v4'
STATUSES = ('已关联', '缺资料', '待确认')
DIALOGUE_KEY = re.compile(r'mdnew_(.+)_mdnew_decoration_(intro|\d+)_(\d+)', re.IGNORECASE)
STORY_KEY = re.compile(r'(story\d+(?:_\d+)?_(?:begin|end|task)|fire\d+)_(\d+)', re.IGNORECASE)
FORESTCONCERT_CLUE_WORDS = {'en': ('treble', 'treble clef'), 'zh-CN': ('高音谱号', '高音谱号')}


@dataclass
class ContextRecord:
    record_id: str
    category: str
    activity_id: str = ''
    group_id: str = ''
    item_type: str = ''
    related_record_ids: list = field(default_factory=list)
    image_evidence: list = field(default_factory=list)
    context_status: str = 'pending'
    context_evidence: list = field(default_factory=list)
    relations: list = field(default_factory=list)
    uncertainties: list = field(default_factory=list)


def text(value):
    return '' if value is None else str(value)


def dialogue_identity(key, speaker, note):
    """兼容客户把完整对白 Key 放在 note、说话人放在 key 的表格。"""
    key, speaker, note = (text(value).strip() for value in (key, speaker, note))
    matches = lambda value: DIALOGUE_KEY.fullmatch(value) or STORY_KEY.fullmatch(value)
    if not matches(key) and matches(note):
        if not key or key.casefold() == speaker.casefold() or (
                not speaker and re.fullmatch(r"[A-Za-z][A-Za-z .'-]{0,60}", key)):
            return note, speaker or key
    return key, speaker


def validate_material_name(name):
    normalized = text(name).replace('\\', '/').rsplit('/', 1)[-1]
    if re.search(r'done(?:\s*\(\d+\))?$', Path(normalized).stem, re.IGNORECASE):
        raise ValueError('不能使用 done 结尾的材料，请选择原始待译文件或原始指导材料')


class ContextProcessor:
    def __init__(self, reference_path=DEFAULT_REFERENCE, guide_path=DEFAULT_GUIDE, materials_dir=MATERIAL_DIR, language=None):
        self.language = language_pair(language)
        self.reference_path = Path(reference_path)
        self.guide_path = Path(guide_path)
        self.materials_dir = Path(materials_dir).resolve()

    def resolve_material(self, name):
        validate_material_name(name)
        candidate = (self.materials_dir / name).resolve()
        validate_material_name(candidate.name)
        if candidate.parent != self.materials_dir or candidate.suffix.lower() != '.xlsx' or candidate.name.startswith(('.', '~$')):
            raise ValueError('参考材料只能选择 WC-UI 目录中的 XLSX 文件，不可使用工具缓存或外部目录')
        return candidate

    def material_options(self):
        result = []
        for path in sorted(self.materials_dir.glob('*.xlsx')):
            try:
                self.resolve_material(path.name)
            except ValueError:
                continue
            result.append(path.name)
        return result

    def material_status(self):
        return [{'name': path.name, 'ready': path.is_file(), 'kind': kind}
                for kind, path in [('本目录待译材料', self.reference_path), ('人设与背景', self.guide_path)]]

    @staticmethod
    def read_records(workbook, source_name, classified=False, language=None):
        pair = language_pair(language)
        records = []
        for sheet in workbook:
            headers = pair.headers(c.value for c in sheet[1])
            en_col = pair.column(headers, 'source')
            if not en_col:
                continue
            if classified and '分类' not in headers:
                raise ValueError(f'“{sheet.title}”缺少“分类”列，请先完成第一步')
            if classified and 'context' in headers:
                raise ValueError('文件已有 Context 列，请上传第一步的分类结果')

            def value(row, *names):
                col = next((headers[name] for name in names if name in headers), None)
                return text(sheet.cell(row, col).value) if col else ''

            width = max(c.column for row in sheet for c in row if c.value is not None)
            for row in range(2, sheet.max_row + 1):
                get = lambda *names: value(row, *names)
                source_text = pair.value(get, 'source')
                if not source_text and not (classified and value(row, '分类')):
                    continue
                key = pair.value(get, 'key').strip()
                speaker = pair.value(get, 'speaker').strip()
                emotion = pair.emotion(get).strip()
                legacy_columns = pair.source_rules['legacy_emotion_columns']
                chinese = value(row, *legacy_columns) if legacy_columns else ''
                note = pair.value(get, 'note')
                if not speaker and chinese.strip().casefold() in EMOTIONS and re.fullmatch(r"[A-Za-z][A-Za-z .'-]{0,60}", key):
                    speaker, emotion, chinese = key, chinese.strip(), ''
                    # 原始混合需求表把对白 Key 放在 note 列，工作副本中则没有。
                    key = note.strip() if re.match(r'^(mdnew_|story|Fire)', note.strip()) else ''
                    if key:
                        note = ''
                key, speaker = dialogue_identity(key, speaker, note)
                if not emotion and chinese.strip().casefold() in EMOTIONS and speaker:
                    emotion = chinese.strip()
                bubble = pair.value(get, 'bubble').strip()
                if re.fullmatch(r'\d+\.0+', bubble):
                    bubble = bubble.split('.')[0]
                records.append({
                    'id': f'{sheet.title}:{row}', 'sheet': sheet.title, 'row': row, 'width': width,
                    'source': f'{source_name} / {sheet.title}!{row}', 'key': key,
                    'source_text': source_text, 'speaker': speaker, 'emotion': emotion,
                    'bubble': bubble, 'chinese': chinese,
                    'note': note.strip(), 'category': value(row, '分类').strip(),
                    'formula': sheet.cell(row, en_col).data_type == 'f',
                })
        return records

    @staticmethod
    def relation(record, activity_hint=''):
        key = record['key'].casefold()
        card = re.fullmatch(r'a(.+)_set(\d+)(?:_(\d+))?_name', key)
        decor = re.fullmatch(r'mdnew_(.+)_mdnew_decoration_name_(\d+)', key)
        dialogue = DIALOGUE_KEY.fullmatch(key)
        story = STORY_KEY.fullmatch(key)
        if card:
            activity, number, item = card.groups()
            return activity, f'{activity}/set/{int(number)}', f'卡册{int(number)}', int(item or 0), ''
        if decor:
            activity, number = decor.groups()
            return activity, f'{activity}/decor/{int(number)}', f'装饰物{int(number)}', int(number), str(int(number))
        if dialogue:
            activity, number, line = dialogue.groups()
            label = '活动开场' if number == 'intro' else f'装饰物{int(number)}对白'
            return activity, f'{activity}/dialogue/{number}', label, int(line), number
        if story:
            prefix, line = story.groups()
            activity = prefix.split('_')[0]
            return activity, f'{activity}/dialogue/{prefix}', prefix, int(line), prefix
        activity = ''
        if key.startswith(('cat_multiboards_', 'catmultiboards_')) or key == 'offer_cat':
            activity = 'cat'
        elif key.startswith('midautumn_multiboards_') or key == 'offer_midautumn':
            activity = 'midautumn'
        elif key.startswith('album_'):
            activity = key.split('_')[1]
        elif key.startswith('event_') and '_mdnew_' in key:
            activity = key.split('_mdnew_')[0][6:]
        elif key.startswith('chain_gift_title_chain_'):
            activity = key[len('chain_gift_title_chain_'):]
        elif key.startswith('forestconcert_'):
            activity = 'forestconcert'
        if not key:
            activity = activity_hint
        if not activity or record.get('category') == CATEGORIES[3]:
            return '', '', '', 0, ''
        module = '商店宣传' if not key else '活动UI'
        if re.match(r'^forestconcert_(road|rank)_', key):
            module = '进度奖励' if '_road_' in key else '排行榜'
        return activity, f'{activity}/ui/{module}', module, record['row'], ''

    def prepare_relations(self, records):
        for record in records:
            activity, group, label, order, number = self.relation(record)
            evidence = ['key_pattern'] if group else []
            if record.get('manual_category'):
                evidence.append('manual_override')
            expected = CATEGORIES[1] if '/set/' in group else CATEGORIES[2] if '/decor/' in group else CATEGORIES[3] if '/dialogue/' in group else CATEGORIES[0]
            if group and expected != record.get('category') and record.get('manual_category'):
                # A corrected category does not prove the old object's identity.
                group, label, number = '', '', ''
                if record['category'] == CATEGORIES[0]:
                    group, label = f'{activity}/ui/manual', '人工确认的活动文案'
            record.update(activity=activity, group=group, label=label, order=order, number=number)
            record['relation_evidence'] = evidence
        for record in records:
            if record['group'] or record['key'] or record['category'] == CATEGORIES[3]:
                continue
            region = record.get('region_id')
            anchors = [r for r in records if region and r.get('region_id') == region and r['sheet'] == record['sheet']
                       and r['group'] and r['key'] and r['category'] == record['category']
                       and r.get('classification_status', 'confirmed') == 'confirmed']
            activities = {r['activity'] for r in anchors}
            if len(anchors) < 2 or len(activities) != 1 or not record.get('region_formatted'):
                continue
            activity = next(iter(activities))
            groups = {r['group'] for r in anchors}
            if record['category'] == CATEGORIES[0]:
                group, label = f'{activity}/ui/region/{region}', '区块内活动文案（具体功能待核对）'
            elif record['category'] == CATEGORIES[1] and len(groups) == 1:
                group, label = next(iter(groups)), '套组内名称（卡片序号未提供）'
            else:
                continue  # A region alone never identifies an individual decoration.
            record.update(activity=activity, group=group, label=label,
                          relation_evidence=[*record['relation_evidence'], 'confirmed_layout_region'],
                          region_anchor_ids=[r['id'] for r in anchors])

    def attach_classification(self, workbook, records, metadata):
        by_id = {r['id']: r for r in (metadata or {}).get('records', [])}
        for sheet in workbook:
            headers = self.language.headers(c.value for c in sheet[1])
            cols = [self.language.column(headers, name) for name in ('key', 'source', 'target', 'note')]
            layout = (metadata or {}).get('layouts', {}).get(sheet.title) or segment_sheet(sheet, cols)
            for r in [r for r in records if r['sheet'] == sheet.title]:
                meta = by_id.get(r['id'], {})
                if metadata and (not meta or meta['category'] != r['category'] or meta['source_text'] != r['source_text']):
                    raise ValueError('分类元数据与工作簿不一致，请重跑分类')
                region = meta.get('region_id') or layout['rows'].get(r['row'], {}).get('region_id', '')
                r.update(region_id=region, region_formatted=layout['regions'].get(region, {}).get('formatted', False),
                         classification_status=meta.get('classification_status', 'confirmed'),
                         manual_category=meta.get('manual_category'), classification_evidence=meta.get('evidence', []))

    @staticmethod
    def anchored_images(workbook, records):
        images, warnings = [], []
        for sheet in workbook:
            for anchor in image_anchors(sheet):
                covered = [r for r in records if r['sheet'] == sheet.title and anchor['extent_known']
                           and anchor['start_row'] <= r['row'] <= anchor['end_row']]
                valid = covered and all(r['group'] and r.get('classification_status') == 'confirmed' for r in covered)
                activities = {r['activity'] for r in covered}
                groups = {r['group'] for r in covered}
                scope, group, activity = 'unresolved', '', ''
                if valid and len(activities) == 1:
                    activity = next(iter(activities))
                    if len(groups) == 1 and '/set/' in next(iter(groups)):
                        scope, group = 'set', next(iter(groups))
                    elif len(covered) == 1 and '/decor/' in covered[0]['group']:
                        scope, group = 'decoration', covered[0]['group']
                    elif all('/ui/' in r['group'] for r in covered):
                        scope = 'activity'
                item = {**anchor, 'scope': scope, 'activity': activity, 'group': group,
                        'record_ids': [r['id'] for r in covered], 'evidence': 'image_anchor'}
                images.append(item)
                if scope == 'unresolved':
                    warnings.append(anchor['image_id'] + '：图片覆盖范围或对象归属不唯一')
        return images, warnings

    @staticmethod
    def load_roles(path, source_name=None):
        if not path.is_file():
            return {}
        workbook = load_workbook(path, data_only=True)
        roles = {}
        if '人设表' in workbook:
            sheet = workbook['人设表']
            for row in range(3, sheet.max_row + 1):
                names = text(sheet.cell(row, 4).value).splitlines()
                if not names:
                    continue
                description = text(sheet.cell(row, 7).value)
                # 仅引用开头两句人物事实，避免混入后续剧透和其他语种列。
                facts = ''.join(re.split(r'(?<=[。！？])', description)[:2]).strip()
                age = text(sheet.cell(row, 5).value).strip()
                role = {'facts': (f'年龄：{age}；' if age and age != '/' else '') + facts,
                        'source': f'{source_name or path.name} / 人设表!G{row}'}
                for name in names:
                    if name.strip():
                        roles[name.strip().casefold()] = role
        workbook.close()
        return roles

    @staticmethod
    def find_reference(record, by_key, by_english):
        if record['key']:
            candidates = by_key.get(record['key'], [])
            if candidates:
                matching = [r for r in candidates if r['source_text'].strip() == record['source_text'].strip()]
                if not matching:
                    return None, '待确认', '相同 Key 的参考源文与本条不一致，请核对版本'
                candidates = matching
            else:
                return None, '缺资料', '参考文件未找到此 Key'
        else:
            candidates = by_english.get(record['source_text'].strip(), [])
        if record['speaker']:
            candidates = [r for r in candidates if r['speaker'] == record['speaker']
                          and (not record['emotion'] or r['emotion'] == record['emotion'])]
        if len(candidates) != 1:
            return None, '待确认' if candidates else '缺资料', (
                f'找到 {len(candidates)} 条参考候选，无法唯一定位' if candidates else '未找到源文及角色一致的参考记录')
        return candidates[0], '已关联', ''

    @staticmethod
    def extract_images(workbook, references, asset_dir, source_name):
        asset_dir.mkdir(parents=True, exist_ok=True)
        anchors, warnings = ContextProcessor.anchored_images(workbook, references)
        by_id = {a['image_id']: a for a in anchors}
        images = []
        for sheet in workbook:
            for index, image in enumerate(sheet._images):
                anchor = by_id.get(f'{sheet.title}:image:{index + 1}')
                if not anchor:
                    continue
                data = image._data()
                image.ref = BytesIO(data)
                with PillowImage.open(BytesIO(data)) as picture:
                    extension = picture.format.lower()
                if extension not in ('png', 'jpeg', 'jpg', 'gif', 'bmp'):
                    warnings.append(f'{sheet.title} 图片{index + 1}格式暂不支持')
                    continue
                name = hashlib.sha256(data).hexdigest()[:16] + '.' + extension
                (asset_dir / name).write_bytes(data)
                caption = f'{source_name} / {sheet.title} 图片{index + 1}（锚点第{anchor["start_row"]}行）'
                images.append({**anchor, 'name': name, 'path': f'images/{name}', 'caption': caption,
                               'source_file': source_name})
        return images, warnings

    @staticmethod
    def ui_unit(record):
        key = record['key'].casefold().replace('catmultiboards_', 'cat_multiboards_')
        if '_book_' in key:
            return key.split('_book_', 1)[0] + '/book'
        if '_guide_' in key:
            return key.split('_guide_', 1)[0] + '/intro'
        if key.endswith('_pre_desc'):
            return key.removesuffix('_pre_desc') + '/intro'
        if key.endswith('_title') and (key.removesuffix('_title').endswith('_multiboards') or key.startswith('album_')):
            return key.removesuffix('_title') + '/intro'
        rank = re.fullmatch(r'(.+_rank)_(title|score_need|start_desc)', key)
        if rank:
            return rank.group(1) + '/entry'
        return re.sub(r'_(?:title|desc\d*)$', '', key)

    def related_records(self, record, groups, activities):
        peers = [r for r in groups.get(record['group'], []) if r['sheet'] == record['sheet']]
        activity_rows = [r for r in activities.get(record['activity'], []) if r['sheet'] == record['sheet']]
        overview = [r for r in activity_rows if '/ui/' in r['group']
                    and r['key'].casefold().endswith(('_pre_desc', '_prediction_desc'))]
        if not record['key'] and record['category'] == CATEGORIES[0]:
            selected = [record, *(r for r in activity_rows if r['id'] in record.get('region_anchor_ids', []))]
        elif '/set/' in record['group'] or '/dialogue/' in record['group']:
            selected = list(peers)
        elif '/decor/' in record['group']:
            dialogue = groups.get(f"{record['activity']}/dialogue/{record['number']}", [])
            selected = [*peers, *dialogue]
            if not dialogue and overview:
                intro = overview[0]
                title = intro['key'].removesuffix('_pre_desc').removesuffix('_desc') + '_title'
                selected.extend(r for r in activity_rows if r['key'] == title)
                selected.append(intro)
        else:
            selected = [r for r in activity_rows if '/ui/' in r['group'] and self.ui_unit(r) == self.ui_unit(record)]
            if len(selected) < 2:
                titles = [r for r in peers if r['key'].casefold().endswith('_title')
                          and not any(part in r['key'].casefold() for part in ('_complete_', '_success_', '_fail_', '_book_'))]
                anchors = overview if record['key'].casefold().startswith(('offer_', 'chain_gift_')) else titles + overview
                anchor = next((r for r in anchors if r['id'] != record['id'] and r['source_text'] != record['source_text']), None)
                if anchor:
                    selected.append(anchor)
        selected = selected or [record]
        if '/dialogue/' in record['group']:
            selected.extend(groups.get(f"{record['activity']}/decor/{record['number']}", []))
        unique = {r['id']: r for r in selected}
        return [r for r in unique.values() if r['sheet'] == record['sheet']]

    @staticmethod
    def line(record):
        label = record['key'] or record['speaker'] or '无 Key'
        details = [record['speaker'], record['emotion']]
        if record['bubble']:
            details.append('思考' if record['bubble'] == '2' else '正常对话' if record['bubble'] == '1' else '气泡' + record['bubble'])
        details = [item for item in details if item]
        return label + (' (' + ' / '.join(details) + ')' if details else '') + ': ' + record['source_text']

    def analyze_images(self, workbook, reference_workbook, records, images, config_path, cache_dir, task):
        jobs, aliases = {}, defaultdict(list)
        for item in images:
            if item['scope'] == 'unresolved':
                continue
            applicable = [r for r in records if r['category'] != CATEGORIES[3]
                          and r.get('classification_status') == 'confirmed'
                          and (r['sheet'] == item['sheet'] or item.get('source_file'))
                          and ((item['group'] and item['group'] == r['group']) or
                               (item['scope'] == 'activity' and r['category'] == CATEGORIES[0]
                                and item['activity'] == r['activity']))]
            if not applicable:
                continue
            source = reference_workbook if item.get('source_file') else workbook
            picture = source[item['sheet']]._images[int(item['image_id'].rsplit(':', 1)[1]) - 1]
            data = picture._data()
            picture.ref = BytesIO(data)
            identity = (hashlib.sha256(data).hexdigest(), item['scope'])
            jobs[identity] = (data, item['scope'])
            aliases[identity].append(item)
        if not jobs:
            return
        analyzer = VisionContextAnalyzer(config_path, cache_dir)
        results, errors = {}, []
        with ThreadPoolExecutor(max_workers=analyzer.max_workers) as pool:
            futures = {pool.submit(analyzer.analyze, *args): identity for identity, args in jobs.items()}
            for future in as_completed(futures):
                identity = futures[future]
                try:
                    results[identity] = future.result()
                    for item in aliases[identity]:
                        item['visual_analysis'] = results[identity]
                except Exception as error:
                    errors.append(f'{aliases[identity][0]["image_id"]}：{error}')
                task.update(progress=30 + int(15 * (len(results) + len(errors)) / len(jobs)),
                            status=f'正在识别图片视觉语义 {len(results) + len(errors)}/{len(jobs)}（并发 {analyzer.max_workers}）')
        task.update(vision_images=[item for items in aliases.values() for item in items],
                    vision_cached_images=sum(r['cache_hit'] for r in results.values()),
                    vision_model_requests=sum(not r['cache_hit'] for r in results.values()))
        if errors:
            raise ValueError('图片视觉分析失败；成功结果已缓存，请重试本步。' + '；'.join(errors))

    def analyze_scenes(self, prepared, records, matches, config_path, materials, cache_dir, metadata, task):
        jobs, blocked, results = {}, {}, {}
        for record, (status, _, _, effective) in prepared:
            group = effective.get('group', '')
            scene_id = f'{record["sheet"]}:{group}'
            if record['category'] != CATEGORIES[3] or not group or scene_id in jobs or scene_id in blocked:
                continue
            peers = sorted((r for r in records if r['group'] == group and r['sheet'] == record['sheet']), key=lambda r:r['order'])
            if any(r['formula'] or not r['source_text'].strip() or r['category'] != CATEGORIES[3]
                   or r.get('classification_status') != 'confirmed' or matches[r['id']][1] == '待确认' for r in peers):
                blocked[scene_id] = '同段存在异常源文、分类或版本冲突，场景分析留空，请核对后重跑'
            elif len({r['key'].casefold() for r in peers}) != len(peers) or len({r['order'] for r in peers}) != len(peers):
                blocked[scene_id] = '同段完整 Key 重复，无法确定对白顺序，场景分析留空'
            elif status != '待确认':
                warnings = []
                if any(not r['bubble'] for r in peers):
                    warnings.append('本组存在未提供气泡类型的台词，交流媒介或内心独白只能依据明确原文判断')
                orders = [r['order'] for r in peers]
                if orders[0] > 1 or any(b - a != 1 for a, b in zip(orders, orders[1:])):
                    warnings.append('本批对白句序存在缺号，仅分析已提供原文，不补写缺句')
                refs = [r for r in effective.get('related_records', []) if r['group'] != group]
                jobs[scene_id] = (group, peers, refs, warnings, {'relation_policy': CONTEXT_POLICY,
                    'sheet': record['sheet'], 'classification': (metadata or {}).get('workbook_sha256')})
        if not jobs:
            return results, blocked
        analyzer = DialogueContextAnalyzer(config_path, materials=materials, cache_dir=cache_dir, language=self.language)
        errors = []
        with ThreadPoolExecutor(max_workers=analyzer.max_workers) as pool:
            futures = {pool.submit(analyzer.analyze, *args): identity for identity, args in jobs.items()}
            for future in as_completed(futures):
                identity = futures[future]
                try:
                    results[identity] = future.result()
                except Exception as error:
                    errors.append(f'{identity}：{error}')
                task.update(progress=45 + int(20 * (len(results) + len(errors)) / len(jobs)),
                            status=f'正在通读对白场景 {len(results) + len(errors)}/{len(jobs)}（并发 {analyzer.max_workers}）')
        task['dialogue_scenes'] = [results[key] for key in jobs if key in results]
        if errors:
            raise ValueError('对白场景分析失败；成功结果已缓存，请重试本步。' + '；'.join(errors))
        return {key: results[key] for key in jobs}, blocked

    def bounded_context(self, record, selected, activity_rows, images, warnings):
        key, group = record['key'].casefold(), record['group']
        if record['category'] == CATEGORIES[1]:
            kind = 'set_name' if re.fullmatch(r'a.+_set\d+_name', key) else 'card_name' if key else 'set_member_unknown'
        elif record['category'] == CATEGORIES[2]:
            kind = 'decoration_name'
        elif '_guide_' in key:
            kind = 'guide'
        elif key.startswith(('offer_', 'chain_gift_')):
            kind = 'shop_copy'
        elif '_book_' in key:
            kind = 'collection_label'
        elif key.endswith('_title'):
            kind = 'title'
        elif any(word in key for word in ('_complete_', '_success_', '_fail_', '_final_')):
            kind = 'result_message'
        elif key.endswith(('_pre_desc', '_prediction_desc')):
            kind = 'activity_intro'
        else:
            kind = 'description' if key else 'activity_copy_unknown'
        names = dict(set_name='套组名称', card_name='卡片名称', set_member_unknown='套组内名称，卡片序号未提供',
                     decoration_name='装饰物名称', guide='操作引导', shop_copy='商店或礼包文案',
                     collection_label='图鉴名称或分类', title='标题', result_message='结果或状态提示',
                     activity_intro='活动介绍', description='活动说明', activity_copy_unknown='活动文案，具体功能未确认')
        background = [r for r in activity_rows if r['sheet'] == record['sheet'] and '/ui/' in r['group']
                      and (r['key'].casefold().endswith(('_pre_desc', '_prediction_desc')) or
                           re.fullmatch(r'(?:album_.+|.+_multiboards)_title', r['key'].casefold()))]
        image_refs = [im for im in images if (im.get('sheet') == record['sheet'] or im.get('source_file')) and
                      ((im.get('group') == group and im.get('scope') in ('set', 'decoration')) or
                       (record['category'] == CATEGORIES[0] and im.get('scope') == 'activity' and
                        im.get('activity') == record['activity']))]
        if not record['key']:
            warnings.append('缺少 Key；归属由已确认区块支持，未推断卡片或物件编号')
        if record['category'] in CATEGORIES[1:3] and not image_refs:
            warnings.append('未找到可唯一归属的图片；不能据此推断外观')
        ambiguous = [im for im in images if im.get('scope') == 'unresolved' and record['id'] in im.get('record_ids', [])]
        if ambiguous:
            warnings.append('存在覆盖当前行但归属不唯一的图片，未作为对象证据')
        refs = {r['id']: r for r in [*background, *selected] if r['id'] != record['id']}
        relations = []
        for r in refs.values():
            reason = ('same_set' if '/set/' in group and r['group'] == group else
                      'same_object_dialogue' if '/decor/' in group and '/dialogue/' in r['group'] else
                      'same_ui_function' if r in selected and record['category'] == CATEGORIES[0] else 'activity_key')
            relations.append({'record_id': r['id'], 'evidence': reason, 'key': r['key']})
        evidence = list(dict.fromkeys([*record.get('relation_evidence', []),
                                      *(['image_anchor'] if image_refs else [])]))
        contract = ContextRecord(record['id'], record['category'], record['activity'], group, kind,
                                 list(refs), image_refs, 'partial' if warnings else 'linked', evidence,
                                 relations, list(dict.fromkeys(warnings)))
        lines = ['【活动背景】', f'活动标识：{record["activity"]}（依据 Key 或已确认区块）']
        lines.extend(self.line(r) for r in background)
        lines.extend(['【当前文本功能】', names[kind], '【关联对象】'])
        objects = [self.line(r) for r in selected if r['id'] not in {b['id'] for b in background}]
        lines.extend(objects or ['相关记录已列在活动背景中，未提供其他对象信息。'])
        if record['category'] == CATEGORIES[1]:
            lines.append('同套名称用于理解主题和命名区分，不是当前记录的翻译答案。')
        elif record['category'] == CATEGORIES[2]:
            lines.append('对应对白仅作背景；以当前名称源意为准，不将角色猜测当作物件事实。')
        if image_refs:
            lines.append('【图片证据】')
            for im in image_refs:
                lines.append(f'{im["image_id"]}，锚点行 {im["start_row"]}–{im["end_row"]}，'
                    + {'set': '套组级参考，未定位单张卡片', 'activity': '活动级参考，未定位单个物件',
                       'decoration': '锚点唯一对应本物件'}[im['scope']]
                    + ('；配图：' + im['path'] if im.get('path') else '；图片在原工作表中'))
                visual = im.get('visual_analysis')
                if visual:
                    lines.append('视觉观察（模型识别，需核对）：' + visual['summary'])
                    lines.extend(f'- {item["location"]}：{item["description"]}' for item in visual['objects'])
                    if visual['visible_text']:
                        lines.append('图中可辨文字（不覆盖待译源文）：' + '；'.join(visual['visible_text']))
                    warnings.extend(im['image_id'] + '：' + w for w in visual['uncertainties'])
                else:
                    lines.append('尚未进行视觉语义分析。')
            lines.append('仅使用上述归属范围内的视觉线索；不把活动图或套组拼图的整体特征分配给单个物件。')
        contract.uncertainties = list(dict.fromkeys(warnings))
        contract.context_status = 'partial' if warnings else 'linked'
        if warnings:
            lines.extend(['【不确定项】', *contract.uncertainties])
        record.update(context_record=asdict(contract), related_count=len(refs),
                      related_records=list(refs.values()), warning='；'.join(contract.uncertainties))
        return '缺资料' if warnings else '已关联', '\n'.join(lines), image_refs, record

    def build_context(self, record, reference, status, warning, groups, activities, roles, images, enhanced=False):
        effective = dict(reference or record)
        effective['source_text'] = record['source_text']
        warnings = [warning] if warning else []
        if not reference:
            self.prepare_relations([effective])
        if record['category'] == PENDING or record['category'] not in CATEGORIES:
            status, warnings = '待确认', ['分类尚未确认，请先核对第一步结果']
        if effective.get('classification_status') in ('pending', 'inferred'):
            status, warnings = '待确认', ['分类证据尚未人工确认，请先完成分类确认']
        if record['formula'] or not record['source_text'].strip():
            status, warnings = '待确认', ['源文为空或为公式，请先核对源文']
        if not effective.get('group'):
            message = ('缺少完整对白 Key，不能按说话人或行号推断分段' if record['speaker']
                       else '无 Key 或活动归属未明确，无法可靠分组')
            status = '待确认'
            warnings = warnings or [message]
        else:
            group = effective['group']
            expected = CATEGORIES[1] if '/set/' in group else CATEGORIES[2] if '/decor/' in group else CATEGORIES[3] if '/dialogue/' in group else CATEGORIES[0]
            if record['category'] in CATEGORIES and record['category'] != expected:
                status, warnings = '待确认', ['分类与参考 Key 结构不一致，请先核对分类']
        effective['warning'] = '；'.join(warnings)
        if status == '待确认':
            return status, '', [], effective

        selected = self.related_records(effective, groups, activities)
        if record['category'] != CATEGORIES[3]:
            return self.bounded_context(effective, selected, activities.get(effective['activity'], []), images, warnings)
        if effective['group'] == 'forestconcert/dialogue/12':
            # 只引用上传文件内已核对的前段线索，不用模型补写。
            prior = [r for r in activities.get(effective['activity'], [])
                     if r['sheet'] == record['sheet'] and r['key'] == 'mdnew_forestconcert_mdnew_decoration_5_10']
            mention, keyword = FORESTCONCERT_CLUE_WORDS.get(self.language.source_lang, ('', ''))
            if mention and any(mention in r['source_text'].casefold() for r in selected) and prior and keyword in prior[0]['source_text'].casefold():
                selected.extend(prior)
            else:
                warnings.append('未找到或无法验证第5段的高音谱号解释')
        selected = list({r['id']: r for r in selected}.values())
        selected_images = []  # Activity artwork does not prove a dialogue scene's medium or setting.
        role_lines = []
        if enhanced:
            speakers = list(dict.fromkeys(r['speaker'] for r in selected if r['speaker']))
            for speaker in speakers:
                role = roles.get(speaker.casefold())
                if role and role['facts']:
                    role_lines.append(speaker + ': ' + role['facts'])
                else:
                    warnings.append(f'未找到 {speaker} 的人设')
            if effective['speaker'] and not effective['bubble']:
                warnings.append('原始材料未提供气泡类型，不能推断')
            if record['category'] in (CATEGORIES[1], CATEGORIES[2]) and not selected_images:
                warnings.append('原始材料未提供可确定归属的图片')
        useful = len(selected) > 1 or role_lines or selected_images
        if not useful:
            warnings.append('未找到可用关联原文，Context 留空')
        if warnings:
            status = '缺资料'
        lines = ['[关联语境参考]', *(self.line(r) for r in selected)] if useful else []
        if role_lines:
            lines.extend(['[人物参考]', *role_lines])
        if selected_images:
            lines.append('[配图] ' + '、'.join(dict.fromkeys(im['path'] for im in selected_images)))
        effective['warning'] = '；'.join(dict.fromkeys(warnings))
        effective['related_count'] = len(selected) if useful else 0
        effective['related_records'] = selected
        return status, '\n'.join(lines), selected_images, effective

    def process(self, input_file, output_file, task_dict, task_id, reference_path=None, guide_path=None, enhance=False,
                analyze_dialogue=False, config_path=None, materials=None, scene_cache_dir=None,
                classification_path=None, rules_path=None, analyze_images=False, vision_cache_dir=None):
        task = task_dict[task_id]
        workbook = reference_workbook = None
        try:
            task['status'] = '正在读取分类结果...'
            reference_path = Path(reference_path or self.reference_path)
            guide_path = Path(guide_path or self.guide_path)
            for path, name in [(Path(input_file), task.get('input_name', '')),
                               (reference_path, task.get('reference_name', '')),
                               (guide_path, task.get('guide_name', ''))]:
                validate_material_name(path.name)
                validate_material_name(path.resolve().name)
                validate_material_name(name)
            if enhance:
                reference_path = self.resolve_material(str(reference_path))
                guide_path = self.resolve_material(str(guide_path))
            workbook = load_workbook(input_file)
            records = self.read_records(workbook, task.get('input_name', Path(input_file).name), classified=True, language=self.language)
            if not records:
                raise ValueError('未找到待处理记录，须包含源文列和“分类”列')
            metadata = None
            if classification_path and not Path(classification_path).is_file():
                raise ValueError('缺少分类证据元数据，请重跑分类节点')
            if classification_path and Path(classification_path).exists():
                metadata = json.loads(Path(classification_path).read_text('utf-8'))
                if (metadata.get('workbook_sha256') != hashlib.sha256(Path(input_file).read_bytes()).hexdigest()
                    or metadata.get('classification_policy') != CLASSIFICATION_POLICY
                    or any(layout.get('policy') != LAYOUT_POLICY for layout in metadata.get('layouts', {}).values())
                    or metadata.get('language_rules') != self.language.cache_identity()
                    or (rules_path and metadata.get('rules_sha256') != hashlib.sha256(Path(rules_path).read_bytes()).hexdigest())):
                    raise ValueError('分类、源文件或规则已变化，请重跑分类后再生成 Context')
            self.attach_classification(workbook, records, metadata)
            references, roles = [], {}
            if enhance:
                if not reference_path.is_file():
                    raise ValueError('所选本目录材料不存在，请重新选择或关闭可选补充')
                reference_workbook = load_workbook(reference_path, data_only=True)
                references = self.read_records(reference_workbook, task.get('reference_name', reference_path.name), language=self.language)
                if not references:
                    raise ValueError('原始参考未找到 当前语向对应的源文列')
                roles = self.load_roles(guide_path, task.get('guide_name'))
            by_key, by_english = defaultdict(list), defaultdict(list)
            for r in references:
                if r['key']:
                    by_key[r['key']].append(r)
                by_english[r['source_text'].strip()].append(r)
            effective_records, matches = [], {}
            for record in records:
                effective = dict(record)
                status, warning = '已关联', ''
                if enhance:
                    match, status, warning = self.find_reference(record, by_key, by_english)
                    if match:
                        effective.update({name: match[name] for name in ('key', 'speaker', 'emotion', 'bubble', 'note')})
                        if not effective['chinese']:
                            effective['chinese'] = match['chinese']
                        effective['reference_source'] = match['source']
                effective_records.append(effective)
                matches[record['id']] = (effective, status, warning)
            self.prepare_relations(effective_records)
            groups, activities = defaultdict(list), defaultdict(list)
            for r in effective_records:
                status = matches[r['id']][1]
                if (r['group'] and status != '待确认' and r['category'] in CATEGORIES and not r['formula']
                        and r['source_text'].strip() and r.get('classification_status') == 'confirmed'):
                    expected = CATEGORIES[1] if '/set/' in r['group'] else CATEGORIES[2] if '/decor/' in r['group'] else CATEGORIES[3] if '/dialogue/' in r['group'] else CATEGORIES[0]
                    if r['category'] != expected:
                        continue
                    groups[r['group']].append(r)
                    activities[r['activity']].append(r)
            for peers in groups.values():
                peers.sort(key=lambda r: (r['order'], r['row']))
            output = Path(output_file)
            asset_dir = output.with_name(output.stem + '_assets') / 'images'
            images, image_warnings = self.anchored_images(workbook, effective_records)
            extracted_images = []
            if enhance:
                self.prepare_relations(references)
                for r in references:
                    r['classification_status'] = 'confirmed' if r['group'] else 'pending'
                extracted_images, extra_warnings = self.extract_images(reference_workbook, references, asset_dir,
                                                             task.get('reference_name', reference_path.name))
                image_warnings.extend(extra_warnings)
                # Match by validated project identity; reference-sheet row IDs stay explicitly external.
                images.extend(extracted_images)
            task.update(progress=30, status='正在按活动、套组和对白段关联...')
            if analyze_images:
                self.analyze_images(workbook, reference_workbook, effective_records, images, config_path,
                                    vision_cache_dir, task)
            prepared = []
            for record in records:
                effective, status, warning = matches[record['id']]
                prepared.append((record, self.build_context(record, effective, status, warning,
                    groups, activities, roles, images, enhanced=enhance)))
            result = []
            context_columns = {}
            scene_results, blocked_scenes = {}, {}
            if analyze_dialogue:
                scene_results, blocked_scenes = self.analyze_scenes(prepared, effective_records, matches,
                    config_path, materials, scene_cache_dir, metadata, task)
            for record, (status, context, selected, effective) in prepared:
                sheet = workbook[record['sheet']]
                if sheet.title not in context_columns:
                    column = record['width'] + 1
                    context_columns[sheet.title] = column
                    header = sheet.cell(1, column, 'Context')
                    header.font = Font(bold=True, color='FFFFFF')
                    header.fill = PatternFill('solid', fgColor='334155')
                    sheet.column_dimensions[header.column_letter].width = 90
                group = effective.get('group', '')
                scene_id = f'{record["sheet"]}:{group}'
                if analyze_dialogue and record['category'] == CATEGORIES[3] and group and status != '待确认':
                    if scene_id in blocked_scenes:
                        status, context = '待确认', ''
                        effective['warning'] = blocked_scenes[scene_id]
                    else:
                        scene = scene_results[scene_id]
                        context = scene['context']
                        prior_warning = effective.get('warning', '').replace('未找到可用关联原文，Context 留空', '').strip('；')
                        effective['warning'] = '；'.join(filter(None, [prior_warning, *scene['warnings']]))
                        effective['related_count'] = len(scene['sources']) + len(scene['references'])
                        status = '缺资料' if effective['warning'] else '已关联'
                if len(context) > 32767:
                    raise ValueError(f'{sheet.title} 第{record["row"]}行 Context 超过 Excel 单元格上限，请缩小关联范围')
                cell = sheet.cell(record['row'], context_columns[sheet.title], context)
                cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                contract = effective.get('context_record') or asdict(ContextRecord(
                    record['id'], record['category'], effective.get('activity', ''), group,
                    'dialogue' if record['category'] == CATEGORIES[3] else '',
                    [r['id'] for r in effective.get('related_records', []) if r['id'] != record['id']],
                    [], 'pending' if not context else 'partial' if status == '缺资料' else 'linked',
                    effective.get('relation_evidence', []), uncertainties=[effective['warning']] if effective.get('warning') else []))
                result.append({'id': record['id'], 'sheet': record['sheet'], 'row': record['row'],
                               'key': effective['key'], 'speaker': effective['speaker'],
                               'source_text': record['source_text'], 'category': record['category'],
                               'group': effective.get('group', ''), 'status': status, 'context': context,
                               'images': selected, 'warning': effective.get('warning', ''),
                               'source': record['source'], 'reference_source': effective.get('reference_source', ''),
                               'related_count': effective.get('related_count', 0), 'context_record': contract})
                task['progress'] = 65 + int(25 * len(result) / len(records))
            KeyClassifier.save_workbook(workbook, output_file)
            output.with_suffix('.context.json').write_text(json.dumps({'policy': CONTEXT_POLICY,
                'classification_policy': CLASSIFICATION_POLICY, 'layout_policy': LAYOUT_POLICY, 'vision_policy': VISION_POLICY,
                'input_sha256': hashlib.sha256(Path(input_file).read_bytes()).hexdigest(),
                'records': [r['context_record'] for r in result], 'image_warnings': image_warnings},
                ensure_ascii=False, indent=2), 'utf-8')
            bundle_file = None
            if extracted_images:
                bundle_file = str(output.with_suffix('.zip'))
                with ZipFile(bundle_file, 'w', ZIP_DEFLATED) as archive:
                    archive.write(output, task.get('dl_name', 'WC_context.xlsx'))
                    for name in {im['name'] for im in extracted_images}:
                        archive.write(asset_dir / name, f'images/{name}')
                    manifest = '\n'.join(f"{im['path']}\n{im['caption']}\n{im['scope']}\n" for im in extracted_images)
                    archive.writestr('图片来源.txt', manifest)
            counts = Counter(r['status'] for r in result)
            task.update(progress=100, state='completed', status=f'关联完成，共 {len(result)} 条；待确认 {counts["待确认"]} 条，缺资料 {counts["缺资料"]} 条',
                        result_file=output_file, records=result, counts={s: counts[s] for s in STATUSES},
                        bundle_file=bundle_file, image_warnings=image_warnings,
                        asset_paths={im['name']: str(asset_dir / im['name']) for im in extracted_images},
                        reference_name=task.get('reference_name', reference_path.name) if enhance else '未启用可选补充',
                        enhanced=enhance, context_policy=DIALOGUE_CONTEXT_POLICY if analyze_dialogue else CONTEXT_POLICY,
                        dialogue_scenes=list(scene_results.values()),
                        dialogue_cached_scenes=sum(s.get('cache_hit', False) for s in scene_results.values()),
                        dialogue_model_requests=sum(not s.get('cache_hit', False) for s in scene_results.values()))
        except Exception as error:
            task.update(status=f'关联失败：{error}', state='error')
        finally:
            if workbook:
                workbook.close()
            if reference_workbook:
                reference_workbook.close()
