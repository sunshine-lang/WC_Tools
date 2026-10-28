#!/usr/bin/env python3
"""流水线编排：一份待译文件进，六个阶段依次跑，每阶段落一个产物。

模块本身不改——每个阶段都是对现有 `process(input, output, task_dict, task_id)`
的适配。编排层负责：批次目录、产物命名、进度、异常包装、预览统计、清单落盘。

用法：
    python pipeline.py <待译文件.xlsx> [--game wc] [--lang en-de]
                       [--stages classify,context] [--batch 名称]
"""
from __future__ import annotations

import argparse
import json
import sys
import time
import traceback
from core_tools.language import language_pair, language_profile, language_profiles
from dataclasses import asdict, dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parent
WORKSPACE = ROOT / 'workspace'


# ---------------------------------------------------------------- 结果与预览

@dataclass
class StageResult:
    key: str
    title: str
    state: str                 # completed | failed | skipped
    artifact: str | None
    seconds: float
    message: str = ''
    preview: dict = field(default_factory=dict)


def workbook_preview(path: Path, sample_rows: int = 3) -> dict:
    """读产物的表头、行数、每列非空数，以及前几行样本。"""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return {'rows': 0, 'columns': [], 'filled': {}, 'sample': []}

    columns = [str(c).strip() if c is not None else '' for c in header]
    body = [r for r in rows if any(v not in (None, '') for v in r)]

    filled = {}
    for index, name in enumerate(columns):
        if name:
            filled[name] = sum(1 for r in body if index < len(r) and r[index] not in (None, ''))

    sample = [
        {columns[i]: (r[i] if i < len(r) else None) for i in range(len(columns)) if columns[i]}
        for r in body[:sample_rows]
    ]
    return {'rows': len(body), 'columns': columns, 'filled': filled, 'sample': sample}


# ---------------------------------------------------------------- 阶段适配器

def _task(task: dict | None = None) -> tuple[dict, dict]:
    """模块要求 (task_dict, task_id)，返回容器与容器里的那个任务。"""
    task = task if task is not None else {}
    return {'t': task}, task


def run_normalize(ctx: dict, source: Path, target: Path, task: dict | None = None) -> dict:
    from core_tools.normalizer import Normalizer

    module = Normalizer(speakers=ctx['game'].get('speakers', ()), language=ctx.get('language'))
    tasks, task = _task(task)
    module.process(str(source), str(target), tasks, 't')
    return task


def run_classify(ctx: dict, source: Path, target: Path, task: dict | None = None) -> dict:
    from core_tools.key_classifier import KeyClassifier

    rules = ROOT / ctx['game']['key_rules']
    module = KeyClassifier(str(rules), language=ctx.get('language'))
    tasks, task = _task(task)
    module.process(str(source), str(target), tasks, 't')
    return task


def run_context(ctx: dict, source: Path, target: Path, task: dict | None = None) -> dict:
    from core_tools.context_processor import ContextProcessor

    module = ContextProcessor(language=ctx.get('language'))
    reference = ctx['game'].get('context', {}).get('reference')
    guide = ctx['game'].get('context', {}).get('guide')
    tasks, task = _task(task)
    kwargs = {'analyze_dialogue': True, 'config_path': ctx.get('config_path'), 'materials': ctx.get('materials'),
              'scene_cache_dir': Path(ctx.get('batch_dir') or target.parent) / 'context_scenes'}
    if reference:
        kwargs['reference_path'] = ROOT / reference
    if guide:
        kwargs['guide_path'] = ROOT / guide
    module.process(str(source), str(target), tasks, 't', **kwargs)
    return task


def run_term(ctx: dict, source: Path, target: Path, task: dict | None = None) -> dict:
    from core_tools.term_scanner import TermScanner

    termbases = ctx['language'].get('termbases', [])
    module = TermScanner(str(ROOT / termbases[0]) if termbases else None, language=ctx.get('language'))
    tasks, task = _task(task)
    module.process(str(source), str(target), tasks, 't')
    return task


def run_term_filter(ctx: dict, source: Path, target: Path, task: dict | None = None) -> dict:
    from core_tools.term_filter import TermFilter

    module = TermFilter(config_path=ctx.get('config_path'), materials=ctx.get('materials'), language=ctx.get('language'))
    tasks, task = _task(task)
    skip_terms = 'termbases' in ctx.get('language', {}) and not ctx['language']['termbases']
    module.process(str(source), str(target), tasks, 't', skip_terms=skip_terms)
    if task.get('state') == 'completed' and skip_terms:
        reason = '未提供术语库，已跳过术语筛选，生效术语和无效术语列留空'
        task.update(skipped=True, skip_reason=reason, status=reason)
    return task


def run_tm(ctx: dict, source: Path, target: Path, task: dict | None = None) -> dict:
    from core_tools.tm_matcher import TMMatcher

    memories = ctx['language'].get('tms', []) if ctx['language'].get('use_tm', True) else []
    module = TMMatcher(str(ROOT / memories[0]) if memories else None, semantic=True,
                       config_path=ctx.get('config_path'), cache_dir=ctx.get('tm_cache_dir'), language=ctx.get('language'))
    tasks, task = _task(task)
    module.process(str(source), str(target), tasks, 't')
    if task.get('state') == 'completed' and not memories:
        reason = ('未提供记忆库，已跳过 TM 检索，TM参考列留空' if ctx['language'].get('use_tm', True)
                  else '本项目选择不使用 TM，已跳过检索，TM参考列留空')
        task.update(skipped=True, skip_reason=reason, status=reason)
    return task


def run_preflight(ctx: dict, source: Path, target: Path, task: dict | None = None) -> dict:
    """可选节点：工作台暂未启用，pipeline_service.PREFLIGHT_ENABLED=True 时使用。"""
    from core_tools.preflight import process as preflight_process

    tasks, task = _task(task)
    batch_dir = Path(ctx.get('batch_dir') or target.parent)
    preflight_process(batch_dir / '00-规范化.xlsx', source, tasks, 't',
                      batch_name=ctx.get('batch_name') or batch_dir.name, language=ctx.get('language'))
    return task


STAGES = (
    ('normalize', '规范化', '00-规范化.xlsx', run_normalize),
    ('classify', '分类', '01-分类.xlsx', run_classify),
    ('context', '上下文', '02-上下文.xlsx', run_context),
    ('term', '术语', '03-术语.xlsx', run_term),
    ('term_filter', '术语筛选', '04-术语筛选.xlsx', run_term_filter),
    ('tm', 'TM', '05-TM.xlsx', run_tm),
)


# ---------------------------------------------------------------- 编排

class Pipeline:
    def __init__(self, game: str, language: str, batch: str):
        self.game = json.loads((ROOT / 'profiles' / 'game' / f'{game}.json').read_text('utf-8'))
        self.language = language_profile(language)
        self.language_name = self.language['key']
        self.batch = batch
        self.dir = WORKSPACE / batch
        self.dir.mkdir(parents=True, exist_ok=True)

    def context(self) -> dict:
        return {'game': self.game, 'language': self.language, 'batch_dir': str(self.dir)}

    def run(self, source: Path, enabled: set[str] | None = None,
            keep_going: bool = False, start_after: str | None = None) -> list[StageResult]:
        results: list[StageResult] = []
        current = source
        ctx = self.context()

        skipping = start_after is not None
        for key, title, artifact, runner in STAGES:
            if skipping:
                if artifact == start_after:
                    skipping = False
                    current = self.dir / artifact
                continue
            if enabled and key not in enabled:
                continue
            out = self.dir / artifact
            started = time.time()
            try:
                task = runner(ctx, current, out)
                state = task.get('state', 'completed')
                message = task.get('status', '')
                if state == 'error':
                    raise RuntimeError(message or '模块返回 error')
                results.append(StageResult(
                    key, title, 'completed', artifact, round(time.time() - started, 1),
                    message, workbook_preview(out) if out.suffix.lower() in ('.xlsx', '.xlsm') and out.exists() else {},
                ))
                current = out
            except Exception as error:  # noqa: BLE001 — 编排层兜住一切
                results.append(StageResult(
                    key, title, 'failed', artifact if out.exists() else None,
                    round(time.time() - started, 1), f'{error}',
                    workbook_preview(out) if out.suffix.lower() in ('.xlsx', '.xlsm') and out.exists() else {},
                ))
                if not keep_going:
                    break

        manifest = self.dir / 'run.json'
        manifest.write_text(json.dumps({
            'batch': self.batch,
            'game': self.game.get('name'),
            'language': self.language_name,
            'source': str(source),
            'stages': [asdict(r) for r in results],
        }, ensure_ascii=False, indent=2), 'utf-8')

        return results


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description='翻译前准备流水线')
    parser.add_argument('source', help='待译文件（xlsx/csv）')
    parser.add_argument('--game', default='wc')
    profiles = language_profiles()
    parser.add_argument('--lang', choices=tuple(profiles),
                        default=next(key for key, profile in profiles.items() if profile['default']))
    parser.add_argument('--termbase', help='当前语向术语库路径')
    tm_options = parser.add_mutually_exclusive_group()
    tm_options.add_argument('--tm', help='当前语向 TMX 路径')
    tm_options.add_argument('--no-tm', action='store_true', help='明确不使用 TM，参考列留空')
    parser.add_argument('--source-column')
    parser.add_argument('--target-column')
    parser.add_argument('--emotion-column')
    parser.add_argument('--batch', default=None, help='批次名，默认取文件名')
    parser.add_argument('--stages', default=None, help='只跑这些阶段，逗号分隔')
    parser.add_argument('--keep-going', action='store_true', help='某阶段失败也继续')
    parser.add_argument('--from', dest='resume', default=None,
                        help='从已有产物续跑，如 04-术语筛选.xlsx')
    args = parser.parse_args(argv)

    source = Path(args.source).expanduser().resolve()
    if not source.exists():
        print(f'找不到文件：{source}', file=sys.stderr)
        return 2

    enabled = set(args.stages.split(',')) if args.stages else None
    pipeline = Pipeline(args.game, args.lang, args.batch or source.stem)
    if args.termbase:
        pipeline.language['termbases'] = [str(Path(args.termbase).expanduser().resolve())]
    if args.tm:
        pipeline.language['tms'] = [str(Path(args.tm).expanduser().resolve())]
    if args.no_tm:
        pipeline.language.update(tms=[], use_tm=False)
    mapping = {k: v for k, v in {'source': args.source_column, 'target': args.target_column,
                                 'emotion': args.emotion_column}.items() if v}
    pipeline.language.update(language_pair(pipeline.language, mapping).as_dict())
    print(f'批次 {pipeline.batch}  →  {pipeline.dir}\n')

    results = pipeline.run(source, enabled=enabled, keep_going=args.keep_going,
                           start_after=args.resume)

    width = max((len(r.title) for r in results), default=4)
    for r in results:
        mark = {'completed': '完成', 'failed': '失败', 'skipped': '跳过'}[r.state]
        rows = r.preview.get('rows')
        print(f'  {r.title:<{width}}  {mark}  {r.seconds:>5.1f}s  '
              f'{rows if rows is not None else "-":>5} 行  {r.message[:60]}')

    failed = [r for r in results if r.state == 'failed']
    if failed:
        print(f'\n阻断于：{failed[0].title}')
        return 1

    last = results[-1] if results else None
    if last:
        print(f'\n产物：{pipeline.dir / last.artifact}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
